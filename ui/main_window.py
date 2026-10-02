
import os
import json
import logging
from pathlib import Path
from datetime import datetime

import ipaddress
import re
import socket
import subprocess
import threading
from collections import deque
from concurrent.futures import FIRST_COMPLETED, ThreadPoolExecutor, wait

from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

from PySide6.QtCore import Qt, QObject, QThread, Signal

from PySide6.QtWidgets import (
    QAbstractItemView,
    QComboBox,
    QCheckBox,
    QDialog,
    QDialogButtonBox,
    QFormLayout,
    QGroupBox,
    QHBoxLayout,
    QLabel,
    QLineEdit,
    QMainWindow,
    QMessageBox,
    QPushButton,
    QSplitter,
    QStackedWidget,
    QTabWidget,
    QTextEdit,
    QScrollArea,
    QTreeWidget,
    QTreeWidgetItem,
    QVBoxLayout,
    QWidget,
    QFileDialog,
    QTableWidget,
    QTableWidgetItem,
)

from database import SessionLocal
from models import (
    Admin,
    Computer,
    ComputerStatus,
    Department,
    EquipmentAsset,
    Maintenance,
    Phone,
    PhoneStatus,
    Telephone,
    TelephoneStatus,
)
from security import set_admin_password
from sqlalchemy.orm import joinedload
from werkzeug.security import generate_password_hash

import pandas as pd

from services.device_enrichment_service import collect_device_enrichment
from services.excel_mapping_service import (
    build_mapping_from_profile,
    detect_auto_mapping,
    detect_mapping_profile,
    read_excel_dataframe,
)
from services.phone_service import (
    assign_phone,
    create_phone,
    delete_phone,
    get_phone,
    get_phones,
    move_phone_to_stock,
    reserve_phone,
    update_phone,
)
from services.equipment_type_service import (
    EQUIPMENT_STATUSES,
    create_equipment_type,
    get_equipment_filter_options,
    get_enabled_equipment_types,
    get_equipment_types,
    get_equipment_page_kind,
    set_equipment_type_enabled,
)
from services.telephone_service import (
    assign_telephone,
    create_telephone,
    delete_telephone,
    get_telephone,
    get_telephones,
    move_telephone_to_stock,
    import_telephones,
    reserve_telephone,
    update_telephone,
)
from services.equipment_asset_service import (
    STATUSES as EQUIPMENT_ASSET_STATUSES,
    assign_equipment_asset,
    create_equipment_asset,
    delete_equipment_asset,
    get_equipment_asset,
    get_equipment_assets,
    import_equipment_assets,
    move_equipment_asset_to_stock,
    reserve_equipment_asset,
    update_equipment_asset,
)
from services.equipment_report_service import (
    get_enabled_department_summary,
    get_enabled_equipment_summary,
    get_enabled_report_types,
    get_equipment_report_rows,
)
from ui.excel_mapping_dialog import ExcelMappingDialog


def normalize_host_name(value):

    if value is None:
        return None

    normalized = str(value).strip().strip(".").lower()

    if not normalized:
        return None

    return normalized.split(".")[0]


def normalize_ip_address(value):

    if value is None:
        return None

    normalized = str(value).strip()

    if not normalized:
        return None

    try:
        return str(ipaddress.ip_address(normalized))
    except ValueError:
        return normalized.lower()


def normalize_mac_address(value):

    if not value:
        return None

    cleaned = re.sub(
        r"[^0-9A-Fa-f]",
        "",
        value,
    )

    if len(cleaned) != 12:
        return None

    return "-".join(
        cleaned[i:i + 2].upper()
        for i in range(0, len(cleaned), 2)
    )


def run_windows_command(args, timeout=None):

    kwargs = {
        "capture_output": True,
        "text": True,
        "check": False,
    }

    if timeout is not None:
        kwargs["timeout"] = timeout

    if os.name == "nt":
        startup_info = subprocess.STARTUPINFO()
        startup_info.dwFlags |= subprocess.STARTF_USESHOWWINDOW
        startup_info.wShowWindow = subprocess.SW_HIDE
        kwargs["startupinfo"] = startup_info
        kwargs["creationflags"] = getattr(
            subprocess,
            "CREATE_NO_WINDOW",
            0,
        )

    return subprocess.run(
        args,
        **kwargs,
    )


def is_virtual_adapter_name(adapter_name):

    lowered = (adapter_name or "").lower()

    virtual_markers = (
        "docker",
        "hyper-v",
        "vethernet",
        "virtualbox",
        "vmware",
        "loopback",
        "isatap",
        "teredo",
        "tunnel",
        "bluetooth",
        "hamachi",
        "wireguard",
        "ppp",
    )

    return any(
        marker in lowered
        for marker in virtual_markers
    )


def get_ipconfig_adapter_details():

    try:
        result = run_windows_command(
            ["ipconfig", "/all"]
        )
    except Exception:
        return []

    adapters = []
    current_adapter = None

    for raw_line in result.stdout.splitlines():

        stripped = raw_line.strip()

        if stripped.endswith(":") and any(
            stripped.lower().startswith(prefix)
            for prefix in (
                "ethernet adapter ",
                "wireless lan adapter ",
                "wi-fi",
                "wi-fi adapter ",
            )
        ):
            if current_adapter is not None:
                adapters.append(current_adapter)

            current_adapter = {
                "name": stripped[:-1],
                "lines": [],
            }
            continue

        if current_adapter is not None:
            current_adapter["lines"].append(
                raw_line
            )

    if current_adapter is not None:
        adapters.append(current_adapter)

    details = []

    for adapter in adapters:

        adapter_info = {
            "name": adapter["name"],
            "description": None,
            "ipv4": None,
            "subnet_mask": None,
            "gateway": None,
            "mac": None,
            "media_state": None,
            "is_virtual": is_virtual_adapter_name(adapter["name"]),
        }

        for line in adapter["lines"]:

            stripped = line.strip()

            if stripped.startswith("Description"):
                adapter_info["description"] = (
                    stripped.split(":", 1)[1].strip()
                )

            elif stripped.startswith("Physical Address"):
                value = stripped.split(":", 1)[1].strip()
                adapter_info["mac"] = normalize_mac_address(
                    value
                )

            elif stripped.startswith("IPv4 Address"):
                value = stripped.split(":", 1)[1].strip()
                adapter_info["ipv4"] = normalize_ip_address(
                    value.split("(", 1)[0].strip()
                )

            elif stripped.startswith("Subnet Mask"):
                value = stripped.split(":", 1)[1].strip()
                adapter_info["subnet_mask"] = normalize_ip_address(
                    value.split("(", 1)[0].strip()
                )

            elif stripped.startswith("Default Gateway"):
                value = stripped.split(":", 1)[1].strip()
                if value and value != "Not Available":
                    adapter_info["gateway"] = normalize_ip_address(
                        value.split("(", 1)[0].strip()
                    )

            elif stripped.startswith("Media State"):
                adapter_info["media_state"] = (
                    stripped.split(":", 1)[1].strip()
                )

        if adapter_info["ipv4"] and adapter_info["subnet_mask"]:
            try:
                adapter_info["network"] = str(
                    ipaddress.IPv4Network(
                        f"{adapter_info['ipv4']}/{adapter_info['subnet_mask']}",
                        strict=False,
                    )
                )
            except Exception:
                adapter_info["network"] = None

        details.append(adapter_info)

    return details


def get_active_network_adapter():

    adapters = get_ipconfig_adapter_details()

    if not adapters:
        return None

    candidates = []

    for adapter in adapters:

        ipv4 = adapter.get("ipv4")
        media_state = (adapter.get("media_state") or "").lower()

        if not ipv4:
            continue

        if "disconnected" in media_state:
            continue

        if adapter.get("is_virtual"):
            continue

        candidates.append(adapter)

    if not candidates:
        for adapter in adapters:
            ipv4 = adapter.get("ipv4")
            media_state = (adapter.get("media_state") or "").lower()

            if not ipv4:
                continue

            if "disconnected" in media_state:
                continue

            candidates.append(adapter)

    if not candidates:
        return None

    candidates.sort(
        key=lambda item: (
            0 if item.get("gateway") else 1,
            0 if not item.get("is_virtual") else 1,
            item.get("name", ""),
        )
    )

    return candidates[0]


def get_default_network_range():

    active_adapter = get_active_network_adapter()

    if active_adapter and active_adapter.get("network"):
        return active_adapter["network"]

    try:
        hostname = socket.gethostname()
        local_ip = socket.gethostbyname(hostname)
    except Exception:
        local_ip = None

    if local_ip and local_ip != "127.0.0.1":
        try:
            gateway_ip = ipaddress.ip_address(local_ip)
            base_address = ".".join(
                str(gateway_ip).split(".")[:-1]
            )
            return f"{base_address}.0/24"
        except Exception:
            pass

    session = SessionLocal()

    try:
        computers = session.query(
            Computer
        ).filter(
            Computer.ip_address.isnot(None)
        ).all()

        for computer in computers:
            if computer.ip_address:
                try:
                    ip_obj = ipaddress.ip_address(computer.ip_address)
                    base_address = ".".join(
                        str(ip_obj).split(".")[:-1]
                    )
                    return f"{base_address}.0/24"
                except Exception:
                    continue

    finally:
        session.close()

    return "192.168.1.0/24"


def ping_ip(ip_address):

    try:
        result = run_windows_command(
            [
                "ping",
                "-n",
                "1",
                "-w",
                "300",
                ip_address,
            ]
        )
        return result.returncode == 0
    except Exception:
        return False


def get_arp_entries():

    try:
        result = run_windows_command(
            ["arp", "-a"]
        )
    except Exception:
        result = None

    entries = {}

    if result is not None:

        for line in result.stdout.splitlines():

            stripped = line.strip()

            if not stripped:
                continue

            if stripped.startswith("Interface:") or stripped.startswith("Internet Address"):
                continue

            parts = stripped.split()

            if len(parts) >= 2 and parts[0].count(".") == 3:
                ip_address = normalize_ip_address(parts[0])

                if ip_address:
                    mac_address = normalize_mac_address(parts[1])

                    if mac_address:
                        entries[ip_address] = mac_address

    adapter_entries = {}

    for adapter in get_ipconfig_adapter_details():

        if adapter.get("ipv4") and adapter.get("mac"):
            adapter_entries[adapter["ipv4"]] = adapter["mac"]

    if adapter_entries:
        entries.update(adapter_entries)

    return entries


def parse_ping_hostname(output):

    if not output:
        return None

    match = re.search(
        r"Pinging\s+([^\s]+)\s+\[([^\]]+)\]",
        output,
        re.IGNORECASE,
    )

    if not match:
        return None

    return normalize_host_name(match.group(1))


def parse_nbtstat_hostname(output):

    if not output:
        return None

    match = re.search(
        r"\b([A-Za-z0-9][A-Za-z0-9-]*)\s+<00>\s+UNIQUE",
        output,
        re.IGNORECASE,
    )

    if not match:
        return None

    return normalize_host_name(match.group(1))


def parse_ipconfig_dns_cache(output, ip_address=None):

    if not output:
        return None

    record_name = None

    for raw_line in output.splitlines():

        stripped = raw_line.strip()

        if not stripped:
            continue

        lowered = stripped.lower()

        if lowered.startswith("record name"):
            record_name = stripped.split(":", 1)[1].strip()
            continue

        if lowered.startswith("a (host) record"):
            host_ip = stripped.split(":", 1)[1].strip()

            if record_name and (
                ip_address is None
                or normalize_ip_address(host_ip) == normalize_ip_address(ip_address)
            ):
                return normalize_host_name(record_name)

            record_name = None

    return None


def parse_nslookup_hostname(output):

    if not output:
        return None

    matches = re.findall(
        r"(?:Name|name)\s*[:=]\s*([A-Za-z0-9][A-Za-z0-9.-]+)",
        output,
        flags=re.IGNORECASE,
    )

    for candidate in matches:
        normalized = normalize_host_name(candidate)
        if normalized:
            return normalized

    return None


def get_hostname_for_ip(ip_address):

    normalized_ip = normalize_ip_address(ip_address)

    if not normalized_ip:
        return None

    try:
        host_info = socket.gethostbyaddr(normalized_ip)
        if host_info and host_info[0]:
            return normalize_host_name(host_info[0])
    except Exception:
        pass

    try:
        nslookup_result = run_windows_command(
            ["nslookup", normalized_ip],
            timeout=1,
        )
    except Exception:
        nslookup_result = None

    if nslookup_result is not None:
        nslookup_output = (nslookup_result.stdout or "") + (nslookup_result.stderr or "")
        parsed_hostname = parse_nslookup_hostname(nslookup_output)
        if parsed_hostname:
            return parsed_hostname

    try:
        result = run_windows_command(
            [
                "ping",
                "-a",
                "-n",
                "1",
                "-w",
                "400",
                normalized_ip,
            ],
            timeout=1,
        )
    except Exception:
        result = None

    if result is not None:
        output = (result.stdout or "") + (result.stderr or "")
        parsed_hostname = parse_ping_hostname(output)
        if parsed_hostname:
            return parsed_hostname

    try:
        nbt_result = run_windows_command(
            ["nbtstat", "-A", normalized_ip],
            timeout=1,
        )
    except Exception:
        nbt_result = None

    if nbt_result is not None:
        nbt_output = (nbt_result.stdout or "") + (nbt_result.stderr or "")
        parsed_hostname = parse_nbtstat_hostname(nbt_output)
        if parsed_hostname:
            return parsed_hostname

    try:
        dns_result = run_windows_command(
            ["ipconfig", "/displaydns"],
            timeout=2,
        )
    except Exception:
        dns_result = None

    if dns_result is not None:
        dns_output = (dns_result.stdout or "") + (dns_result.stderr or "")
        parsed_hostname = parse_ipconfig_dns_cache(dns_output, normalized_ip)
        if parsed_hostname:
            return parsed_hostname

    return None


def get_network_targets(network_range):

    try:
        network = ipaddress.ip_network(network_range, strict=False)
    except Exception:
        return []

    return [str(ip) for ip in network.hosts()]


def classify_discovered_device(
    ip_address,
    hostname,
    mac_address,
    existing_computers,
):

    normalized_ip = normalize_ip_address(ip_address)
    normalized_hostname = normalize_host_name(hostname)
    normalized_mac = normalize_mac_address(mac_address)

    computer_by_ip = {
        normalize_ip_address(computer.ip_address): computer
        for computer in existing_computers
        if normalize_ip_address(computer.ip_address)
    }

    computer_by_mac = {
        normalize_mac_address(computer.mac_address): computer
        for computer in existing_computers
        if normalize_mac_address(computer.mac_address)
    }

    computer_by_name = {
        normalize_host_name(computer.computer_name): computer
        for computer in existing_computers
        if normalize_host_name(computer.computer_name)
    }

    matched_computer = None
    match_basis = None

    if normalized_mac and normalized_mac in computer_by_mac:
        matched_computer = computer_by_mac[normalized_mac]
        match_basis = "MAC"
    elif normalized_ip and normalized_ip in computer_by_ip:
        matched_computer = computer_by_ip[normalized_ip]
        match_basis = "IP"
    elif normalized_hostname and normalized_hostname in computer_by_name:
        matched_computer = computer_by_name[normalized_hostname]
        match_basis = "Hostname"

    if matched_computer is None:

        if hostname or mac_address or ip_address:
            return {
                "status": "NEW_DEVICE",
                "reason": "No matching computer found in the registered inventory for the discovered device.",
                "matched_computer": None,
            }

        return {
            "status": "UNKNOWN",
            "reason": "No identifiable device data was discovered.",
            "matched_computer": None,
        }

    registered_mac = normalize_mac_address(
        matched_computer.mac_address
    )
    registered_ip = normalize_ip_address(
        matched_computer.ip_address
    )
    registered_hostname = normalize_host_name(
        matched_computer.computer_name
    )

    mismatches = []

    if normalized_mac and registered_mac and normalized_mac != registered_mac:
        mismatches.append("MAC address differs")

    if normalized_ip and registered_ip and normalized_ip != registered_ip:
        mismatches.append("IP address differs")

    if normalized_hostname and registered_hostname and normalized_hostname != registered_hostname:
        if match_basis in {"MAC", "IP"}:
            pass
        else:
            mismatches.append("Computer name differs")

    if mismatches:
        return {
            "status": "MISMATCH",
            "reason": "; ".join(mismatches),
            "matched_computer": matched_computer,
        }

    evidence = []

    if normalized_mac and registered_mac and normalized_mac == registered_mac:
        evidence.append("MAC")

    if normalized_ip and registered_ip and normalized_ip == registered_ip:
        evidence.append("IP")

    if normalized_hostname and registered_hostname and normalized_hostname == registered_hostname:
        evidence.append("hostname")

    if evidence:

        if set(evidence) == {"MAC", "IP", "hostname"}:
            reason = "MAC, hostname and IP match the registered computer."
        else:
            safe_evidence = [item for item in ["MAC", "IP", "hostname"] if item in evidence]
            reason = f"{', '.join(safe_evidence)} match the registered computer."

        return {
            "status": "VERIFIED",
            "reason": reason,
            "matched_computer": matched_computer,
        }

    return {
        "status": "VERIFIED",
        "reason": "The discovered device matches a registered computer record.",
        "matched_computer": matched_computer,
    }


def enrich_discovery_results(results):

    if not results:
        return results

    enrichment_fields = [
        "hostname",
        "device_type",
        "manufacturer",
        "model",
        "serial_number",
        "operating_system",
        "windows_version",
        "windows_build",
        "cpu",
        "ram",
        "architecture",
        "domain",
        "workgroup",
        "logged_in_user",
        "bios",
        "adapter_name",
        "adapter_vendor",
        "link_speed",
        "installed_software",
        "access_status",
        "enrichment_status",
    ]

    for item in results:

        ip_address = item.get("ip_address")

        if not ip_address:
            continue

        try:
            payload = collect_device_enrichment(
                ip_address=item.get("ip_address"),
                mac_address=item.get("mac_address"),
                hostname=item.get("hostname"),
            )
        except Exception:
            payload = {
                "access_status": "Unavailable",
                "enrichment_status": "Unavailable",
            }

        device_enrichment = {
            key: payload.get(key)
            for key in enrichment_fields
            if payload.get(key) is not None
        }

        item["device_enrichment"] = device_enrichment

        for key in enrichment_fields:
            if key in device_enrichment:
                item[key] = device_enrichment[key]

        item["access_status"] = device_enrichment.get(
            "access_status",
            "Unavailable",
        )
        item["enrichment_status"] = device_enrichment.get(
            "enrichment_status",
            "Unavailable",
        )

    return results


def _merge_discovery_entry(
    discovered_entries,
    ip_address,
    hostname=None,
    mac_address=None,
    computers=None,
):

    normalized_ip = normalize_ip_address(ip_address)
    normalized_hostname = normalize_host_name(hostname)
    normalized_mac = normalize_mac_address(mac_address)

    if not normalized_ip and not normalized_mac and not normalized_hostname:
        return discovered_entries

    existing_entry = None
    existing_key = None

    if normalized_ip:
        for key, entry in discovered_entries.items():
            if entry.get("ip_address") == normalized_ip:
                existing_entry = entry
                existing_key = key
                break

    if existing_entry is None and normalized_mac:
        for key, entry in discovered_entries.items():
            if entry.get("mac_address") == normalized_mac:
                existing_entry = entry
                existing_key = key
                break

    if existing_entry is None and normalized_hostname:
        for key, entry in discovered_entries.items():
            if entry.get("hostname") == normalized_hostname:
                existing_entry = entry
                existing_key = key
                break

    if existing_entry is None:
        if normalized_ip:
            existing_key = normalized_ip
        elif normalized_mac:
            existing_key = normalized_mac
        else:
            existing_key = normalized_hostname

        discovered_entries[existing_key] = {
            "computer_name": None,
            "hostname": None,
            "ip_address": normalized_ip,
            "mac_address": normalized_mac,
            "status": "UNKNOWN",
            "reason": "",
            "matched_computer": None,
            "department": "Unassigned",
            "notes": "",
            "is_known_record": False,
        }

    entry = discovered_entries[existing_key]

    if normalized_ip:
        entry["ip_address"] = normalized_ip

    if normalized_mac:
        entry["mac_address"] = normalized_mac

    if normalized_hostname:
        entry["hostname"] = normalized_hostname

    verification = classify_discovered_device(
        entry.get("ip_address"),
        entry.get("hostname"),
        entry.get("mac_address"),
        computers or [],
    )

    matched_computer = verification.get("matched_computer")

    entry["status"] = verification.get("status", "UNKNOWN")
    entry["reason"] = verification.get("reason", "")
    entry["matched_computer"] = (
        matched_computer.computer_name
        if matched_computer else None
    )
    entry["department"] = (
        matched_computer.department.name
        if matched_computer and matched_computer.department
        else "Unassigned"
    )
    entry["notes"] = (
        matched_computer.notes
        if matched_computer
        else "Discovered on the local network."
    )
    entry["is_known_record"] = matched_computer is not None
    entry["computer_name"] = (
        matched_computer.computer_name
        if matched_computer
        else entry.get("hostname") or f"Unknown Device {entry.get('ip_address') or 'Unknown'}"
    )

    return discovered_entries


def _snapshot_discovery_entries(discovered_entries):

    snapshot = []

    for entry in discovered_entries.values():
        snapshot.append(
            {
                "computer_name": entry.get("computer_name") or "",
                "hostname": entry.get("hostname") or "",
                "ip_address": entry.get("ip_address") or "",
                "mac_address": entry.get("mac_address") or "",
                "status": entry.get("status") or "UNKNOWN",
                "reason": entry.get("reason") or "",
                "matched_computer": entry.get("matched_computer"),
                "department": entry.get("department") or "Unassigned",
                "notes": entry.get("notes") or "",
                "is_known_record": entry.get("is_known_record", False),
            }
        )

    snapshot.sort(
        key=lambda item: (
            item.get("ip_address") or "",
            item.get("hostname") or "",
            item.get("computer_name") or "",
        )
    )

    return snapshot


def collect_network_discovery_results(
    network_range,
    progress_callback=None,
    device_callback=None,
    stop_event=None,
    max_workers=32,
):

    session = SessionLocal()

    try:
        computers = (
            session.query(
                Computer
            )
            .options(
                joinedload(Computer.department)
            )
            .order_by(
                Computer.computer_name
            )
            .all()
        )
    finally:
        session.close()

    resolved_network_range = network_range

    try:
        ipaddress.ip_network(resolved_network_range, strict=False)
    except Exception:
        resolved_network_range = get_default_network_range()

    targets = get_network_targets(resolved_network_range)
    total_targets = len(targets)
    arp_entries = get_arp_entries()

    if stop_event is None:
        stop_event = threading.Event()

    max_workers = max(1, min(int(max_workers or 32), 32))
    discovered_entries = {}

    for ip_address in sorted(arp_entries):
        if ip_address not in targets:
            continue

        _merge_discovery_entry(
            discovered_entries,
            ip_address,
            None,
            arp_entries.get(ip_address),
            computers,
        )

    if device_callback is not None:
        device_callback(
            _snapshot_discovery_entries(discovered_entries)
        )

    if progress_callback is not None:
        progress_value = (
            int((len(arp_entries) / total_targets) * 100)
            if total_targets
            else 100
        )
        progress_callback(
            "ARP Discovery",
            min(progress_value, 100),
            len(discovered_entries),
        )

    pending_targets = deque(
        ip_address
        for ip_address in targets
        if ip_address not in arp_entries
    )

    scanned_count = len(arp_entries)
    running = {}

    def ping_ip_task(ip_address):
        return ping_ip(ip_address)

    with ThreadPoolExecutor(max_workers=max_workers) as executor:

        while pending_targets or running:

            if stop_event.is_set():
                break

            while len(running) < max_workers and pending_targets:
                ip_address = pending_targets.popleft()
                running[executor.submit(ping_ip_task, ip_address)] = ip_address

            if not running:
                break

            done, _ = wait(
                running,
                return_when=FIRST_COMPLETED,
            )

            for future in done:
                ip_address = running.pop(future)
                scanned_count += 1

            try:
                reachable = future.result(timeout=0)
            except Exception:
                reachable = False

            if reachable:
                hostname = get_hostname_for_ip(ip_address)

                refreshed_arp_entries = get_arp_entries()
                mac_address = refreshed_arp_entries.get(ip_address)

                _merge_discovery_entry(
                    discovered_entries,
                    ip_address,
                    hostname,
                    mac_address,
                    computers,
                )

                if device_callback is not None:
                    device_callback(
                        _snapshot_discovery_entries(discovered_entries)
                    )

                if progress_callback is not None:
                    progress_value = (
                        int((scanned_count / total_targets) * 100)
                        if total_targets
                        else 100
                    )
                    progress_callback(
                        "Active Scan",
                        min(progress_value, 100),
                        len(discovered_entries),
                    )

    active_results = _snapshot_discovery_entries(discovered_entries)
    active_ips = {
        item["ip_address"]
        for item in active_results
        if item.get("ip_address")
    }

    offline_results = []

    for computer in computers:

        if not computer.ip_address:
            continue

        if normalize_ip_address(computer.ip_address) in active_ips:
            continue

        offline_results.append(
            {
                "computer_name": computer.computer_name,
                "hostname": computer.computer_name,
                "ip_address": computer.ip_address,
                "mac_address": computer.mac_address,
                "status": "OFFLINE",
                "reason": "Inventory record did not respond during the scan; cannot verify because the device is offline.",
                "department": (
                    computer.department.name
                    if computer.department
                    else "Unassigned"
                ),
                "notes": "Inventory record did not respond during the scan.",
                "is_known_record": True,
            }
        )

    results = active_results + offline_results
    results = enrich_discovery_results(results)

    summary = {
        "verified": sum(1 for item in results if item["status"] == "VERIFIED"),
        "mismatch": sum(1 for item in results if item["status"] == "MISMATCH"),
        "offline": sum(1 for item in results if item["status"] == "OFFLINE"),
        "new_device": sum(1 for item in results if item["status"] == "NEW_DEVICE"),
        "unknown": sum(1 for item in results if item["status"] == "UNKNOWN"),
        "total": len(results),
    }

    if progress_callback is not None:
        progress_callback(
            "Phase: Complete" if not stop_event.is_set() else "Phase: Stopped",
            100,
            len(active_results),
        )

    return results, summary


class NetworkScanWorker(QObject):

    progress = Signal(str, int, int)
    live_results = Signal(object)
    finished = Signal(object, object)
    error = Signal(str)

    def __init__(self, network_range):
        super().__init__()
        self.network_range = network_range
        self.stop_event = threading.Event()

    def request_stop(self):

        self.stop_event.set()

    def run(self):

        try:
            results, summary = collect_network_discovery_results(
                self.network_range,
                progress_callback=self.progress.emit,
                device_callback=self.live_results.emit,
                stop_event=self.stop_event,
            )
            self.finished.emit(results, summary)
        except Exception as exc:
            self.error.emit(str(exc))


from services.computer_service import (
    assign_computer,
    create_computer,
    delete_computer,
    approve_pending_device,
    get_pending_devices,
    get_computer_data,
    get_computers,
    move_to_stock,
    reserve_computer,
    update_computer,
)

from services.department_service import (
    create_department,
    delete_department,
    get_department,
    get_departments,
    update_department,
)


# =========================================================
# DEPARTMENT DIALOG
# =========================================================

class DepartmentDialog(QDialog):

    def __init__(self, department=None, parent=None):

        super().__init__(parent)

        self.department = department

        self.setWindowTitle(
            "Edit Department" if department else "Add Department"
        )

        self.setMinimumWidth(420)

        self.name_edit = QLineEdit()

        self.description_edit = QTextEdit()
        self.description_edit.setMinimumHeight(100)

        form = QFormLayout()

        form.addRow(
            "Name:",
            self.name_edit,
        )

        form.addRow(
            "Description:",
            self.description_edit,
        )

        if department:

            self.name_edit.setText(
                department.name
            )

            self.description_edit.setPlainText(
                department.description or ""
            )

        buttons = QDialogButtonBox(
            QDialogButtonBox.Save
            | QDialogButtonBox.Cancel
        )

        buttons.accepted.connect(
            self.accept
        )

        buttons.rejected.connect(
            self.reject
        )

        layout = QVBoxLayout(self)

        layout.addLayout(form)
        layout.addWidget(buttons)

    def get_data(self):

        return {
            "name": self.name_edit.text().strip(),
            "description": (
                self.description_edit
                .toPlainText()
                .strip()
            ),
        }


# =========================================================
# COMPUTER DIALOG
# =========================================================

class ComputerDialog(QDialog):

    def __init__(
        self,
        departments,
        computer=None,
        parent=None,
    ):

        super().__init__(parent)

        self.computer = computer

        self.setWindowTitle(
            "Edit Computer"
            if computer
            else "Add Computer"
        )

        self.setMinimumWidth(560)

        # -------------------------------------------------
        # BASIC
        # -------------------------------------------------

        self.name_edit = QLineEdit()
        self.manufacturer_edit = QLineEdit()
        self.model_edit = QLineEdit()
        self.serial_edit = QLineEdit()

        # -------------------------------------------------
        # OPERATING SYSTEM
        # -------------------------------------------------

        self.os_edit = QLineEdit()
        self.os_version_edit = QLineEdit()

        # -------------------------------------------------
        # HARDWARE
        # -------------------------------------------------

        self.cpu_edit = QLineEdit()
        self.ram_edit = QLineEdit()
        self.storage_edit = QLineEdit()

        # -------------------------------------------------
        # NETWORK
        # -------------------------------------------------

        self.ip_edit = QLineEdit()
        self.mac_edit = QLineEdit()

        # -------------------------------------------------
        # DEPARTMENT
        # -------------------------------------------------

        self.department_combo = QComboBox()

        self.department_combo.addItem(
            "No Department",
            None,
        )

        for department in departments:

            self.department_combo.addItem(
                department["name"],
                department["id"],
            )

        # -------------------------------------------------
        # ASSIGNMENT
        # -------------------------------------------------

        self.user_edit = QLineEdit()
        self.location_edit = QLineEdit()

        # -------------------------------------------------
        # STATUS
        # -------------------------------------------------

        self.status_combo = QComboBox()

        self.status_combo.addItem(
            "In Stock",
            ComputerStatus.IN_STOCK.value,
        )

        self.status_combo.addItem(
            "Assigned",
            ComputerStatus.ASSIGNED.value,
        )

        self.status_combo.addItem(
            "Reserved",
            ComputerStatus.RESERVED.value,
        )

        self.status_combo.addItem(
            "Maintenance",
            ComputerStatus.MAINTENANCE.value,
        )

        self.status_combo.addItem(
            "Retired",
            ComputerStatus.RETIRED.value,
        )

        # -------------------------------------------------
        # OTHER
        # -------------------------------------------------

        self.warranty_edit = QLineEdit()

        self.notes_edit = QTextEdit()
        self.notes_edit.setMinimumHeight(90)

        # -------------------------------------------------
        # FORM
        # -------------------------------------------------

        form = QFormLayout()

        form.addRow(
            "Computer Name:",
            self.name_edit,
        )

        form.addRow(
            "Manufacturer:",
            self.manufacturer_edit,
        )

        form.addRow(
            "Model:",
            self.model_edit,
        )

        form.addRow(
            "Serial Number:",
            self.serial_edit,
        )

        form.addRow(
            "Operating System:",
            self.os_edit,
        )

        form.addRow(
            "OS Version:",
            self.os_version_edit,
        )

        form.addRow(
            "CPU:",
            self.cpu_edit,
        )

        form.addRow(
            "RAM:",
            self.ram_edit,
        )

        form.addRow(
            "Storage:",
            self.storage_edit,
        )

        form.addRow(
            "IP Address:",
            self.ip_edit,
        )

        form.addRow(
            "MAC Address:",
            self.mac_edit,
        )

        form.addRow(
            "Department:",
            self.department_combo,
        )

        form.addRow(
            "Assigned User:",
            self.user_edit,
        )

        form.addRow(
            "Location:",
            self.location_edit,
        )

        form.addRow(
            "Status:",
            self.status_combo,
        )

        form.addRow(
            "Warranty:",
            self.warranty_edit,
        )

        form.addRow(
            "Notes:",
            self.notes_edit,
        )

        if computer:

            self.load_computer(
                computer
            )

        buttons = QDialogButtonBox(
            QDialogButtonBox.Save
            | QDialogButtonBox.Cancel
        )

        buttons.accepted.connect(
            self.accept
        )

        buttons.rejected.connect(
            self.reject
        )

        layout = QVBoxLayout(self)

        layout.addLayout(form)
        layout.addWidget(buttons)

    # -----------------------------------------------------
    # LOAD COMPUTER
    # -----------------------------------------------------

    def load_computer(self, computer):

        self.name_edit.setText(
            computer.get("computer_name") or ""
        )

        self.manufacturer_edit.setText(
            computer.get("manufacturer") or ""
        )

        self.model_edit.setText(
            computer.get("model") or ""
        )

        self.serial_edit.setText(
            computer.get("serial_number") or ""
        )

        self.os_edit.setText(
            computer.get("operating_system") or ""
        )

        self.os_version_edit.setText(
            computer.get("os_version") or ""
        )

        self.cpu_edit.setText(
            computer.get("cpu") or ""
        )

        self.ram_edit.setText(
            computer.get("ram") or ""
        )

        self.storage_edit.setText(
            computer.get("storage") or ""
        )

        self.ip_edit.setText(
            computer.get("ip_address") or ""
        )

        self.mac_edit.setText(
            computer.get("mac_address") or ""
        )

        department_id = computer.get(
            "department_id"
        )

        index = self.department_combo.findData(
            department_id
        )

        if index >= 0:

            self.department_combo.setCurrentIndex(
                index
            )

        self.user_edit.setText(
            computer.get("assigned_user") or ""
        )

        self.location_edit.setText(
            computer.get("location") or ""
        )

        status = computer.get(
            "status"
        )

        index = self.status_combo.findData(
            status
        )

        if index >= 0:

            self.status_combo.setCurrentIndex(
                index
            )

        self.warranty_edit.setText(
            computer.get("warranty") or ""
        )

        self.notes_edit.setPlainText(
            computer.get("notes") or ""
        )

    # -----------------------------------------------------
    # GET DATA
    # -----------------------------------------------------

    def get_data(self):

        return {

            "computer_name":
                self.name_edit
                .text()
                .strip(),

            "manufacturer":
                self.manufacturer_edit
                .text()
                .strip()
                or None,

            "model":
                self.model_edit
                .text()
                .strip()
                or None,

            "serial_number":
                self.serial_edit
                .text()
                .strip()
                or None,

            "operating_system":
                self.os_edit
                .text()
                .strip()
                or None,

            "os_version":
                self.os_version_edit
                .text()
                .strip()
                or None,

            "cpu":
                self.cpu_edit
                .text()
                .strip()
                or None,

            "ram":
                self.ram_edit
                .text()
                .strip()
                or None,

            "storage":
                self.storage_edit
                .text()
                .strip()
                or None,

            "ip_address":
                self.ip_edit
                .text()
                .strip()
                or None,

            "mac_address":
                self.mac_edit
                .text()
                .strip()
                or None,

            "department_id":
                self.department_combo
                .currentData(),

            "assigned_user":
                self.user_edit
                .text()
                .strip()
                or None,

            "location":
                self.location_edit
                .text()
                .strip()
                or "IT Storage",

            "status":
                self.status_combo
                .currentData(),

            "purchase_date":
                None,

            "warranty":
                self.warranty_edit
                .text()
                .strip()
                or None,

            "notes":
                self.notes_edit
                .toPlainText()
                .strip()
                or None,
        }


class PendingApprovalDialog(QDialog):

    def __init__(self, departments, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Approve Pending Device")
        self.setMinimumWidth(420)

        self.department_combo = QComboBox()
        self.department_combo.addItem("No Department", None)
        for department in departments:
            self.department_combo.addItem(
                department["name"],
                department["id"],
            )

        self.user_edit = QLineEdit()
        self.location_edit = QLineEdit()
        self.status_combo = QComboBox()
        for status in ComputerStatus:
            self.status_combo.addItem(status.value, status.value)
        self.status_combo.setCurrentIndex(
            self.status_combo.findData(ComputerStatus.IN_STOCK.value)
        )

        form = QFormLayout()
        form.addRow("Department:", self.department_combo)
        form.addRow("Assigned User:", self.user_edit)
        form.addRow("Location:", self.location_edit)
        form.addRow("Status:", self.status_combo)

        buttons = QDialogButtonBox(
            QDialogButtonBox.Ok | QDialogButtonBox.Cancel
        )
        buttons.accepted.connect(self.accept)
        buttons.rejected.connect(self.reject)

        layout = QVBoxLayout(self)
        layout.addWidget(
            QLabel("Choose the inventory fields for this device.")
        )
        layout.addLayout(form)
        layout.addWidget(buttons)

    def get_data(self):
        return {
            "department_id": self.department_combo.currentData(),
            "assigned_user": self.user_edit.text().strip() or None,
            "location": self.location_edit.text().strip() or None,
            "status": self.status_combo.currentData(),
        }


class PhoneDialog(QDialog):

    def __init__(
        self,
        departments,
        phone=None,
        parent=None,
    ):

        super().__init__(parent)

        self.phone = phone

        self.setWindowTitle(
            "Edit Phone"
            if phone
            else "Add Phone"
        )

        self.setMinimumWidth(620)

        self.asset_tag_edit = QLineEdit()
        self.phone_number_edit = QLineEdit()
        self.brand_edit = QLineEdit()
        self.model_edit = QLineEdit()
        self.imei_1_edit = QLineEdit()
        self.imei_2_edit = QLineEdit()
        self.serial_number_edit = QLineEdit()

        self.os_edit = QLineEdit()
        self.os_version_edit = QLineEdit()
        self.mac_edit = QLineEdit()
        self.ip_edit = QLineEdit()

        self.department_combo = QComboBox()
        self.department_combo.addItem("No Department", None)
        for department in departments:
            self.department_combo.addItem(department["name"], department["id"])

        self.user_edit = QLineEdit()
        self.location_edit = QLineEdit()

        self.status_combo = QComboBox()
        for status in [
            PhoneStatus.IN_STOCK,
            PhoneStatus.ASSIGNED,
            PhoneStatus.RESERVED,
            PhoneStatus.MAINTENANCE,
            PhoneStatus.LOST,
            PhoneStatus.RETIRED,
        ]:
            self.status_combo.addItem(status.value, status.value)

        self.purchase_date_edit = QLineEdit()
        self.warranty_expiry_edit = QLineEdit()
        self.notes_edit = QTextEdit()
        self.notes_edit.setMinimumHeight(90)

        form = QFormLayout()

        form.addRow("Asset Tag:", self.asset_tag_edit)
        form.addRow("Phone Number:", self.phone_number_edit)
        form.addRow("Brand:", self.brand_edit)
        form.addRow("Model:", self.model_edit)
        form.addRow("IMEI 1:", self.imei_1_edit)
        form.addRow("IMEI 2:", self.imei_2_edit)
        form.addRow("Serial Number:", self.serial_number_edit)
        form.addRow("Operating System:", self.os_edit)
        form.addRow("OS Version:", self.os_version_edit)
        form.addRow("MAC Address:", self.mac_edit)
        form.addRow("IP Address:", self.ip_edit)
        form.addRow("Department:", self.department_combo)
        form.addRow("Assigned User:", self.user_edit)
        form.addRow("Location:", self.location_edit)
        form.addRow("Status:", self.status_combo)
        form.addRow("Purchase Date:", self.purchase_date_edit)
        form.addRow("Warranty Expiry:", self.warranty_expiry_edit)
        form.addRow("Notes:", self.notes_edit)

        if phone:
            self.load_phone(phone)

        buttons = QDialogButtonBox(
            QDialogButtonBox.Save
            | QDialogButtonBox.Cancel
        )

        buttons.accepted.connect(self.accept)
        buttons.rejected.connect(self.reject)

        layout = QVBoxLayout(self)
        layout.addLayout(form)
        layout.addWidget(buttons)

    def load_phone(self, phone):
        self.asset_tag_edit.setText(phone.get("asset_tag") or "")
        self.phone_number_edit.setText(phone.get("phone_number") or "")
        self.brand_edit.setText(phone.get("brand") or "")
        self.model_edit.setText(phone.get("model") or "")
        self.imei_1_edit.setText(phone.get("imei_1") or "")
        self.imei_2_edit.setText(phone.get("imei_2") or "")
        self.serial_number_edit.setText(phone.get("serial_number") or "")
        self.os_edit.setText(phone.get("operating_system") or "")
        self.os_version_edit.setText(phone.get("os_version") or "")
        self.mac_edit.setText(phone.get("mac_address") or "")
        self.ip_edit.setText(phone.get("ip_address") or "")

        department_id = phone.get("department_id")
        index = self.department_combo.findData(department_id)
        if index >= 0:
            self.department_combo.setCurrentIndex(index)

        self.user_edit.setText(phone.get("assigned_user") or "")
        self.location_edit.setText(phone.get("location") or "")

        status = phone.get("status")
        index = self.status_combo.findData(status)
        if index >= 0:
            self.status_combo.setCurrentIndex(index)

        self.purchase_date_edit.setText(phone.get("purchase_date") or "")
        self.warranty_expiry_edit.setText(phone.get("warranty_expiry") or "")
        self.notes_edit.setPlainText(phone.get("notes") or "")

    def get_data(self):
        return {
            "asset_tag": self.asset_tag_edit.text().strip() or None,
            "phone_number": self.phone_number_edit.text().strip() or None,
            "brand": self.brand_edit.text().strip() or None,
            "model": self.model_edit.text().strip() or None,
            "imei_1": self.imei_1_edit.text().strip() or None,
            "imei_2": self.imei_2_edit.text().strip() or None,
            "serial_number": self.serial_number_edit.text().strip() or None,
            "operating_system": self.os_edit.text().strip() or None,
            "os_version": self.os_version_edit.text().strip() or None,
            "mac_address": self.mac_edit.text().strip() or None,
            "ip_address": self.ip_edit.text().strip() or None,
            "department_id": self.department_combo.currentData(),
            "assigned_user": self.user_edit.text().strip() or None,
            "location": self.location_edit.text().strip() or None,
            "status": self.status_combo.currentData(),
            "purchase_date": self.purchase_date_edit.text().strip() or None,
            "warranty_expiry": self.warranty_expiry_edit.text().strip() or None,
            "notes": self.notes_edit.toPlainText().strip() or None,
        }


class TelephoneDialog(QDialog):

    def __init__(self, departments, telephone=None, parent=None):
        super().__init__(parent)
        self.telephone = telephone
        self.setWindowTitle("Edit Telephone" if telephone else "Add Telephone")
        self.setMinimumWidth(560)

        self.asset_tag_edit = QLineEdit()
        self.number_edit = QLineEdit()
        self.brand_edit = QLineEdit()
        self.model_edit = QLineEdit()
        self.serial_edit = QLineEdit()
        self.mac_edit = QLineEdit()
        self.ip_edit = QLineEdit()
        self.user_edit = QLineEdit()
        self.location_edit = QLineEdit()
        self.department_combo = QComboBox()
        self.department_combo.addItem("No Department", None)
        for department in departments:
            self.department_combo.addItem(department["name"], department["id"])
        self.status_combo = QComboBox()
        for status in TelephoneStatus:
            self.status_combo.addItem(status.value, status.value)
        self.purchase_date_edit = QLineEdit()
        self.notes_edit = QTextEdit()
        self.notes_edit.setMinimumHeight(90)

        form = QFormLayout()
        for label, field in (
            ("Asset Tag:", self.asset_tag_edit),
            ("Telephone Number / Extension:", self.number_edit),
            ("Brand:", self.brand_edit),
            ("Model:", self.model_edit),
            ("Serial Number:", self.serial_edit),
            ("MAC Address:", self.mac_edit),
            ("IP Address:", self.ip_edit),
        ):
            form.addRow(label, field)
        form.addRow("Department:", self.department_combo)
        form.addRow("Assigned User:", self.user_edit)
        form.addRow("Location:", self.location_edit)
        form.addRow("Status:", self.status_combo)
        form.addRow("Purchase Date:", self.purchase_date_edit)
        form.addRow("Notes:", self.notes_edit)

        if telephone:
            self.load_telephone(telephone)

        buttons = QDialogButtonBox(QDialogButtonBox.Save | QDialogButtonBox.Cancel)
        buttons.accepted.connect(self.accept)
        buttons.rejected.connect(self.reject)
        layout = QVBoxLayout(self)
        layout.addLayout(form)
        layout.addWidget(buttons)

    def load_telephone(self, telephone):
        self.asset_tag_edit.setText(telephone.get("asset_tag") or "")
        self.number_edit.setText(telephone.get("telephone_number") or "")
        self.brand_edit.setText(telephone.get("brand") or "")
        self.model_edit.setText(telephone.get("model") or "")
        self.serial_edit.setText(telephone.get("serial_number") or "")
        self.mac_edit.setText(telephone.get("mac_address") or "")
        self.ip_edit.setText(telephone.get("ip_address") or "")
        index = self.department_combo.findData(telephone.get("department_id"))
        if index >= 0:
            self.department_combo.setCurrentIndex(index)
        self.user_edit.setText(telephone.get("assigned_user") or "")
        self.location_edit.setText(telephone.get("location") or "")
        index = self.status_combo.findData(telephone.get("status"))
        if index >= 0:
            self.status_combo.setCurrentIndex(index)
        self.purchase_date_edit.setText(telephone.get("purchase_date") or "")
        self.notes_edit.setPlainText(telephone.get("notes") or "")

    def get_data(self):
        return {
            "asset_tag": self.asset_tag_edit.text().strip() or None,
            "telephone_number": self.number_edit.text().strip() or None,
            "brand": self.brand_edit.text().strip() or None,
            "model": self.model_edit.text().strip() or None,
            "serial_number": self.serial_edit.text().strip() or None,
            "mac_address": self.mac_edit.text().strip() or None,
            "ip_address": self.ip_edit.text().strip() or None,
            "department_id": self.department_combo.currentData(),
            "assigned_user": self.user_edit.text().strip() or None,
            "location": self.location_edit.text().strip() or None,
            "status": self.status_combo.currentData(),
            "purchase_date": self.purchase_date_edit.text().strip() or None,
            "notes": self.notes_edit.toPlainText().strip() or None,
        }


class GenericEquipmentDialog(QDialog):

    def __init__(self, equipment_name, departments, asset=None, parent=None):
        super().__init__(parent)
        self.setWindowTitle(f"Edit {equipment_name}" if asset else f"Add {equipment_name}")
        self.setMinimumWidth(600)
        self.asset_tag_edit = QLineEdit()
        self.name_edit = QLineEdit()
        self.brand_edit = QLineEdit()
        self.model_edit = QLineEdit()
        self.serial_edit = QLineEdit()
        self.ip_edit = QLineEdit()
        self.mac_edit = QLineEdit()
        self.user_edit = QLineEdit()
        self.location_edit = QLineEdit()
        self.department_combo = QComboBox()
        self.department_combo.addItem("No Department", None)
        for department in departments:
            self.department_combo.addItem(department["name"], department["id"])
        self.status_combo = QComboBox()
        for status in EQUIPMENT_ASSET_STATUSES:
            self.status_combo.addItem(status, status)
        self.notes_edit = QTextEdit()
        self.notes_edit.setMinimumHeight(80)
        form = QFormLayout()
        for label, field in (
            ("Asset Tag:", self.asset_tag_edit),
            ("Name / Identifier:", self.name_edit),
            ("Brand:", self.brand_edit),
            ("Model:", self.model_edit),
            ("Serial Number:", self.serial_edit),
            ("IP Address:", self.ip_edit),
            ("MAC Address:", self.mac_edit),
        ):
            form.addRow(label, field)
        form.addRow("Department:", self.department_combo)
        form.addRow("Assigned User:", self.user_edit)
        form.addRow("Location:", self.location_edit)
        form.addRow("Status:", self.status_combo)
        form.addRow("Notes:", self.notes_edit)
        if asset:
            self.load_asset(asset)
        buttons = QDialogButtonBox(QDialogButtonBox.Save | QDialogButtonBox.Cancel)
        buttons.accepted.connect(self.accept)
        buttons.rejected.connect(self.reject)
        layout = QVBoxLayout(self)
        layout.addLayout(form)
        layout.addWidget(buttons)

    def load_asset(self, asset):
        self.asset_tag_edit.setText(asset.get("asset_tag") or "")
        self.name_edit.setText(asset.get("name") or "")
        self.brand_edit.setText(asset.get("brand") or "")
        self.model_edit.setText(asset.get("model") or "")
        self.serial_edit.setText(asset.get("serial_number") or "")
        self.ip_edit.setText(asset.get("ip_address") or "")
        self.mac_edit.setText(asset.get("mac_address") or "")
        index = self.department_combo.findData(asset.get("department_id"))
        if index >= 0:
            self.department_combo.setCurrentIndex(index)
        self.user_edit.setText(asset.get("assigned_user") or "")
        self.location_edit.setText(asset.get("location") or "")
        index = self.status_combo.findData(asset.get("status"))
        if index >= 0:
            self.status_combo.setCurrentIndex(index)
        self.notes_edit.setPlainText(asset.get("notes") or "")

    def get_data(self):
        return {
            "asset_tag": self.asset_tag_edit.text().strip() or None,
            "name": self.name_edit.text().strip() or None,
            "brand": self.brand_edit.text().strip() or None,
            "model": self.model_edit.text().strip() or None,
            "serial_number": self.serial_edit.text().strip() or None,
            "ip_address": self.ip_edit.text().strip() or None,
            "mac_address": self.mac_edit.text().strip() or None,
            "department_id": self.department_combo.currentData(),
            "assigned_user": self.user_edit.text().strip() or None,
            "location": self.location_edit.text().strip() or None,
            "status": self.status_combo.currentData(),
            "notes": self.notes_edit.toPlainText().strip() or None,
        }


# =========================================================
# MAIN WINDOW
# =========================================================

class MainWindow(QMainWindow):

    def __init__(
        self,
        admin,
        parent=None,
    ):

        super().__init__(parent)

        self.admin = admin
        self._computer_department_id = None
        self._computer_department_view = False
        self._phone_department_id = None
        self._phone_department_view = False

        self.setWindowTitle(
            "MANAGEMENT CONSOLE"
        )

        self.resize(
            1400,
            820,
        )

        # -------------------------------------------------
        # TREE
        # -------------------------------------------------

        self.tree = QTreeWidget()

        self.tree.setHeaderLabel(
            "Management Console"
        )

        self.tree.setMinimumWidth(
            280
        )

        self.tree.itemClicked.connect(
            self.show_item
        )

        # -------------------------------------------------
        # WORKSPACE
        # -------------------------------------------------

        self.workspace = QStackedWidget()

        self.empty_workspace = QTextEdit()

        self.empty_workspace.setReadOnly(
            True
        )

        self.empty_workspace.setPlainText(
            "Select an item from the navigation tree."
        )

        self.workspace.addWidget(
            self.empty_workspace
        )

        # -------------------------------------------------
        # SPLITTER
        # -------------------------------------------------

        splitter = QSplitter(
            Qt.Horizontal
        )

        splitter.addWidget(
            self.tree
        )

        splitter.addWidget(
            self.workspace
        )

        splitter.setStretchFactor(
            1,
            1,
        )

        self.setCentralWidget(
            splitter
        )

        # -------------------------------------------------
        # STATUS BAR
        # -------------------------------------------------

        self.statusBar().showMessage(
            f"Signed in as {admin.username}"
        )

        # -------------------------------------------------
        # LOAD TREE
        # -------------------------------------------------

        self.refresh_tree()

    # =====================================================
    # TREE
    # =====================================================

    def refresh_tree(self):

        self.tree.clear()

        root = self.create_tree_item(
            "MANAGEMENT CONSOLE",
            "root",
        )

        self.tree.addTopLevelItem(
            root
        )

        # -------------------------------------------------
        # DASHBOARD
        # -------------------------------------------------

        dashboard = self.create_tree_item(
            "Dashboard",
            "dashboard",
        )

        root.addChild(
            dashboard
        )

        # -------------------------------------------------
        # DEPARTMENTS
        # -------------------------------------------------

        departments_node = self.create_tree_item(
            "Departments",
            "departments",
        )

        root.addChild(
            departments_node
        )

        # -------------------------------------------------
        # GLOBAL ASSETS
        # -------------------------------------------------

        enabled_keys = set()
        session = SessionLocal()
        try:
            enabled_keys = {
                item.key for item in get_enabled_equipment_types(session)
            }
        finally:
            session.close()

        session = SessionLocal()
        try:
            equipment_types = get_equipment_types(session)
        finally:
            session.close()
        equipment_node = self.create_tree_item("Equipment", "equipment")
        stock_node = self.create_tree_item("Stock", "stock")
        root.addChild(equipment_node)
        root.addChild(stock_node)
        for equipment_type in equipment_types:
            if equipment_type.key not in enabled_keys:
                continue
            kind = get_equipment_page_kind(equipment_type.key)
            equipment_node.addChild(
                self.create_tree_item(
                    equipment_type.name,
                    kind,
                    equipment_type.key if kind == "equipment_type" else None,
                )
            )
            stock_node.addChild(
                self.create_tree_item(
                    equipment_type.name,
                    "stock_equipment",
                    equipment_type.key,
                )
            )

        # -------------------------------------------------
        # NETWORK DISCOVERY
        # -------------------------------------------------

        network_node = self.create_tree_item(
            "Network Discovery",
            "network_discovery",
        )

        root.addChild(
            network_node
        )







        # -------------------------------------------------
        # MAINTENANCE
        # -------------------------------------------------

        maintenance_node = self.create_tree_item(
            "Maintenance",
            "maintenance",
        )

        root.addChild(
            maintenance_node
        )

        # -------------------------------------------------
        # REPORTS
        # -------------------------------------------------

        reports_node = self.create_tree_item(
            "Reports",
            "reports",
        )

        root.addChild(
            reports_node
        )

        reports_overview = self.create_tree_item(
            "Overview",
            "reports_overview",
        )

        reports_departments = self.create_tree_item(
            "Departments",
            "reports_departments",
        )

        equipment_reports = self.create_tree_item(
            "Equipment Reports",
            "equipment_reports",
        )

        stock_reports = self.create_tree_item(
            "Stock Reports",
            "stock_reports",
        )

        for equipment_type in equipment_types:
            if equipment_type.key not in enabled_keys:
                continue
            equipment_reports.addChild(
                self.create_tree_item(
                    equipment_type.name,
                    "equipment_report",
                    equipment_type.key,
                )
            )
            stock_reports.addChild(
                self.create_tree_item(
                    equipment_type.name,
                    "stock_report_equipment",
                    equipment_type.key,
                )
            )

        reports_maintenance = self.create_tree_item(
            "Maintenance",
            "reports_maintenance",
        )

        reports_node.addChild(
            reports_overview
        )

        reports_node.addChild(
            reports_departments
        )

        reports_node.addChild(
            equipment_reports
        )

        reports_node.addChild(
            stock_reports
        )

        reports_node.addChild(
            reports_maintenance
        )

        # -------------------------------------------------
        # ADMIN
        # -------------------------------------------------

        admin_node = self.create_tree_item(
            "Admin Management",
            "admins",
        )

        administrators_item = self.create_tree_item(
            "Administrators",
            "admin_list",
        )

        add_admin_item = self.create_tree_item(
            "Add Administrator",
            "admin_add",
        )

        admin_node.addChild(
            administrators_item
        )

        admin_node.addChild(
            add_admin_item
        )

        root.addChild(
            admin_node
        )

        # -------------------------------------------------
        # SETTINGS
        # -------------------------------------------------

        settings_node = self.create_tree_item(
            "Settings",
            "settings",
        )

        root.addChild(
            settings_node
        )

        # =================================================
        # DATABASE
        # =================================================

        session = SessionLocal()

        try:

            departments = (
                session.query(
                    Department
                )
                .order_by(
                    Department.name
                )
                .all()
            )
            generic_assets = session.query(EquipmentAsset).all()

            # ---------------------------------------------
            # DEPARTMENT TREE
            # ---------------------------------------------

            for department in departments:

                department_item = (
                    self.create_tree_item(
                        f"{department.name} "
                        f"({len(department.computers)} PCs)",
                        "department",
                        department.id,
                    )
                )

                departments_node.addChild(
                    department_item
                )

                department_types = [
                    item for item in equipment_types
                    if item.key in enabled_keys
                    and (
                        (
                            item.key == "computers" and bool(department.computers)
                        )
                        or (
                            item.key == "phones" and bool(department.phones)
                        )
                        or (
                            item.key == "telephones" and bool(department.telephones)
                        )
                        or any(
                            asset.equipment_type_id == item.id
                            and asset.department_id == department.id
                            for asset in generic_assets
                        )
                    )
                ]
                for equipment_type in department_types:
                    page_kind = get_equipment_page_kind(equipment_type.key)
                    if page_kind == "equipment_type":
                        item_kind = "department_equipment"
                        item_id = (department.id, equipment_type.key)
                    else:
                        item_kind = f"department_{page_kind}"
                        item_id = department.id
                    department_item.addChild(
                        self.create_tree_item(
                            equipment_type.name,
                            item_kind,
                            item_id,
                        )
                    )

            departments_node.setText(
                0,
                f"Departments ({len(departments)})",
            )

        finally:

            session.close()

        root.setExpanded(
            True
        )

        departments_node.setExpanded(
            True
        )

        self.tree.setCurrentItem(
            dashboard
        )

        self.show_item(
            dashboard,
            0,
        )

    # =====================================================
    # CREATE TREE ITEM
    # =====================================================

    def create_tree_item(
        self,
        text,
        kind,
        object_id=None,
    ):

        item = QTreeWidgetItem(
            [text]
        )

        item.setData(
            0,
            Qt.UserRole,
            (
                kind,
                object_id,
            ),
        )

        return item

    # =====================================================
    # TREE SELECTION
    # =====================================================

    def show_item(
        self,
        item,
        _column,
    ):

        data = item.data(
            0,
            Qt.UserRole,
        )

        if not data:
            return

        kind, object_id = data

        if kind == "dashboard":

            self.show_dashboard()

        elif kind == "computer":

            self.show_computer(
                object_id
            )

        elif kind == "department":

            self.show_department(
                object_id
            )

        elif kind == "departments":

            self.show_departments()

        elif kind == "computers":

            self.show_computers()

        elif kind == "equipment":
            return

        elif kind == "stock_equipment":

            self.show_stock_equipment(object_id)

        elif kind == "department_phones":

            self.show_department_phones(
                object_id
            )

        elif kind == "department_computers":

            self.show_department_computers(
                object_id
            )

        elif kind == "phones":

            self.show_phones()

        elif kind == "telephones":

            self.show_telephones()

        elif kind == "department_telephones":

            self.show_telephones(object_id, True)

        elif kind == "department_equipment":

            department_id, equipment_key = object_id
            self.show_equipment_management(equipment_key, department_id)

        elif kind == "equipment_type":

            self.show_equipment_management(object_id)

        elif kind == "network_discovery":

            self.show_network_discovery()

        elif kind == "stock":

            self.show_stock()

        elif kind == "maintenance":

            self.show_text(
                "maintenance"
            )

        elif kind == "reports":

            self.show_reports()

        elif kind == "equipment_reports":

            return

        elif kind == "equipment_report":

            self.show_equipment_report(object_id)

        elif kind == "stock_reports":

            self.show_dynamic_stock_report()

        elif kind == "stock_report_equipment":

            self.show_equipment_report(object_id, stock_only=True)

        elif kind == "reports_overview":

            self.show_reports()

        elif kind == "reports_departments":

            self.show_dynamic_departments_report()

        elif kind == "reports_computers":

            self.show_computers_report()

        elif kind == "reports_stock":

            self.show_stock_report()

        elif kind == "reports_maintenance":

            self.show_maintenance_report()

        elif kind == "admins":

            self.show_admins()

        elif kind == "admin_list":

            self.show_admins()

        elif kind == "admin_add":

            self.show_add_admin_form()

        elif kind == "settings":

            self.show_settings()

    # =====================================================
    # DASHBOARD
    # =====================================================

    def show_dashboard(self):

        widget = QWidget()
        layout = QVBoxLayout(widget)

        title = QLabel(
            "Dashboard"
        )

        title.setStyleSheet(
            "font-size: 28px; "
            "font-weight: 700;"
        )

        subtitle = QLabel(
            "Management Console overview"
        )

        session = SessionLocal()

        try:
            departments = session.query(
                Department
            ).count()

        finally:

            session.close()

        # -------------------------------------------------
        # STATISTICS
        # -------------------------------------------------

        stats = QGroupBox(
            "System Overview"
        )

        stats_layout = QFormLayout(
            stats
        )

        stats_layout.addRow(
            "Departments:",
            QLabel(str(departments)),
        )

        managed_group = QGroupBox("Managed Equipment")
        managed_layout = QFormLayout(managed_group)
        session = SessionLocal()
        try:
            managed_summary = get_enabled_equipment_summary(session)
        finally:
            session.close()
        for summary in managed_summary:
            statuses = summary["by_status"]
            managed_layout.addRow(
                f"{summary['name']}:",
                QLabel(
                    f"Total {summary['total']} | "
                    f"In Stock {statuses.get('IN_STOCK', 0)} | "
                    f"Assigned {statuses.get('ASSIGNED', 0)} | "
                    f"Maintenance {statuses.get('MAINTENANCE', 0)}"
                ),
            )

        session = SessionLocal()
        try:
            department_summary = get_enabled_department_summary(session)
            enabled_report_types = get_enabled_report_types(session)
        finally:
            session.close()
        distribution_group = QGroupBox("Department Equipment Distribution")
        distribution_layout = QVBoxLayout(distribution_group)
        distribution_columns = ["Department"] + [item.name for item in enabled_report_types]
        distribution_table = QTableWidget()
        distribution_table.setColumnCount(len(distribution_columns))
        distribution_table.setHorizontalHeaderLabels(distribution_columns)
        distribution_table.setRowCount(len(department_summary))
        for row_index, row in enumerate(department_summary):
            for column_index, column in enumerate(distribution_columns):
                value = row["Department"] if column == "Department" else row.get(column, 0)
                distribution_table.setItem(row_index, column_index, QTableWidgetItem(str(value)))
        distribution_table.resizeColumnsToContents()
        distribution_table.horizontalHeader().setStretchLastSection(True)
        distribution_layout.addWidget(distribution_table)

        # -------------------------------------------------
        # QUICK ACTIONS
        # -------------------------------------------------

        quick_group = QGroupBox(
            "Quick Actions"
        )

        quick_layout = QHBoxLayout(
            quick_group
        )

        computers_button = QPushButton(
            "Manage Computers"
        )

        phones_button = QPushButton("Manage Phones")

        departments_button = QPushButton(
            "Manage Departments"
        )

        stock_button = QPushButton(
            "View Stock"
        )

        computers_button.clicked.connect(
            self.show_computers
        )

        phones_button.clicked.connect(
            self.show_phones
        )

        session = SessionLocal()
        try:
            enabled_equipment = get_enabled_equipment_types(session)
        finally:
            session.close()

        if not any(item.key == "phones" for item in enabled_equipment):
            phones_button.hide()

        for equipment_type in enabled_equipment:
            if equipment_type.key == "telephones":
                telephone_button = QPushButton("Manage Telephones")
                telephone_button.clicked.connect(self.show_telephones)
                quick_layout.addWidget(telephone_button)
            elif equipment_type.key not in {"computers", "phones"}:
                equipment_button = QPushButton(f"Manage {equipment_type.name}")
                equipment_button.clicked.connect(
                    lambda _checked=False, key=equipment_type.key: self.show_equipment_management(key)
                )
                quick_layout.addWidget(equipment_button)

        departments_button.clicked.connect(
            self.show_departments
        )

        stock_button.clicked.connect(
            self.show_stock
        )

        quick_layout.addWidget(
            computers_button
        )

        quick_layout.addWidget(
            phones_button
        )

        quick_layout.addWidget(
            departments_button
        )

        quick_layout.addWidget(
            stock_button
        )

        quick_layout.addStretch()

        layout.addWidget(
            title
        )

        layout.addWidget(
            subtitle
        )

        layout.addWidget(
            stats
        )

        layout.addWidget(managed_group)
        layout.addWidget(distribution_group)

        layout.addWidget(
            quick_group
        )

        layout.addStretch()

        self.workspace.addWidget(
            widget
        )

        self.workspace.setCurrentWidget(
            widget
        )

    # =====================================================
    # DEPARTMENTS
    # =====================================================

    def show_departments(self):

        widget = QWidget()

        layout = QVBoxLayout(
            widget
        )

        title = QLabel(
            "Departments"
        )

        title.setStyleSheet(
            "font-size: 24px; "
            "font-weight: 600;"
        )

        subtitle = QLabel(
            "Manage hotel departments and their computers."
        )

        buttons_layout = QHBoxLayout()

        add_button = QPushButton(
            "+ Add Department"
        )

        edit_button = QPushButton(
            "Edit"
        )

        delete_button = QPushButton(
            "Delete"
        )

        refresh_button = QPushButton(
            "Refresh"
        )

        buttons_layout.addWidget(
            add_button
        )

        buttons_layout.addWidget(
            edit_button
        )

        buttons_layout.addWidget(
            delete_button
        )

        buttons_layout.addWidget(
            refresh_button
        )

        buttons_layout.addStretch()

        add_button.clicked.connect(
            self.add_department
        )

        edit_button.clicked.connect(
            self.edit_selected_department
        )

        delete_button.clicked.connect(
            self.delete_selected_department
        )

        refresh_button.clicked.connect(
            self.refresh_departments_view
        )

        self.department_tree = QTreeWidget()

        self.department_tree.setHeaderLabels(
            [
                "Department",
                "Computers",
            ]
        )

        self.load_department_table()

        layout.addWidget(
            title
        )

        layout.addWidget(
            subtitle
        )

        layout.addLayout(
            buttons_layout
        )

        layout.addWidget(
            self.department_tree
        )

        self.workspace.addWidget(
            widget
        )

        self.workspace.setCurrentWidget(
            widget
        )

    # =====================================================
    # LOAD DEPARTMENTS
    # =====================================================

    def load_department_table(self):

        if not hasattr(
            self,
            "department_tree",
        ):
            return

        self.department_tree.clear()

        session = SessionLocal()

        try:

            departments = get_departments(
                session
            )

            for department in departments:

                item = QTreeWidgetItem(
                    [
                        department["name"],
                        str(
                            department[
                                "computer_count"
                            ]
                        ),
                    ]
                )

                item.setData(
                    0,
                    Qt.UserRole,
                    department["id"],
                )

                self.department_tree.addTopLevelItem(
                    item
                )

        finally:

            session.close()

    # =====================================================
    # ADD DEPARTMENT
    # =====================================================

    def add_department(self):

        dialog = DepartmentDialog(
            parent=self
        )

        if (
            dialog.exec()
            != QDialog.Accepted
        ):
            return

        data = dialog.get_data()

        if not data["name"]:

            QMessageBox.warning(
                self,
                "Invalid Department",
                "Department name is required.",
            )

            return

        session = SessionLocal()

        try:

            create_department(
                session,
                data,
            )

            session.commit()

        except Exception as exc:

            session.rollback()

            QMessageBox.critical(
                self,
                "Error",
                f"Could not create department:\n\n{exc}",
            )

            return

        finally:

            session.close()

        self.refresh_tree()

        self.show_departments()

    # =====================================================
    # EDIT DEPARTMENT
    # =====================================================

    def edit_selected_department(self):

        item = (
            self.department_tree.currentItem()
        )

        if item is None:

            QMessageBox.information(
                self,
                "Select Department",
                "Please select a department first.",
            )

            return

        department_id = item.data(
            0,
            Qt.UserRole,
        )

        session = SessionLocal()

        try:

            department = session.get(
                Department,
                department_id,
            )

            if department is None:

                QMessageBox.warning(
                    self,
                    "Not Found",
                    "Department no longer exists.",
                )

                return

            dialog = DepartmentDialog(
                department,
                self,
            )

            if (
                dialog.exec()
                != QDialog.Accepted
            ):
                return

            data = dialog.get_data()

            if not data["name"]:

                QMessageBox.warning(
                    self,
                    "Invalid Department",
                    "Department name is required.",
                )

                return

            update_department(
                session,
                department_id,
                data,
            )

            session.commit()

        except Exception as exc:

            session.rollback()

            QMessageBox.critical(
                self,
                "Error",
                f"Could not update department:\n\n{exc}",
            )

            return

        finally:

            session.close()

        self.refresh_tree()

        self.show_departments()

    # =====================================================
    # DELETE DEPARTMENT
    # =====================================================

    def delete_selected_department(self):

        item = (
            self.department_tree.currentItem()
        )

        if item is None:

            QMessageBox.information(
                self,
                "Select Department",
                "Please select a department first.",
            )

            return

        department_id = item.data(
            0,
            Qt.UserRole,
        )

        department_name = item.text(
            0
        )

        confirmation = QMessageBox.question(
            self,
            "Delete Department",
            (
                "Are you sure you want to delete "
                f"'{department_name}'?"
            ),
            QMessageBox.Yes
            | QMessageBox.No,
        )

        if (
            confirmation
            != QMessageBox.Yes
        ):
            return

        session = SessionLocal()

        try:

            delete_department(
                session,
                department_id,
            )

            session.commit()

        except Exception as exc:

            session.rollback()

            QMessageBox.warning(
                self,
                "Cannot Delete",
                str(exc),
            )

            return

        finally:

            session.close()

        self.refresh_tree()

        self.show_departments()

    # =====================================================
    # REFRESH DEPARTMENTS
    # =====================================================

    def refresh_departments_view(self):

        self.refresh_tree()

        self.show_departments()

    # =====================================================
    # DEPARTMENT DETAILS
    # =====================================================

    def show_department(
        self,
        department_id,
    ):

        session = SessionLocal()

        try:

            department = get_department(
                session,
                department_id,
            )

        except Exception as exc:

            QMessageBox.warning(
                self,
                "Error",
                str(exc),
            )

            return

        finally:

            session.close()

        widget = QWidget()
        outer_layout = QVBoxLayout(widget)
        content_widget = QWidget()
        layout = QVBoxLayout(content_widget)
        department_scroll = QScrollArea()
        department_scroll.setWidgetResizable(True)
        department_scroll.setWidget(content_widget)
        outer_layout.addWidget(department_scroll)

        title = QLabel(
            department["name"]
        )

        title.setStyleSheet(
            "font-size: 24px; "
            "font-weight: 600;"
        )

        description = QLabel(
            department["description"]
            or "No description."
        )

        group = QGroupBox(
            "Department Information"
        )

        group_layout = QFormLayout(
            group
        )

        group_layout.addRow(
            "Name:",
            QLabel(
                department["name"]
            ),
        )

        group_layout.addRow(
            "Computers:",
            QLabel(
                str(
                    department[
                        "computer_count"
                    ]
                )
            ),
        )

        group_layout.addRow(
            "Phones:",
            QLabel(
                str(
                    department[
                        "phone_count"
                    ]
                )
            ),
        )

        session = SessionLocal()
        try:
            enabled_equipment_types = get_enabled_equipment_types(session)
            enabled_keys = {item.key for item in enabled_equipment_types}
        finally:
            session.close()
        telephone_enabled = "telephones" in enabled_keys
        if telephone_enabled:
            group_layout.addRow(
                "Telephones:",
                QLabel(str(department["telephone_count"])),
            )

        # Department equipment is rendered from the central registry. The
        # legacy per-model widgets below remain unused for compatibility.
        layout.addWidget(title)
        layout.addWidget(description)
        layout.addWidget(group)
        filters_layout = QHBoxLayout()
        department_search = QLineEdit()
        department_search.setPlaceholderText("Search equipment in this department...")
        department_status = QComboBox()
        department_status.addItem("All Statuses", None)
        for status in EQUIPMENT_ASSET_STATUSES:
            department_status.addItem(status.replace("_", " ").title(), status)
        filters_layout.addWidget(QLabel("Search:"))
        filters_layout.addWidget(department_search, 1)
        filters_layout.addWidget(QLabel("Status:"))
        filters_layout.addWidget(department_status)
        layout.addLayout(filters_layout)

        section_widgets = []
        for section in department["equipment_sections"]:
            section_group = QGroupBox(section["name"])
            section_layout = QVBoxLayout(section_group)
            section_table = QTreeWidget()
            section_table.setRootIsDecorated(False)
            section_table.setAlternatingRowColors(True)
            section_table.setHeaderLabels(
                ["Asset Tag", "Name", "Brand", "Model", "Serial", "IP", "Status", "User", "Location"]
            )
            section_layout.addWidget(section_table)
            layout.addWidget(section_group)
            section_widgets.append((section, section_group, section_table))

        def refresh_department_sections():
            search = department_search.text().strip().lower()
            status = department_status.currentData()
            for section, section_group, section_table in section_widgets:
                section_table.clear()
                visible = []
                for record in section["records"]:
                    values = [
                        record.get("asset_tag"), record.get("name"), record.get("brand"),
                        record.get("model"), record.get("serial_number"), record.get("ip_address"),
                        record.get("status"), record.get("assigned_user"), record.get("location"),
                    ]
                    if status and record.get("status") != status:
                        continue
                    if search and search not in " ".join(str(value or "").lower() for value in values):
                        continue
                    visible.append(values)
                for values in visible:
                    item = QTreeWidgetItem([str(value or "-") for value in values])
                    section_table.addTopLevelItem(item)
                section_group.setVisible(bool(visible))

        department_search.textChanged.connect(refresh_department_sections)
        department_status.currentIndexChanged.connect(refresh_department_sections)
        refresh_department_sections()
        self.workspace.addWidget(widget)
        self.workspace.setCurrentWidget(widget)
        return

        computers_group = QGroupBox(
            "Computers"
        )

        computers_layout = QVBoxLayout(
            computers_group
        )

        computer_tree = QTreeWidget()

        computer_tree.setHeaderLabels(
            [
                "Computer",
                "Status",
                "User",
                "Location",
            ]
        )

        for computer in department[
            "computers"
        ]:

            item = QTreeWidgetItem(
                [
                    computer[
                        "computer_name"
                    ],
                    computer[
                        "status"
                    ]
                    or "-",
                    computer[
                        "assigned_user"
                    ]
                    or "-",
                    computer[
                        "location"
                    ]
                    or "-",
                ]
            )

            item.setData(
                0,
                Qt.UserRole,
                computer["id"],
            )

            computer_tree.addTopLevelItem(
                item
            )

        computer_tree.itemClicked.connect(
            lambda item, _column:
                self.show_computer(
                    item.data(
                        0,
                        Qt.UserRole,
                    )
                )
        )

        computers_layout.addWidget(
            computer_tree
        )

        phones_group = QGroupBox(
            "Phones"
        )

        phones_layout = QVBoxLayout(
            phones_group
        )

        phone_tree = QTreeWidget()

        phone_tree.setHeaderLabels(
            [
                "Asset Tag",
                "Phone Number",
                "Status",
                "User",
                "Location",
                "Brand",
                "Model",
            ]
        )

        for phone in department[
            "phones"
        ]:

            item = QTreeWidgetItem(
                [
                    phone[
                        "asset_tag"
                    ]
                    or "-",
                    phone[
                        "phone_number"
                    ]
                    or "-",
                    phone[
                        "status"
                    ]
                    or "-",
                    phone[
                        "assigned_user"
                    ]
                    or "-",
                    phone[
                        "location"
                    ]
                    or "-",
                    phone[
                        "brand"
                    ]
                    or "-",
                    phone[
                        "model"
                    ]
                    or "-",
                ]
            )

            item.setData(
                0,
                Qt.UserRole,
                phone["id"],
            )

            phone_tree.addTopLevelItem(
                item
            )

        phones_layout.addWidget(
            phone_tree
        )

        telephone_group = QGroupBox("Telephones")
        telephone_layout = QVBoxLayout(telephone_group)
        telephone_tree = QTreeWidget()
        telephone_tree.setHeaderLabels(["Asset Tag", "Number / Extension", "Status", "User", "Location", "Brand", "Model"])
        for telephone in department["telephones"]:
            item = QTreeWidgetItem([
                telephone["asset_tag"] or "-",
                telephone["telephone_number"] or "-",
                telephone["status"] or "-",
                telephone["assigned_user"] or "-",
                telephone["location"] or "-",
                telephone["brand"] or "-",
                telephone["model"] or "-",
            ])
            item.setData(0, Qt.UserRole, telephone["id"])
            telephone_tree.addTopLevelItem(item)
        telephone_tree.itemClicked.connect(
            lambda _item, _column: self.show_telephones(department_id, True)
        )
        telephone_layout.addWidget(telephone_tree)

        layout.addWidget(
            title
        )

        layout.addWidget(
            description
        )

        layout.addWidget(
            group
        )

        if "computers" in enabled_keys and department["computers"]:
            layout.addWidget(computers_group)

        if "phones" in enabled_keys and department["phones"]:
            layout.addWidget(phones_group)

        if telephone_enabled and department["telephones"]:
            layout.addWidget(telephone_group)

        for equipment_type in enabled_equipment_types:
            if equipment_type.key in {"computers", "phones", "telephones"}:
                continue
            assets = [
                asset for asset in department["equipment_assets"]
                if asset["equipment_type_key"] == equipment_type.key
            ]
            if not assets:
                continue
            asset_group = QGroupBox(equipment_type.name)
            asset_layout = QVBoxLayout(asset_group)
            asset_tree = QTreeWidget()
            asset_tree.setHeaderLabels(["Asset Tag", "Name", "Brand", "Model", "Serial", "IP", "Status", "User", "Location"])
            for asset in assets:
                item = QTreeWidgetItem([
                    asset["asset_tag"] or "-", asset["name"] or "-", asset["brand"] or "-",
                    asset["model"] or "-", asset["serial_number"] or "-", asset["ip_address"] or "-",
                    asset["status"] or "-", asset["assigned_user"] or "-", asset["location"] or "-",
                ])
                item.setData(0, Qt.UserRole, asset["id"])
                asset_tree.addTopLevelItem(item)
            asset_tree.itemClicked.connect(
                lambda _item, _column, key=equipment_type.key: self.show_equipment_management(key, department_id)
            )
            asset_layout.addWidget(asset_tree)
            layout.addWidget(asset_group)

        self.workspace.addWidget(
            widget
        )

        self.workspace.setCurrentWidget(
            widget
        )

    # =====================================================
    # COMPUTERS
    # =====================================================

    def show_department_computers(self, department_id):

        self.show_computers(
            department_id,
            department_view=True,
        )

    def show_computers(
        self,
        department_id=None,
        department_view=False,
    ):

        self._computer_department_id = department_id
        self._computer_department_view = department_view

        session = SessionLocal()

        try:

            departments = get_departments(
                session
            )

        finally:
            session.close()

        department_name = next(
            (
                department["name"]
                for department in departments
                if department["id"] == department_id
            ),
            None,
        )

        widget = QWidget()

        layout = QVBoxLayout(
            widget
        )

        title = QLabel(
            (
                f"PCs — {department_name}"
                if department_name
                else "Computers"
            )
        )

        title.setStyleSheet(
            "font-size: 24px; "
            "font-weight: 600;"
        )

        subtitle = QLabel(
            "Manage hotel computers, assignments, stock and status."
        )

        # -------------------------------------------------
        # BUTTONS
        # -------------------------------------------------

        buttons_layout = QHBoxLayout()

        add_button = QPushButton(
            "+ Add Computer"
        )

        edit_button = QPushButton(
            "Edit"
        )

        delete_button = QPushButton(
            "Delete"
        )

        assign_button = QPushButton(
            "Assign"
        )

        stock_button = QPushButton(
            "Move to Stock"
        )

        reserve_button = QPushButton(
            "Reserve"
        )

        refresh_button = QPushButton(
            "Refresh"
        )

        upload_button = QPushButton(
                    "Upload"
                )

        if department_id is not None:
            assign_button.hide()
            stock_button.hide()
            reserve_button.hide()
            upload_button.hide()

        buttons_layout.addWidget(
            add_button
        )

        buttons_layout.addWidget(
            edit_button
        )

        buttons_layout.addWidget(
            delete_button
        )

        buttons_layout.addWidget(
            assign_button
        )

        buttons_layout.addWidget(
            stock_button
        )

        buttons_layout.addWidget(
            reserve_button
        )

        buttons_layout.addWidget(
            refresh_button
        )

        buttons_layout.addWidget(
                    upload_button
                )

        buttons_layout.addStretch()

        # -------------------------------------------------
        # SEARCH
        # -------------------------------------------------

        search_layout = QHBoxLayout()

        status_label = QLabel(
            "Status:"
        )

        self.computer_status_filter = QComboBox()

        self.computer_status_filter.addItem(
            "All",
            None,
        )

        for status in EQUIPMENT_STATUSES:
            self.computer_status_filter.addItem(
                status.replace("_", " ").title(),
                status,
            )

        department_label = QLabel(
            "Department:"
        )

        self.computer_department_filter = QComboBox()

        self.computer_department_filter.addItem(
            "All Departments",
            None,
        )

        for department in departments:

            self.computer_department_filter.addItem(
                department["name"],
                department["id"],
            )

        if department_id is not None:
            department_index = self.computer_department_filter.findData(
                department_id
            )
            if department_index >= 0:
                self.computer_department_filter.setCurrentIndex(
                    department_index
                )

        search_layout.addWidget(
            status_label
        )

        search_layout.addWidget(
            self.computer_status_filter
        )

        search_layout.addWidget(
            department_label
        )

        search_layout.addWidget(
            self.computer_department_filter
        )

        # -------------------------------------------------
        # TABLE
        # -------------------------------------------------

        self.computer_table = QTreeWidget()

        self.computer_table.setHeaderLabels(
            [
                "Computer",
                "Status",
                "Department",
                "User",
                "Location",
                "Last Seen",
                "Connection",
                "Manufacturer",
            ]
        )

        self.computer_table.setRootIsDecorated(
            False
        )

        self.computer_table.setAlternatingRowColors(
            True
        )

        # -------------------------------------------------
        # SIGNALS
        # -------------------------------------------------

        add_button.clicked.connect(
            self.add_computer
        )

        edit_button.clicked.connect(
            self.edit_selected_computer
        )

        delete_button.clicked.connect(
            self.delete_selected_computer
        )

        assign_button.clicked.connect(
            self.assign_selected_computer
        )

        stock_button.clicked.connect(
            self.move_selected_computer_to_stock
        )

        reserve_button.clicked.connect(
            self.reserve_selected_computer
        )

        refresh_button.clicked.connect(
            self.refresh_computers_view
        )

        upload_button.clicked.connect(
            self.upload_computers
        )

        self.computer_status_filter.currentIndexChanged.connect(
            self.load_computer_table
        )

        self.computer_department_filter.currentIndexChanged.connect(
            self.load_computer_table
        )

        self.computer_table.itemClicked.connect(
            self.computer_table_clicked
        )

        # -------------------------------------------------
        # LAYOUT
        # -------------------------------------------------

        computers_tab = QWidget()
        computers_layout = QVBoxLayout(computers_tab)
        computers_layout.addWidget(title)
        computers_layout.addWidget(subtitle)
        computers_layout.addLayout(buttons_layout)
        computers_layout.addLayout(search_layout)
        computers_layout.addWidget(self.computer_table)

        tabs = QTabWidget()
        tabs.addTab(computers_tab, "Computers")

        pending_tab = QWidget()
        pending_layout = QVBoxLayout(pending_tab)

        pending_actions = QHBoxLayout()
        approve_button = QPushButton("Approve")
        pending_refresh_button = QPushButton("Refresh")
        pending_actions.addWidget(approve_button)
        pending_actions.addWidget(pending_refresh_button)
        pending_actions.addStretch()

        self.pending_table = QTreeWidget()
        self.pending_table.setHeaderLabels(
            [
                "Computer",
                "Manufacturer",
                "Model",
                "CPU",
                "RAM",
                "Storage",
                "Windows",
                "OS Version",
                "IP Address",
                "MAC Address",
                "Last Seen",
            ]
        )
        self.pending_table.setRootIsDecorated(False)
        self.pending_table.setAlternatingRowColors(True)

        pending_layout.addWidget(QLabel("Review agent-discovered devices before adding them to Computers."))
        pending_layout.addLayout(pending_actions)
        pending_layout.addWidget(self.pending_table)
        tabs.addTab(pending_tab, "Pending Devices")

        approve_button.clicked.connect(self.approve_selected_pending_device)
        pending_refresh_button.clicked.connect(self.load_pending_devices)
        tabs.currentChanged.connect(
            lambda index: self.load_pending_devices() if index == 1 else None
        )

        layout.addWidget(tabs)

        self.workspace.addWidget(
            widget
        )

        self.workspace.setCurrentWidget(
            widget
        )

        self.load_computer_table()

    def load_pending_devices(self):
        if not hasattr(self, "pending_table"):
            return

        self.pending_table.clear()
        session = SessionLocal()
        try:
            for device in get_pending_devices(session):
                last_seen = device["last_seen"]
                if last_seen:
                    last_seen = datetime.fromisoformat(last_seen).strftime(
                        "%d %b %Y %H:%M"
                    )
                else:
                    last_seen = "NEVER CONNECTED"

                self.pending_table.addTopLevelItem(
                    QTreeWidgetItem(
                        [
                            device["computer_name"] or "-",
                            device["manufacturer"] or "-",
                            device["model"] or "-",
                            device["cpu"] or "-",
                            device["ram"] or "-",
                            device["storage"] or "-",
                            device["operating_system"] or "-",
                            device["os_version"] or "-",
                            device["ip_address"] or "-",
                            device["mac_address"] or "-",
                            last_seen,
                        ]
                    )
                )
                self.pending_table.topLevelItem(
                    self.pending_table.topLevelItemCount() - 1
                ).setData(0, Qt.UserRole, device["id"])
        finally:
            session.close()

    def approve_selected_pending_device(self):
        item = self.pending_table.currentItem()
        if item is None:
            QMessageBox.information(
                self,
                "Pending Devices",
                "Select a pending device first.",
            )
            return

        device_id = item.data(0, Qt.UserRole)
        session = SessionLocal()
        try:
            departments = get_departments(session)
        finally:
            session.close()

        dialog = PendingApprovalDialog(
            departments,
            parent=self,
        )
        if dialog.exec() != QDialog.Accepted:
            return

        approval_data = dialog.get_data()
        session = SessionLocal()
        try:
            approve_pending_device(
                session,
                device_id,
                approval_data,
            )
            session.commit()
        except Exception as exc:
            session.rollback()
            logging.exception(
                "Failed to approve pending device %s",
                device_id,
            )
            QMessageBox.critical(
                self,
                "Approval Failed",
                str(exc),
            )
            return
        finally:
            session.close()

        self.load_pending_devices()
        self.load_computer_table()
        QMessageBox.information(
            self,
            "Device Approved",
            "The device is now available in Computers.",
        )

    # =====================================================
    # LOAD COMPUTERS
    # =====================================================

    def load_computer_table(self):

        if not hasattr(
            self,
            "computer_table",
        ):
            return

        self.computer_table.clear()

        filters = {}

        status = (
            self.computer_status_filter
            .currentData()
        )

        if status:

            filters["status"] = status

        department_id = (
            self.computer_department_filter
            .currentData()
        )

        if department_id:

            filters["department_id"] = department_id

        session = SessionLocal()

        try:

            computers = get_computers(
                session,
                filters,
            )

            for computer in computers:

                last_seen = computer["last_seen"]
                if last_seen:
                    last_seen = datetime.fromisoformat(last_seen).strftime(
                        "%d %b %Y %H:%M"
                    )
                else:
                    last_seen = "NEVER CONNECTED"

                item = QTreeWidgetItem(
                    [
                        computer[
                            "computer_name"
                        ]
                        or "-",

                        computer[
                            "status"
                        ]
                        or "-",

                        computer[
                            "department_name"
                        ]
                        or "Stock",

                        computer[
                            "assigned_user"
                        ]
                        or "-",

                        computer[
                            "location"
                        ]
                        or "-",

                        last_seen,

                        computer[
                            "connection"
                        ]
                        or "NEVER CONNECTED",

                        computer[
                            "manufacturer"
                        ]
                        or "-",
                    ]
                )

                item.setData(
                    0,
                    Qt.UserRole,
                    computer["id"],
                )

                self.computer_table.addTopLevelItem(
                    item
                )

        finally:

            session.close()

    # =====================================================
    # COMPUTER CLICK
    # =====================================================

    def computer_table_clicked(
        self,
        item,
        _column,
    ):

        computer_id = item.data(
            0,
            Qt.UserRole,
        )

        if computer_id:

            self.show_computer(
                computer_id
            )

    # =====================================================
    # ADD COMPUTER
    # =====================================================

    def add_computer(self):

        session = SessionLocal()

        try:

            departments = get_departments(
                session
            )

        finally:

            session.close()

        dialog = ComputerDialog(
            departments,
            parent=self,
        )

        if self._computer_department_id is not None:
            department_index = dialog.department_combo.findData(
                self._computer_department_id
            )
            if department_index >= 0:
                dialog.department_combo.setCurrentIndex(
                    department_index
                )
                dialog.department_combo.setEnabled(False)

        if (
            dialog.exec()
            != QDialog.Accepted
        ):
            return

        data = dialog.get_data()

        if self._computer_department_id is not None:
            data["department_id"] = self._computer_department_id

        if not data["computer_name"]:

            QMessageBox.warning(
                self,
                "Invalid Computer",
                "Computer name is required.",
            )

            return

        session = SessionLocal()

        try:

            create_computer(
                session,
                data,
            )

            session.commit()

        except Exception as exc:

            session.rollback()

            QMessageBox.critical(
                self,
                "Error",
                f"Could not create computer:\n\n{exc}",
            )

            return

        finally:

            session.close()

        self.refresh_tree()

        self.show_computers(
            self._computer_department_id,
            self._computer_department_view,
        )

    # =====================================================
    # EDIT COMPUTER
    # =====================================================

    def edit_selected_computer(self):

        item = (
            self.computer_table.currentItem()
        )

        if item is None:

            QMessageBox.information(
                self,
                "Select Computer",
                "Please select a computer first.",
            )

            return

        computer_id = item.data(
            0,
            Qt.UserRole,
        )

        session = SessionLocal()

        try:

            computer = get_computer_data(
                session,
                computer_id,
            )

            departments = get_departments(
                session
            )

        finally:

            session.close()

        dialog = ComputerDialog(
            departments,
            computer,
            self,
        )

        if (
            dialog.exec()
            != QDialog.Accepted
        ):
            return

        data = dialog.get_data()

        if not data["computer_name"]:

            QMessageBox.warning(
                self,
                "Invalid Computer",
                "Computer name is required.",
            )

            return

        session = SessionLocal()

        try:

            update_computer(
                session,
                computer_id,
                data,
            )

            session.commit()

        except Exception as exc:

            session.rollback()

            QMessageBox.critical(
                self,
                "Error",
                f"Could not update computer:\n\n{exc}",
            )

            return

        finally:

            session.close()

        self.refresh_tree()

        self.show_computers(
            self._computer_department_id,
            self._computer_department_view,
        )

    # =====================================================
    # DELETE COMPUTER
    # =====================================================

    def delete_selected_computer(self):

        item = (
            self.computer_table.currentItem()
        )

        if item is None:

            QMessageBox.information(
                self,
                "Select Computer",
                "Please select a computer first.",
            )

            return

        computer_id = item.data(
            0,
            Qt.UserRole,
        )

        computer_name = item.text(
            0
        )

        confirmation = QMessageBox.question(
            self,
            "Delete Computer",
            (
                "Are you sure you want to delete "
                f"'{computer_name}'?"
            ),
            QMessageBox.Yes
            | QMessageBox.No,
        )

        if (
            confirmation
            != QMessageBox.Yes
        ):
            return

        session = SessionLocal()

        try:

            delete_computer(
                session,
                computer_id,
            )

            session.commit()

        except Exception as exc:

            session.rollback()

            QMessageBox.critical(
                self,
                "Error",
                f"Could not delete computer:\n\n{exc}",
            )

            return

        finally:

            session.close()

        self.refresh_tree()

        self.show_computers(
            self._computer_department_id,
            self._computer_department_view,
        )

    # =====================================================
    # ASSIGN COMPUTER
    # =====================================================

    def assign_selected_computer(self):

        item = (
            self.computer_table.currentItem()
        )

        if item is None:

            QMessageBox.information(
                self,
                "Select Computer",
                "Please select a computer first.",
            )

            return

        computer_id = item.data(
            0,
            Qt.UserRole,
        )

        session = SessionLocal()

        try:

            departments = get_departments(
                session
            )

        finally:

            session.close()

        if not departments:

            QMessageBox.warning(
                self,
                "No Departments",
                (
                    "Create a department before "
                    "assigning a computer."
                ),
            )

            return

        dialog = QDialog(
            self
        )

        dialog.setWindowTitle(
            "Assign Computer"
        )

        dialog.setMinimumWidth(
            420
        )

        department_combo = QComboBox()

        user_edit = QLineEdit()

        location_edit = QLineEdit()

        for department in departments:

            department_combo.addItem(
                department["name"],
                department["id"],
            )

        form = QFormLayout()

        form.addRow(
            "Department:",
            department_combo,
        )

        form.addRow(
            "Assigned User:",
            user_edit,
        )

        form.addRow(
            "Location:",
            location_edit,
        )

        buttons = QDialogButtonBox(
            QDialogButtonBox.Save
            | QDialogButtonBox.Cancel
        )

        buttons.accepted.connect(
            dialog.accept
        )

        buttons.rejected.connect(
            dialog.reject
        )

        layout = QVBoxLayout(
            dialog
        )

        layout.addLayout(
            form
        )

        layout.addWidget(
            buttons
        )

        if (
            dialog.exec()
            != QDialog.Accepted
        ):
            return

        department_id = (
            department_combo.currentData()
        )

        assigned_user = (
            user_edit.text()
            .strip()
        )

        location = (
            location_edit.text()
            .strip()
        )

        if not assigned_user:

            QMessageBox.warning(
                self,
                "Invalid Assignment",
                "Assigned user is required.",
            )

            return

        if not location:

            location = "Assigned Location"

        session = SessionLocal()

        try:

            assign_computer(
                session,
                computer_id,
                department_id,
                assigned_user,
                location,
            )

            session.commit()

        except Exception as exc:

            session.rollback()

            QMessageBox.critical(
                self,
                "Error",
                f"Could not assign computer:\n\n{exc}",
            )

            return

        finally:

            session.close()

        self.refresh_tree()

        self.show_computers()

    # =====================================================
    # MOVE TO STOCK
    # =====================================================

    def move_selected_computer_to_stock(
        self
    ):

        item = (
            self.computer_table.currentItem()
        )

        if item is None:

            QMessageBox.information(
                self,
                "Select Computer",
                "Please select a computer first.",
            )

            return

        computer_id = item.data(
            0,
            Qt.UserRole,
        )

        confirmation = QMessageBox.question(
            self,
            "Move to Stock",
            (
                "Move the selected computer "
                "back to IT Stock?"
            ),
            QMessageBox.Yes
            | QMessageBox.No,
        )

        if (
            confirmation
            != QMessageBox.Yes
        ):
            return

        session = SessionLocal()

        try:

            move_to_stock(
                session,
                computer_id,
            )

            session.commit()

        except Exception as exc:

            session.rollback()

            QMessageBox.critical(
                self,
                "Error",
                (
                    "Could not move computer "
                    f"to stock:\n\n{exc}"
                ),
            )

            return

        finally:

            session.close()

        self.refresh_tree()

        self.show_computers()

    # =====================================================
    # RESERVE COMPUTER
    # =====================================================

    def reserve_selected_computer(self):

        item = (
            self.computer_table.currentItem()
        )

        if item is None:

            QMessageBox.information(
                self,
                "Select Computer",
                "Please select a computer first.",
            )

            return

        computer_id = item.data(
            0,
            Qt.UserRole,
        )

        session = SessionLocal()

        try:

            reserve_computer(
                session,
                computer_id,
            )

            session.commit()

        except Exception as exc:

            session.rollback()

            QMessageBox.critical(
                self,
                "Error",
                f"Could not reserve computer:\n\n{exc}",
            )

            return

        finally:

            session.close()

        self.refresh_tree()

        self.show_computers()

    # =====================================================
    # REFRESH COMPUTERS
    # =====================================================

    def refresh_computers_view(self):

        self.refresh_tree()

        self.show_computers(
            self._computer_department_id,
            self._computer_department_view,
        )



    # =====================================================
# UPLOAD COMPUTERS
# =====================================================

    # =====================================================
    # UPLOAD COMPUTERS
    # =====================================================
    def upload_computers(self):

        file_path, _ = QFileDialog.getOpenFileName(
            self,
            "Select Computer Excel File",
            "",
            "Excel Files (*.xlsx *.xls)"
        )

        if not file_path:
            return

        try:
            dataframe = read_excel_dataframe(file_path)
        except Exception as exc:
            QMessageBox.critical(
                self,
                "Excel Error",
                (
                    "Could not read the Excel file.\n\n"
                    f"Error:\n{exc}"
                ),
            )
            return

        columns = list(dataframe.columns)
        detected_mappings = detect_auto_mapping(columns)
        profile = detect_mapping_profile(columns)

        if profile:
            profile_mapping = build_mapping_from_profile(profile, columns)
            required_target_present = any(
                target == "computer_name"
                for target in profile_mapping.values()
            )

            if required_target_present:
                choice = QMessageBox.question(
                    self,
                    f"Existing mapping profile detected: {profile.get('name')}",
                    (
                        f"An existing mapping profile was detected for this Excel file.\n\n"
                        f"Would you like to apply the profile '{profile.get('name')}' or review the mapping manually?"
                    ),
                    QMessageBox.Yes | QMessageBox.No,
                    QMessageBox.No,
                )

                if choice == QMessageBox.Yes:
                    detected_mappings = profile_mapping
            else:
                QMessageBox.warning(
                    self,
                    "Saved profile does not fully match this Excel file.",
                    "The saved profile does not fully match this Excel file.\n\nThe Mapping Wizard will open for manual review.",
                )

        dialog = ExcelMappingDialog(
            self,
            dataframe,
            detected_mappings=detected_mappings,
            profile=profile,
        )

        result = dialog.exec_()

        if result == 1:
            self.refresh_tree()
            self.show_computers()



    # =====================================================
    # PHONES
    # =====================================================

    def show_department_phones(self, department_id):

        self.show_phones(
            department_id,
            department_view=True,
        )

    def show_phones(
        self,
        department_id=None,
        department_view=False,
    ):

        self._phone_department_id = department_id
        self._phone_department_view = department_view

        session = SessionLocal()

        try:

            departments = get_departments(
                session
            )

        finally:

            session.close()

        department_name = next(
            (
                department["name"]
                for department in departments
                if department["id"] == department_id
            ),
            None,
        )

        widget = QWidget()

        layout = QVBoxLayout(
            widget
        )

        title = QLabel(
            (
                f"Phones — {department_name}"
                if department_name
                else "Phones"
            )
        )

        title.setStyleSheet(
            "font-size: 24px; "
            "font-weight: 600;"
        )

        subtitle = QLabel(
            "Manage hotel phones, assignments, stock and status."
        )

        # -------------------------------------------------
        # BUTTONS
        # -------------------------------------------------

        buttons_layout = QHBoxLayout()

        add_button = QPushButton(
            "+ Add Phone"
        )

        edit_button = QPushButton(
            "Edit"
        )

        delete_button = QPushButton(
            "Delete"
        )

        assign_button = QPushButton(
            "Assign"
        )

        stock_button = QPushButton(
            "Move to Stock"
        )

        reserve_button = QPushButton(
            "Reserve"
        )

        refresh_button = QPushButton(
            "Refresh"
        )

        upload_button = QPushButton(
            "Upload"
        )

        if department_view:
            assign_button.hide()
            stock_button.hide()
            reserve_button.hide()
            upload_button.hide()

        buttons_layout.addWidget(
            add_button
        )

        buttons_layout.addWidget(
            edit_button
        )

        buttons_layout.addWidget(
            delete_button
        )

        buttons_layout.addWidget(
            assign_button
        )

        buttons_layout.addWidget(
            stock_button
        )

        buttons_layout.addWidget(
            reserve_button
        )

        buttons_layout.addWidget(
            refresh_button
        )

        buttons_layout.addWidget(
            upload_button
        )

        buttons_layout.addStretch()

        # -------------------------------------------------
        # SEARCH
        # -------------------------------------------------

        search_layout = QHBoxLayout()

        status_label = QLabel(
            "Status:"
        )

        self.phone_status_filter = QComboBox()

        self.phone_status_filter.addItem(
            "All",
            None,
        )

        for status in EQUIPMENT_STATUSES:
            self.phone_status_filter.addItem(
                status.replace("_", " ").title(),
                status,
            )

        department_label = QLabel(
            "Department:"
        )

        self.phone_department_filter = QComboBox()

        self.phone_department_filter.addItem(
            "All Departments",
            None,
        )

        for department in departments:

            self.phone_department_filter.addItem(
                department["name"],
                department["id"],
            )

        if department_id is not None:
            department_index = self.phone_department_filter.findData(
                department_id
            )
            if department_index >= 0:
                self.phone_department_filter.setCurrentIndex(
                    department_index
                )

        search_layout.addWidget(
            status_label
        )

        search_layout.addWidget(
            self.phone_status_filter
        )

        search_layout.addWidget(
            department_label
        )

        search_layout.addWidget(
            self.phone_department_filter
        )

        # -------------------------------------------------
        # TABLE
        # -------------------------------------------------

        self.phone_table = QTreeWidget()

        self.phone_table.setHeaderLabels(
            [
                "Asset Tag",
                "Phone Number",
                "Department",
                "Assigned User",
                "Location",
                "Brand",
                "Model",
                "Status",
            ]
        )

        self.phone_table.setRootIsDecorated(
            False
        )

        self.phone_table.setAlternatingRowColors(
            True
        )

        # -------------------------------------------------
        # SIGNALS
        # -------------------------------------------------

        add_button.clicked.connect(
            self.add_phone
        )

        edit_button.clicked.connect(
            self.edit_selected_phone
        )

        delete_button.clicked.connect(
            self.delete_selected_phone
        )

        assign_button.clicked.connect(
            self.assign_selected_phone
        )

        stock_button.clicked.connect(
            self.move_selected_phone_to_stock
        )

        reserve_button.clicked.connect(
            self.reserve_selected_phone
        )

        upload_button.clicked.connect(
            self.upload_phones
        )

        refresh_button.clicked.connect(
            self.refresh_phones_view
        )

        self.phone_status_filter.currentIndexChanged.connect(
            self.load_phone_table
        )

        self.phone_department_filter.currentIndexChanged.connect(
            self.load_phone_table
        )

        self.phone_table.itemClicked.connect(
            self.phone_table_clicked
        )

        # -------------------------------------------------
        # LAYOUT
        # -------------------------------------------------

        layout.addWidget(
            title
        )

        layout.addWidget(
            subtitle
        )

        layout.addLayout(
            buttons_layout
        )

        layout.addLayout(
            search_layout
        )

        layout.addWidget(
            self.phone_table
        )

        self.workspace.addWidget(
            widget
        )

        self.workspace.setCurrentWidget(
            widget
        )

        self.load_phone_table()


    # =====================================================
    # LOAD PHONES
    # =====================================================

    def load_phone_table(self):

        if not hasattr(
            self,
            "phone_table",
        ):
            return

        self.phone_table.clear()

        filters = {}

        status = (
            self.phone_status_filter
            .currentData()
        )

        if status:

            filters["status"] = status

        department_id = (
            self.phone_department_filter
            .currentData()
        )

        if department_id:

            filters["department_id"] = department_id

        session = SessionLocal()

        try:

            phones = get_phones(
                session,
                filters,
            )

            for phone in phones:

                item = QTreeWidgetItem(
                    [
                        phone["asset_tag"] or "-",

                        phone["phone_number"] or "-",

                        phone["department_name"]
                        or "Stock",

                        phone["assigned_user"]
                        or "-",

                        phone["location"]
                        or "-",

                        phone["brand"]
                        or "-",

                        phone["model"]
                        or "-",

                        phone["status"]
                        or "-",
                    ]
                )

                item.setData(
                    0,
                    Qt.UserRole,
                    phone["id"],
                )

                self.phone_table.addTopLevelItem(
                    item
                )

        finally:

            session.close()


    # =====================================================
    # PHONE CLICK
    # =====================================================

    def phone_table_clicked(
        self,
        item,
        _column,
    ):

        phone_id = item.data(
            0,
            Qt.UserRole,
        )

        if phone_id:

            # For now, open the edit dialog
            self.edit_selected_phone()


    # =====================================================
    # ADD PHONE
    # =====================================================

    def _selected_phone_id(self):

        item = self.phone_table.currentItem()

        if item is None:
            QMessageBox.information(
                self,
                "Select Phone",
                "Please select a phone first.",
            )
            return None

        return item.data(
            0,
            Qt.UserRole,
        )

    def assign_selected_phone(self):

        phone_id = self._selected_phone_id()
        if phone_id is None:
            return

        session = SessionLocal()
        try:
            departments = get_departments(session)
        finally:
            session.close()

        if not departments:
            QMessageBox.warning(
                self,
                "No Departments",
                "Create a department before assigning a phone.",
            )
            return

        dialog = QDialog(self)
        dialog.setWindowTitle("Assign Phone")
        department_combo = QComboBox()
        user_edit = QLineEdit()
        location_edit = QLineEdit()

        for department in departments:
            department_combo.addItem(
                department["name"],
                department["id"],
            )

        form = QFormLayout()
        form.addRow("Department:", department_combo)
        form.addRow("Assigned User:", user_edit)
        form.addRow("Location:", location_edit)

        buttons = QDialogButtonBox(
            QDialogButtonBox.Save | QDialogButtonBox.Cancel
        )
        buttons.accepted.connect(dialog.accept)
        buttons.rejected.connect(dialog.reject)

        layout = QVBoxLayout(dialog)
        layout.addLayout(form)
        layout.addWidget(buttons)

        if dialog.exec() != QDialog.Accepted:
            return

        assigned_user = user_edit.text().strip()
        if not assigned_user:
            QMessageBox.warning(
                self,
                "Invalid Assignment",
                "Assigned user is required.",
            )
            return

        location = location_edit.text().strip() or "Assigned Location"
        session = SessionLocal()
        try:
            assign_phone(
                session,
                phone_id,
                department_combo.currentData(),
                assigned_user,
                location,
            )
            session.commit()
        except Exception as exc:
            session.rollback()
            QMessageBox.critical(
                self,
                "Error",
                f"Could not assign phone:\n\n{exc}",
            )
            return
        finally:
            session.close()

        self.refresh_tree()
        self.show_phones(
            self._phone_department_id,
            self._phone_department_view,
        )

    def move_selected_phone_to_stock(self):

        phone_id = self._selected_phone_id()
        if phone_id is None:
            return

        confirmation = QMessageBox.question(
            self,
            "Move to Stock",
            "Move the selected phone back to IT Stock?",
            QMessageBox.Yes | QMessageBox.No,
        )
        if confirmation != QMessageBox.Yes:
            return

        session = SessionLocal()
        try:
            move_phone_to_stock(session, phone_id)
            session.commit()
        except Exception as exc:
            session.rollback()
            QMessageBox.critical(
                self,
                "Error",
                f"Could not move phone to stock:\n\n{exc}",
            )
            return
        finally:
            session.close()

        self.refresh_tree()
        self.show_phones(
            self._phone_department_id,
            self._phone_department_view,
        )

    def reserve_selected_phone(self):

        phone_id = self._selected_phone_id()
        if phone_id is None:
            return

        session = SessionLocal()
        try:
            reserve_phone(session, phone_id)
            session.commit()
        except Exception as exc:
            session.rollback()
            QMessageBox.critical(
                self,
                "Error",
                f"Could not reserve phone:\n\n{exc}",
            )
            return
        finally:
            session.close()

        self.refresh_tree()
        self.show_phones(
            self._phone_department_id,
            self._phone_department_view,
        )

    def upload_phones(self):

        file_path, _ = QFileDialog.getOpenFileName(
            self,
            "Select Phone Excel File",
            "",
            "Excel Files (*.xlsx *.xls)",
        )
        if not file_path:
            return

        try:
            dataframe = read_excel_dataframe(file_path)
        except Exception as exc:
            QMessageBox.critical(
                self,
                "Excel Error",
                f"Could not read the Excel file.\n\n{exc}",
            )
            return

        columns = {
            str(column).strip().lower().replace(" ", "_"): column
            for column in dataframe.columns
        }
        if "asset_tag" not in columns or "phone_number" not in columns:
            QMessageBox.warning(
                self,
                "Invalid Phone File",
                "The file must contain asset_tag and phone_number columns.",
            )
            return

        phone_fields = [
            "asset_tag", "phone_number", "brand", "model", "imei_1",
            "imei_2", "serial_number", "operating_system", "os_version",
            "mac_address", "ip_address", "assigned_user", "location",
            "status", "notes",
        ]

        def value(row, field):
            source = columns.get(field)
            item = row[source] if source is not None else None
            return None if pd.isna(item) else str(item).strip() or None

        session = SessionLocal()
        created = 0
        try:
            for _, row in dataframe.iterrows():
                payload = {
                    field: value(row, field)
                    for field in phone_fields
                }
                if not payload["asset_tag"] or not payload["phone_number"]:
                    continue
                create_phone(session, payload)
                created += 1
            session.commit()
        except Exception as exc:
            session.rollback()
            QMessageBox.critical(
                self,
                "Phone Upload Error",
                f"Could not import phones:\n\n{exc}",
            )
            return
        finally:
            session.close()

        QMessageBox.information(
            self,
            "Phone Upload Complete",
            f"Imported {created} phone(s).",
        )
        self.refresh_tree()
        self.show_phones(
            self._phone_department_id,
            self._phone_department_view,
        )

    def add_phone(self):

        session = SessionLocal()

        try:

            departments = get_departments(
                session
            )

        finally:

            session.close()

        dialog = PhoneDialog(
            departments,
            parent=self,
        )

        if self._phone_department_id is not None:
            department_index = dialog.department_combo.findData(
                self._phone_department_id
            )
            if department_index >= 0:
                dialog.department_combo.setCurrentIndex(
                    department_index
                )
                dialog.department_combo.setEnabled(False)

        if (
            dialog.exec()
            != QDialog.Accepted
        ):
            return

        data = dialog.get_data()

        if self._phone_department_id is not None:
            data["department_id"] = self._phone_department_id

        if not data["asset_tag"]:

            QMessageBox.warning(
                self,
                "Invalid Phone",
                "Asset tag is required.",
            )

            return

        if not data["phone_number"]:

            QMessageBox.warning(
                self,
                "Invalid Phone",
                "Phone number is required.",
            )

            return

        session = SessionLocal()

        try:

            create_phone(
                session,
                data,
            )

            session.commit()

        except Exception as exc:

            session.rollback()

            QMessageBox.critical(
                self,
                "Error",
                f"Could not create phone:\n\n{exc}",
            )

            return

        finally:

            session.close()

        self.refresh_tree()

        self.show_phones(
            self._phone_department_id,
            self._phone_department_view,
        )


    # =====================================================
    # EDIT PHONE
    # =====================================================

    def edit_selected_phone(self):

        item = (
            self.phone_table.currentItem()
        )

        if item is None:

            QMessageBox.information(
                self,
                "Select Phone",
                "Please select a phone first.",
            )

            return

        phone_id = item.data(
            0,
            Qt.UserRole,
        )

        session = SessionLocal()

        try:

            phone = get_phone(
                session,
                phone_id,
            )

            departments = get_departments(
                session
            )

        finally:

            session.close()

        if not phone:

            QMessageBox.warning(
                self,
                "Phone Not Found",
                "The selected phone could not be found.",
            )

            return

        dialog = PhoneDialog(
            departments,
            phone,
            self,
        )

        if (
            dialog.exec()
            != QDialog.Accepted
        ):
            return

        data = dialog.get_data()

        if not data["asset_tag"]:

            QMessageBox.warning(
                self,
                "Invalid Phone",
                "Asset tag is required.",
            )

            return

        if not data["phone_number"]:

            QMessageBox.warning(
                self,
                "Invalid Phone",
                "Phone number is required.",
            )

            return

        session = SessionLocal()

        try:

            update_phone(
                session,
                phone_id,
                data,
            )

            session.commit()

        except Exception as exc:

            session.rollback()

            QMessageBox.critical(
                self,
                "Error",
                f"Could not update phone:\n\n{exc}",
            )

            return

        finally:

            session.close()

        self.refresh_tree()

        self.show_phones(
            self._phone_department_id,
            self._phone_department_view,
        )


    # =====================================================
    # DELETE PHONE
    # =====================================================

    def delete_selected_phone(self):

        item = (
            self.phone_table.currentItem()
        )

        if item is None:

            QMessageBox.information(
                self,
                "Select Phone",
                "Please select a phone first.",
            )

            return

        phone_id = item.data(
            0,
            Qt.UserRole,
        )

        asset_tag = item.text(
            0
        )

        confirmation = QMessageBox.question(
            self,
            "Delete Phone",
            (
                "Are you sure you want to delete "
                f"phone '{asset_tag}'?"
            ),
            QMessageBox.Yes
            | QMessageBox.No,
        )

        if (
            confirmation
            != QMessageBox.Yes
        ):
            return

        session = SessionLocal()

        try:

            delete_phone(
                session,
                phone_id,
            )

            session.commit()

        except Exception as exc:

            session.rollback()

            QMessageBox.critical(
                self,
                "Error",
                f"Could not delete phone:\n\n{exc}",
            )

            return

        finally:

            session.close()

        self.refresh_tree()

        self.show_phones(
            self._phone_department_id,
            self._phone_department_view,
        )


    # =====================================================
    # REFRESH PHONES
    # =====================================================

    def refresh_phones_view(self):

        self.refresh_tree()

        self.show_phones(
            self._phone_department_id,
            self._phone_department_view,
        )




    # =====================================================
    # TELEPHONES
    # =====================================================

    def show_telephones(self, department_id=None, department_view=False):
        self._telephone_department_id = department_id
        self._telephone_department_view = department_view
        widget = QWidget()
        layout = QVBoxLayout(widget)
        title = QLabel("Telephones")
        title.setStyleSheet("font-size: 24px; font-weight: 600;")
        buttons = QHBoxLayout()
        add_button = QPushButton("+ Add Telephone")
        edit_button = QPushButton("Edit")
        delete_button = QPushButton("Delete")
        stock_button = QPushButton("Move to Stock")
        assign_button = QPushButton("Assign")
        reserve_button = QPushButton("Reserve")
        upload_button = QPushButton("Upload Excel")
        refresh_button = QPushButton("Refresh")
        for button in (add_button, edit_button, delete_button, assign_button, stock_button, reserve_button, upload_button, refresh_button):
            buttons.addWidget(button)
        buttons.addStretch()
        status_filter = QComboBox()
        status_filter.addItem("All Statuses", None)
        for status in EQUIPMENT_STATUSES:
            status_filter.addItem(status.replace("_", " ").title(), status)
        department_filter = QComboBox()
        department_filter.addItem("All Departments", None)
        session = SessionLocal()
        try:
            for department in get_departments(session):
                department_filter.addItem(department["name"], department["id"])
        finally:
            session.close()
        table = QTreeWidget()
        table.setRootIsDecorated(False)
        table.setAlternatingRowColors(True)
        table.setHeaderLabels(["Asset Tag", "Number / Extension", "Department", "Assigned User", "Location", "Brand", "Model", "Status"])
        self.telephone_status_filter = status_filter
        self.telephone_department_filter = department_filter
        self.telephone_table = table
        add_button.clicked.connect(self.add_telephone)
        edit_button.clicked.connect(self.edit_selected_telephone)
        delete_button.clicked.connect(self.delete_selected_telephone)
        stock_button.clicked.connect(self.move_selected_telephone_to_stock)
        assign_button.clicked.connect(self.assign_selected_telephone)
        reserve_button.clicked.connect(self.reserve_selected_telephone)
        upload_button.clicked.connect(self.upload_telephones)
        refresh_button.clicked.connect(lambda: self.show_telephones(department_id, department_view))
        status_filter.currentIndexChanged.connect(self.load_telephone_table)
        department_filter.currentIndexChanged.connect(self.load_telephone_table)
        table.itemDoubleClicked.connect(lambda *_: self.edit_selected_telephone())
        layout.addWidget(title)
        layout.addWidget(QLabel("Manage desk phones, extensions, assignments, stock and status."))
        layout.addLayout(buttons)
        filter_layout = QHBoxLayout()
        filter_layout.addWidget(QLabel("Department:"))
        filter_layout.addWidget(department_filter)
        filter_layout.addWidget(QLabel("Status:"))
        filter_layout.addWidget(status_filter)
        layout.addLayout(filter_layout)
        layout.addWidget(table)
        self.workspace.addWidget(widget)
        self.workspace.setCurrentWidget(widget)
        self.load_telephone_table()

    def load_telephone_table(self):
        if not hasattr(self, "telephone_table"):
            return
        self.telephone_table.clear()
        session = SessionLocal()
        try:
            filters = {}
            status = self.telephone_status_filter.currentData()
            if status:
                filters["status"] = status
            selected_department = self.telephone_department_filter.currentData()
            if selected_department:
                filters["department_id"] = selected_department
            if getattr(self, "_telephone_department_id", None) is not None:
                filters["department_id"] = self._telephone_department_id
            for telephone in get_telephones(session, filters):
                self.telephone_table.addTopLevelItem(QTreeWidgetItem([
                    telephone["asset_tag"] or "-",
                    telephone["telephone_number"] or "-",
                    telephone["department_name"] or "Stock",
                    telephone["assigned_user"] or "-",
                    telephone["location"] or "-",
                    telephone["brand"] or "-",
                    telephone["model"] or "-",
                    telephone["status"] or "-",
                ]))
                self.telephone_table.topLevelItem(self.telephone_table.topLevelItemCount() - 1).setData(0, Qt.UserRole, telephone["id"])
        finally:
            session.close()

    def _selected_telephone_id(self):
        item = self.telephone_table.currentItem()
        if item is None:
            QMessageBox.information(self, "Select Telephone", "Please select a telephone first.")
            return None
        return item.data(0, Qt.UserRole)

    def add_telephone(self):
        session = SessionLocal()
        try:
            departments = get_departments(session)
        finally:
            session.close()
        dialog = TelephoneDialog(departments, parent=self)
        if dialog.exec() != QDialog.Accepted:
            return
        data = dialog.get_data()
        if self._telephone_department_id is not None:
            data["department_id"] = self._telephone_department_id
        if not data["asset_tag"] or not data["telephone_number"]:
            QMessageBox.warning(self, "Invalid Telephone", "Asset tag and telephone number are required.")
            return
        session = SessionLocal()
        try:
            create_telephone(session, data)
            session.commit()
        except Exception as exc:
            session.rollback()
            QMessageBox.critical(self, "Error", f"Could not create telephone:\n\n{exc}")
            return
        finally:
            session.close()
        self.refresh_tree()
        self.show_telephones(self._telephone_department_id, self._telephone_department_view)

    def edit_selected_telephone(self):
        telephone_id = self._selected_telephone_id()
        if telephone_id is None:
            return
        session = SessionLocal()
        try:
            telephone = get_telephone(session, telephone_id)
            departments = get_departments(session)
        finally:
            session.close()
        dialog = TelephoneDialog(departments, telephone, self)
        if dialog.exec() != QDialog.Accepted:
            return
        data = dialog.get_data()
        session = SessionLocal()
        try:
            update_telephone(session, telephone_id, data)
            session.commit()
        except Exception as exc:
            session.rollback()
            QMessageBox.critical(self, "Error", f"Could not update telephone:\n\n{exc}")
            return
        finally:
            session.close()
        self.refresh_tree()
        self.show_telephones(self._telephone_department_id, self._telephone_department_view)

    def delete_selected_telephone(self):
        telephone_id = self._selected_telephone_id()
        if telephone_id is None or QMessageBox.question(self, "Delete Telephone", "Delete the selected telephone?", QMessageBox.Yes | QMessageBox.No) != QMessageBox.Yes:
            return
        session = SessionLocal()
        try:
            delete_telephone(session, telephone_id)
            session.commit()
        except Exception as exc:
            session.rollback()
            QMessageBox.critical(self, "Error", f"Could not delete telephone:\n\n{exc}")
            return
        finally:
            session.close()
        self.refresh_tree()
        self.show_telephones(self._telephone_department_id, self._telephone_department_view)

    def move_selected_telephone_to_stock(self):
        telephone_id = self._selected_telephone_id()
        if telephone_id is None:
            return
        session = SessionLocal()
        try:
            move_telephone_to_stock(session, telephone_id)
            session.commit()
        except Exception as exc:
            session.rollback()
            QMessageBox.critical(self, "Error", f"Could not move telephone to stock:\n\n{exc}")
            return
        finally:
            session.close()
        self.refresh_tree()
        self.show_telephones(self._telephone_department_id, self._telephone_department_view)

    def assign_selected_telephone(self):
        telephone_id = self._selected_telephone_id()
        if telephone_id is None:
            return
        session = SessionLocal()
        try:
            departments = get_departments(session)
        finally:
            session.close()
        dialog = QDialog(self)
        department_combo = QComboBox()
        user_edit = QLineEdit()
        location_edit = QLineEdit()
        for department in departments:
            department_combo.addItem(department["name"], department["id"])
        form = QFormLayout()
        form.addRow("Department:", department_combo)
        form.addRow("Assigned User:", user_edit)
        form.addRow("Location:", location_edit)
        buttons = QDialogButtonBox(QDialogButtonBox.Save | QDialogButtonBox.Cancel)
        buttons.accepted.connect(dialog.accept)
        buttons.rejected.connect(dialog.reject)
        dialog_layout = QVBoxLayout(dialog)
        dialog_layout.addLayout(form)
        dialog_layout.addWidget(buttons)
        if dialog.exec() != QDialog.Accepted:
            return
        session = SessionLocal()
        try:
            assign_telephone(
                session,
                telephone_id,
                department_combo.currentData(),
                user_edit.text().strip(),
                location_edit.text().strip() or "Assigned Location",
            )
            session.commit()
        except Exception as exc:
            session.rollback()
            QMessageBox.critical(self, "Error", f"Could not assign telephone:\n\n{exc}")
            return
        finally:
            session.close()
        self.refresh_tree()
        self.show_telephones(self._telephone_department_id, self._telephone_department_view)

    def reserve_selected_telephone(self):
        telephone_id = self._selected_telephone_id()
        if telephone_id is None:
            return
        session = SessionLocal()
        try:
            reserve_telephone(session, telephone_id)
            session.commit()
        except Exception as exc:
            session.rollback()
            QMessageBox.critical(self, "Error", f"Could not reserve telephone:\n\n{exc}")
            return
        finally:
            session.close()
        self.refresh_tree()
        self.show_telephones(self._telephone_department_id, self._telephone_department_view)

    def upload_telephones(self):
        file_path, _ = QFileDialog.getOpenFileName(
            self, "Select Telephone Excel File", "", "Excel Files (*.xlsx *.xls)"
        )
        if not file_path:
            return
        try:
            dataframe = read_excel_dataframe(file_path)
            session = SessionLocal()
            try:
                result = import_telephones(session, dataframe)
                session.commit()
            except Exception:
                session.rollback()
                raise
            finally:
                session.close()
        except Exception as exc:
            QMessageBox.critical(self, "Excel Import Error", f"Could not import telephones:\n\n{exc}")
            return
        QMessageBox.information(
            self,
            "Import Complete",
            f"Imported {result['created']} telephone(s); skipped {result['skipped']}.",
        )
        self.refresh_tree()
        self.show_telephones(self._telephone_department_id, self._telephone_department_view)

    def show_equipment_management(self, equipment_key, department_id=None):
        session = SessionLocal()
        try:
            equipment_types = get_equipment_types(session)
            equipment_type = next(item for item in equipment_types if item.key == equipment_key)
            departments = get_departments(session)
            filter_options = get_equipment_filter_options(session)
        finally:
            session.close()
        widget = QWidget()
        layout = QVBoxLayout(widget)
        title = QLabel(equipment_type.name)
        title.setStyleSheet("font-size: 24px; font-weight: 600;")
        buttons = QHBoxLayout()
        add_button = QPushButton(f"+ Add {equipment_type.name[:-1] if equipment_type.name.endswith('s') else equipment_type.name}")
        upload_button = QPushButton("Upload Excel")
        edit_button = QPushButton("Edit")
        delete_button = QPushButton("Delete")
        assign_button = QPushButton("Assign")
        stock_button = QPushButton("Move to Stock")
        reserve_button = QPushButton("Reserve")
        refresh_button = QPushButton("Refresh")
        for button in (add_button, upload_button, edit_button, delete_button, assign_button, stock_button, reserve_button, refresh_button):
            buttons.addWidget(button)
        buttons.addStretch()
        status_filter = QComboBox()
        status_filter.addItem("All Statuses", None)
        for status in filter_options["statuses"]:
            status_filter.addItem(status.replace("_", " ").title(), status)
        department_filter = QComboBox()
        department_filter.addItem("All Departments", None)
        for department in filter_options["departments"]:
            department_filter.addItem(department.name, department.id)
        if department_id is not None:
            department_filter.setCurrentIndex(department_filter.findData(department_id))
        table = QTreeWidget()
        table.setRootIsDecorated(False)
        table.setAlternatingRowColors(True)
        table.setHeaderLabels(["Asset Tag", "Name", "Brand", "Model", "Serial", "Department", "Status", "Location"])
        self.generic_equipment_key = equipment_key
        self.generic_equipment_table = table
        self.generic_equipment_department = department_filter
        self.generic_equipment_status = status_filter
        self._generic_equipment_department_id = department_id
        add_button.clicked.connect(lambda: self.add_generic_equipment(equipment_key, departments, department_id))
        upload_button.clicked.connect(lambda: self.upload_generic_equipment(equipment_key, department_id))
        edit_button.clicked.connect(lambda: self.edit_selected_generic_equipment(equipment_key, departments, department_id))
        delete_button.clicked.connect(lambda: self.delete_selected_generic_equipment(equipment_key, department_id))
        assign_button.clicked.connect(lambda: self.assign_selected_generic_equipment(equipment_key, department_id))
        stock_button.clicked.connect(lambda: self.move_selected_generic_equipment_to_stock(equipment_key, department_id))
        reserve_button.clicked.connect(lambda: self.reserve_selected_generic_equipment(equipment_key, department_id))
        refresh_button.clicked.connect(lambda: self.show_equipment_management(equipment_key, department_id))
        department_filter.currentIndexChanged.connect(self.load_generic_equipment_table)
        status_filter.currentIndexChanged.connect(self.load_generic_equipment_table)
        layout.addWidget(title)
        layout.addLayout(buttons)
        filter_layout = QHBoxLayout()
        filter_layout.addWidget(QLabel("Department:"))
        filter_layout.addWidget(department_filter)
        filter_layout.addWidget(QLabel("Status:"))
        filter_layout.addWidget(status_filter)
        layout.addLayout(filter_layout)
        layout.addWidget(table)
        self.workspace.addWidget(widget)
        self.workspace.setCurrentWidget(widget)
        self.load_generic_equipment_table()

    def load_generic_equipment_table(self):
        if not hasattr(self, "generic_equipment_table"):
            return
        self.generic_equipment_table.clear()
        session = SessionLocal()
        try:
            filters = {}
            status = self.generic_equipment_status.currentData()
            if status:
                filters["status"] = status
            selected_department = self.generic_equipment_department.currentData()
            if selected_department:
                filters["department_id"] = selected_department
            for asset in get_equipment_assets(session, self.generic_equipment_key, filters):
                item = QTreeWidgetItem([
                    asset["asset_tag"] or "-", asset["name"] or "-", asset["brand"] or "-",
                    asset["model"] or "-", asset["serial_number"] or "-",
                    asset["department_name"] or "Stock", asset["status"] or "-",
                    asset["location"] or "-",
                ])
                item.setData(0, Qt.UserRole, asset["id"])
                self.generic_equipment_table.addTopLevelItem(item)
        finally:
            session.close()

    def _selected_generic_equipment_id(self):
        item = self.generic_equipment_table.currentItem()
        if item is None:
            QMessageBox.information(self, "Select Equipment", "Please select an equipment item first.")
            return None
        return item.data(0, Qt.UserRole)

    def add_generic_equipment(self, equipment_key, departments, department_id=None):
        dialog = GenericEquipmentDialog(
            self.generic_equipment_key.replace("_", " ").title(), departments, parent=self
        )
        if dialog.exec() != QDialog.Accepted:
            return
        data = dialog.get_data()
        if department_id is not None:
            data["department_id"] = department_id
        session = SessionLocal()
        try:
            create_equipment_asset(session, equipment_key, data)
            session.commit()
        except Exception as exc:
            session.rollback()
            QMessageBox.critical(self, "Error", f"Could not create equipment:\n\n{exc}")
            return
        finally:
            session.close()
        self.refresh_tree()
        self.show_equipment_management(equipment_key, department_id)

    def edit_selected_generic_equipment(self, equipment_key, departments, department_id=None):
        asset_id = self._selected_generic_equipment_id()
        if asset_id is None:
            return
        session = SessionLocal()
        try:
            asset = get_equipment_asset(session, equipment_key, asset_id)
        finally:
            session.close()
        dialog = GenericEquipmentDialog(equipment_key.replace("_", " ").title(), departments, asset, self)
        if dialog.exec() != QDialog.Accepted:
            return
        session = SessionLocal()
        try:
            update_equipment_asset(session, equipment_key, asset_id, dialog.get_data())
            session.commit()
        except Exception as exc:
            session.rollback()
            QMessageBox.critical(self, "Error", f"Could not update equipment:\n\n{exc}")
            return
        finally:
            session.close()
        self.refresh_tree()
        self.show_equipment_management(equipment_key, department_id)

    def delete_selected_generic_equipment(self, equipment_key, department_id=None):
        asset_id = self._selected_generic_equipment_id()
        if asset_id is None or QMessageBox.question(self, "Delete Equipment", "Delete the selected equipment?", QMessageBox.Yes | QMessageBox.No) != QMessageBox.Yes:
            return
        session = SessionLocal()
        try:
            delete_equipment_asset(session, equipment_key, asset_id)
            session.commit()
        except Exception as exc:
            session.rollback()
            QMessageBox.critical(self, "Error", f"Could not delete equipment:\n\n{exc}")
            return
        finally:
            session.close()
        self.refresh_tree()
        self.show_equipment_management(equipment_key, department_id)

    def _update_selected_generic_equipment(self, action, equipment_key, department_id=None):
        asset_id = self._selected_generic_equipment_id()
        if asset_id is None:
            return
        session = SessionLocal()
        try:
            action(session, equipment_key, asset_id)
            session.commit()
        except Exception as exc:
            session.rollback()
            QMessageBox.critical(self, "Error", f"Could not update equipment:\n\n{exc}")
            return
        finally:
            session.close()
        self.refresh_tree()
        self.show_equipment_management(equipment_key, department_id)

    def move_selected_generic_equipment_to_stock(self, equipment_key, department_id=None):
        self._update_selected_generic_equipment(move_equipment_asset_to_stock, equipment_key, department_id)

    def reserve_selected_generic_equipment(self, equipment_key, department_id=None):
        self._update_selected_generic_equipment(reserve_equipment_asset, equipment_key, department_id)

    def assign_selected_generic_equipment(self, equipment_key, department_id=None):
        asset_id = self._selected_generic_equipment_id()
        if asset_id is None:
            return
        session = SessionLocal()
        try:
            departments = get_departments(session)
        finally:
            session.close()
        dialog = QDialog(self)
        department_combo = QComboBox()
        user_edit = QLineEdit()
        location_edit = QLineEdit()
        for department in departments:
            department_combo.addItem(department["name"], department["id"])
        form = QFormLayout()
        form.addRow("Department:", department_combo)
        form.addRow("Assigned User:", user_edit)
        form.addRow("Location:", location_edit)
        buttons = QDialogButtonBox(QDialogButtonBox.Save | QDialogButtonBox.Cancel)
        buttons.accepted.connect(dialog.accept)
        buttons.rejected.connect(dialog.reject)
        dialog_layout = QVBoxLayout(dialog)
        dialog_layout.addLayout(form)
        dialog_layout.addWidget(buttons)
        if dialog.exec() != QDialog.Accepted:
            return
        self._update_selected_generic_equipment(
            lambda session, key, selected_id: assign_equipment_asset(
                session, key, selected_id, department_combo.currentData(),
                user_edit.text().strip(), location_edit.text().strip() or "Assigned Location",
            ),
            equipment_key,
            department_id,
        )

    def upload_generic_equipment(self, equipment_key, department_id=None):
        file_path, _ = QFileDialog.getOpenFileName(self, "Select Equipment Excel File", "", "Excel Files (*.xlsx *.xls)")
        if not file_path:
            return
        try:
            dataframe = read_excel_dataframe(file_path)
            result = None
            session = SessionLocal()
            try:
                result = import_equipment_assets(session, equipment_key, dataframe)
                session.commit()
            except Exception:
                session.rollback()
                raise
            finally:
                session.close()
        except Exception as exc:
            QMessageBox.critical(self, "Excel Import Error", f"Could not import equipment:\n\n{exc}")
            return
        QMessageBox.information(self, "Import Complete", f"Imported {result['created']} item(s); skipped {result['skipped']}.")
        self.refresh_tree()
        self.show_equipment_management(equipment_key, department_id)

    # =====================================================
    # STOCK
    # =====================================================

    def show_stock(self):

        widget = QWidget()

        layout = QVBoxLayout(
            widget
        )

        title = QLabel(
            "Stock"
        )

        title.setStyleSheet(
            "font-size: 24px; "
            "font-weight: 600;"
        )

        subtitle = QLabel(
            "Select an enabled equipment type to view items "
            "currently available in IT Storage."
        )

        stock_tree = QTreeWidget()

        stock_tree.setHeaderLabels(
            [
                "Computer",
                "Manufacturer",
                "Model",
                "Serial Number",
                "Location",
                "Status",
            ]
        )

        session = SessionLocal()

        try:

            computers = (
                session.query(
                    Computer
                )
                .filter(
                    Computer.status
                    == ComputerStatus.IN_STOCK.value
                )
                .order_by(
                    Computer.computer_name
                )
                .all()
            )

            for computer in computers:

                item = QTreeWidgetItem(
                    [
                        computer.computer_name,
                        computer.manufacturer
                        or "-",
                        computer.model
                        or "-",
                        computer.serial_number
                        or "-",
                        computer.location
                        or "IT Storage",
                        computer.status,
                    ]
                )

                item.setData(
                    0,
                    Qt.UserRole,
                    computer.id,
                )

                stock_tree.addTopLevelItem(
                    item
                )

            session_telephones = (
                session.query(Telephone)
                .filter(Telephone.status == TelephoneStatus.IN_STOCK.value)
                .order_by(Telephone.asset_tag)
                .all()
            )
            for telephone in session_telephones:
                item = QTreeWidgetItem([
                    telephone.asset_tag,
                    telephone.brand or "-",
                    telephone.model or "-",
                    telephone.serial_number or "-",
                    telephone.location or "IT Storage",
                    telephone.status,
                ])
                item.setData(0, Qt.UserRole, ("telephone", telephone.id))
                stock_tree.addTopLevelItem(item)

        finally:

            session.close()

        def open_stock_item(item, _column):
            item_data = item.data(0, Qt.UserRole)
            if isinstance(item_data, tuple) and item_data[0] == "telephone":
                self.show_telephones()
            else:
                self.show_computer(item_data)

        stock_tree.itemClicked.connect(open_stock_item)

        refresh_button = QPushButton(
            "Refresh"
        )

        refresh_button.clicked.connect(
            self.refresh_stock_view
        )

        layout.addWidget(
            title
        )

        layout.addWidget(
            subtitle
        )

        layout.addWidget(
            refresh_button
        )

        layout.addWidget(
            stock_tree
        )

        self.workspace.addWidget(
            widget
        )

        self.workspace.setCurrentWidget(
            widget
        )

    def show_stock_equipment(self, equipment_key):
        widget = QWidget()
        layout = QVBoxLayout(widget)
        title = QLabel(f"{equipment_key.replace('_', ' ').title()} Stock")
        title.setStyleSheet("font-size: 24px; font-weight: 600;")
        refresh_button = QPushButton("Refresh")
        table = QTreeWidget()
        table.setRootIsDecorated(False)
        table.setAlternatingRowColors(True)
        table.setHeaderLabels(["Asset / Name", "Brand", "Model", "Department", "Status", "Location"])

        session = SessionLocal()
        try:
            if equipment_key == "computers":
                rows = session.query(Computer).filter(
                    Computer.status == ComputerStatus.IN_STOCK.value
                ).order_by(Computer.computer_name).all()
                values = [
                    (row.computer_name, row.manufacturer, row.model,
                     row.department.name if row.department else "Stock",
                     row.status, row.location or "IT Storage")
                    for row in rows
                ]
            elif equipment_key == "phones":
                rows = session.query(Phone).filter(
                    Phone.status == PhoneStatus.IN_STOCK.value
                ).order_by(Phone.asset_tag).all()
                values = [
                    (row.asset_tag, row.brand, row.model,
                     row.department.name if row.department else "Stock",
                     row.status, row.location or "IT Storage")
                    for row in rows
                ]
            elif equipment_key == "telephones":
                rows = session.query(Telephone).filter(
                    Telephone.status == TelephoneStatus.IN_STOCK.value
                ).order_by(Telephone.asset_tag).all()
                values = [
                    (row.asset_tag, row.brand, row.model,
                     row.department.name if row.department else "Stock",
                     row.status, row.location or "IT Storage")
                    for row in rows
                ]
            else:
                values = [
                    (
                        asset["asset_tag"], asset["brand"], asset["model"],
                        asset["department_name"] or "Stock", asset["status"],
                        asset["location"] or "IT Storage",
                    )
                    for asset in get_equipment_assets(
                        session,
                        equipment_key,
                        {"status": "IN_STOCK"},
                    )
                ]
            for value in values:
                table.addTopLevelItem(QTreeWidgetItem([str(item or "-") for item in value]))
        finally:
            session.close()

        refresh_button.clicked.connect(lambda: self.show_stock_equipment(equipment_key))
        layout.addWidget(title)
        layout.addWidget(refresh_button)
        layout.addWidget(table)
        self.workspace.addWidget(widget)
        self.workspace.setCurrentWidget(widget)

    # =====================================================
    # REFRESH STOCK
    # =====================================================

    def refresh_stock_view(self):

        self.refresh_tree()

        self.show_stock()

    # =====================================================
    # COMPUTER DETAILS
    # =====================================================

    def show_computer(
        self,
        computer_id,
    ):

        session = SessionLocal()

        try:

            computer = get_computer_data(
                session,
                computer_id,
            )

        except Exception as exc:

            QMessageBox.warning(
                self,
                "Computer Error",
                str(exc),
            )

            return

        finally:

            session.close()

        widget = QWidget()

        layout = QVBoxLayout(
            widget
        )

        title = QLabel(
            computer[
                "computer_name"
            ]
        )

        title.setStyleSheet(
            "font-size: 26px; "
            "font-weight: 700;"
        )

        subtitle = QLabel(
            "Computer Information"
        )

        # -------------------------------------------------
        # ACTIONS
        # -------------------------------------------------

        actions = QHBoxLayout()

        edit_button = QPushButton(
            "Edit"
        )

        assign_button = QPushButton(
            "Assign"
        )

        stock_button = QPushButton(
            "Move to Stock"
        )

        actions.addWidget(
            edit_button
        )

        actions.addWidget(
            assign_button
        )

        actions.addWidget(
            stock_button
        )

        actions.addStretch()

        edit_button.clicked.connect(
            lambda:
                self.edit_computer_from_details(
                    computer["id"]
                )
        )

        assign_button.clicked.connect(
            lambda:
                self.assign_computer_from_details(
                    computer["id"]
                )
        )

        stock_button.clicked.connect(
            lambda:
                self.move_computer_from_details(
                    computer["id"]
                )
        )

        # -------------------------------------------------
        # INFORMATION
        # -------------------------------------------------

        group = QGroupBox(
            "Computer Information"
        )

        form = QFormLayout(
            group
        )

        fields = [
            (
                "Status",
                "status",
            ),
            (
                "Department",
                "department_name",
            ),
            (
                "Assigned User",
                "assigned_user",
            ),
            (
                "Location",
                "location",
            ),
            (
                "Manufacturer",
                "manufacturer",
            ),
            (
                "Model",
                "model",
            ),
            (
                "Serial Number",
                "serial_number",
            ),
            (
                "Operating System",
                "operating_system",
            ),
            (
                "OS Version",
                "os_version",
            ),
            (
                "Installed Software",
                "installed_software",
            ),
            (
                "CPU",
                "cpu",
            ),
            (
                "RAM",
                "ram",
            ),
            (
                "Storage",
                "storage",
            ),
            (
                "IP Address",
                "ip_address",
            ),
            (
                "MAC Address",
                "mac_address",
            ),
            (
                "Purchase Date",
                "purchase_date",
            ),
            (
                "Warranty",
                "warranty",
            ),
            (
                "Notes",
                "notes",
            ),
        ]

        for label, key in fields:

            value = computer.get(key) or "-"
            if key == "installed_software" and value != "-":
                try:
                    software = json.loads(value)
                    value = ", ".join(
                        str(entry.get("application") or entry.get("name") or "Unknown")
                        if isinstance(entry, dict)
                        else str(entry)
                        for entry in software
                    ) or "-"
                except (TypeError, ValueError):
                    value = str(value)

            form.addRow(
                f"{label}:",
                QLabel(str(value)),
            )

        layout.addWidget(
            title
        )

        layout.addWidget(
            subtitle
        )

        layout.addLayout(
            actions
        )

        layout.addWidget(
            group
        )

        layout.addStretch()

        self.workspace.addWidget(
            widget
        )

        self.workspace.setCurrentWidget(
            widget
        )

    # =====================================================
    # EDIT FROM DETAILS
    # =====================================================

    def edit_computer_from_details(
        self,
        computer_id,
    ):

        session = SessionLocal()

        try:

            computer = get_computer_data(
                session,
                computer_id,
            )

            departments = get_departments(
                session
            )

        finally:

            session.close()

        dialog = ComputerDialog(
            departments,
            computer,
            self,
        )

        if (
            dialog.exec()
            != QDialog.Accepted
        ):
            return

        data = dialog.get_data()

        if not data["computer_name"]:

            QMessageBox.warning(
                self,
                "Invalid Computer",
                "Computer name is required.",
            )

            return

        session = SessionLocal()

        try:

            update_computer(
                session,
                computer_id,
                data,
            )

            session.commit()

        except Exception as exc:

            session.rollback()

            QMessageBox.critical(
                self,
                "Error",
                f"Could not update computer:\n\n{exc}",
            )

            return

        finally:

            session.close()

        self.refresh_tree()

        self.show_computer(
            computer_id
        )

    # =====================================================
    # ASSIGN FROM DETAILS
    # =====================================================

    def assign_computer_from_details(
        self,
        computer_id,
    ):

        self.show_computers()

        self.select_computer_in_table(
            computer_id
        )

        self.assign_selected_computer()

    # =====================================================
    # MOVE FROM DETAILS
    # =====================================================

    def move_computer_from_details(
        self,
        computer_id,
    ):

        confirmation = QMessageBox.question(
            self,
            "Move to Stock",
            "Move this computer back to IT Stock?",
            QMessageBox.Yes
            | QMessageBox.No,
        )

        if (
            confirmation
            != QMessageBox.Yes
        ):
            return

        session = SessionLocal()

        try:

            move_to_stock(
                session,
                computer_id,
            )

            session.commit()

        except Exception as exc:

            session.rollback()

            QMessageBox.critical(
                self,
                "Error",
                f"Could not move computer to stock:\n\n{exc}",
            )

            return

        finally:

            session.close()

        self.refresh_tree()

        self.show_computer(
            computer_id
        )

    # =====================================================
    # SELECT COMPUTER IN TABLE
    # =====================================================

    def select_computer_in_table(
        self,
        computer_id,
    ):

        if not hasattr(
            self,
            "computer_table",
        ):
            return

        for index in range(
            self.computer_table.topLevelItemCount()
        ):

            item = (
                self.computer_table
                .topLevelItem(index)
            )

            if (
                item.data(
                    0,
                    Qt.UserRole,
                )
                == computer_id
            ):

                self.computer_table.setCurrentItem(
                    item
                )

                return
            




    # =====================================================
    # REPORTS
    # =====================================================

    def show_equipment_report(self, equipment_key, stock_only=False):
        session = SessionLocal()
        try:
            equipment_type = next(
                item for item in get_enabled_report_types(session)
                if item.key == equipment_key
            )
            departments = get_departments(session)
        except StopIteration:
            session.close()
            QMessageBox.information(self, "Reports", "This equipment type is disabled.")
            return
        finally:
            if session.is_active:
                session.close()

        widget = QWidget()
        layout = QVBoxLayout(widget)
        title = QLabel(
            f"{equipment_type.name} Stock Report"
            if stock_only else f"{equipment_type.name} Report"
        )
        title.setStyleSheet("font-size: 26px; font-weight: 700;")
        layout.addWidget(title)

        filters_layout = QHBoxLayout()
        search_edit = QLineEdit()
        search_edit.setPlaceholderText("Search asset, brand, model, serial, user, location...")
        department_filter = QComboBox()
        department_filter.addItem("All Departments", None)
        for department in departments:
            department_filter.addItem(department["name"], department["id"])
        status_filter = QComboBox()
        status_filter.addItem("All Statuses", None)
        for status in EQUIPMENT_ASSET_STATUSES:
            status_filter.addItem(status.replace("_", " ").title(), status)
        filters_layout.addWidget(QLabel("Search:"))
        filters_layout.addWidget(search_edit, 1)
        filters_layout.addWidget(QLabel("Department:"))
        filters_layout.addWidget(department_filter)
        filters_layout.addWidget(QLabel("Status:"))
        filters_layout.addWidget(status_filter)
        layout.addLayout(filters_layout)

        table = QTableWidget()
        columns = [
            "Asset Tag", "Name", "Brand", "Model", "Serial Number",
            "IP Address", "MAC Address", "Department", "Assigned User",
            "Location", "Status", "Notes", "Created At", "Updated At",
        ]
        table.setColumnCount(len(columns))
        table.setHorizontalHeaderLabels(columns)
        layout.addWidget(table)
        display_rows = []

        def apply_filters():
            filters = {
                "search": search_edit.text().strip(),
                "department_id": department_filter.currentData(),
                "status": status_filter.currentData(),
            }
            if stock_only:
                filters["status"] = EQUIPMENT_ASSET_STATUSES[0]
            session = SessionLocal()
            try:
                rows = get_equipment_report_rows(session, equipment_key, filters)
            finally:
                session.close()
            display_rows.clear()
            display_rows.extend(rows)
            table.setRowCount(len(rows))
            for row_index, row in enumerate(rows):
                for column_index, column in enumerate(columns):
                    table.setItem(row_index, column_index, QTableWidgetItem(str(row.get(column) or "")))
            table.resizeColumnsToContents()
            table.horizontalHeader().setStretchLastSection(True)

        search_edit.textChanged.connect(apply_filters)
        department_filter.currentIndexChanged.connect(apply_filters)
        status_filter.currentIndexChanged.connect(apply_filters)
        apply_filters()

        buttons = QHBoxLayout()
        for label, callback in (
            ("Export Excel", lambda: self.export_report_excel(display_rows, f"{equipment_key}_report")),
            ("Export CSV", lambda: self.export_report_csv(display_rows, f"{equipment_key}_report")),
            ("Export PDF", lambda: self.export_report_pdf(display_rows, f"{equipment_key}_report", f"{equipment_type.name} Report")),
            ("Export Word", lambda: self.export_report_word(display_rows, f"{equipment_key}_report", f"{equipment_type.name} Report")),
            ("Refresh", apply_filters),
        ):
            button = QPushButton(label)
            button.clicked.connect(callback)
            buttons.addWidget(button)
        buttons.addStretch()
        layout.addLayout(buttons)
        self.workspace.addWidget(widget)
        self.workspace.setCurrentWidget(widget)

    def show_reports(self):

        session = SessionLocal()

        try:
            total_departments = session.query(Department).count()
            equipment_summary = get_enabled_equipment_summary(session)

        finally:
            session.close()

        widget = QWidget()

        layout = QVBoxLayout(
            widget
        )

        title = QLabel(
            "Reports"
        )

        title.setStyleSheet(
            "font-size: 28px; "
            "font-weight: 700;"
        )

        layout.addWidget(
            title
        )

        description = QLabel("Overview of enabled equipment, departments, stock and maintenance information.")

        description.setStyleSheet(
            "font-size: 14px; "
            "color: #777;"
        )

        layout.addWidget(
            description
        )

        summary_group = QGroupBox(
            "System Overview"
        )

        summary_layout = QFormLayout(
            summary_group
        )

        summary_layout.addRow(
            "Total Departments:",
            QLabel(
                str(total_departments)
            )
        )

        for summary in equipment_summary:
            summary_layout.addRow(
                f"{summary['name']}:",
                QLabel(str(summary["total"])),
            )

        layout.addWidget(
            summary_group
        )

        reports_group = QGroupBox(
            "Available Reports"
        )

        reports_layout = QVBoxLayout(
            reports_group
        )

        departments_button = QPushButton(
            "Departments Report"
        )

        departments_button.clicked.connect(
            self.show_departments_report
        )

        equipment_reports_button = QPushButton("Equipment Reports")
        equipment_reports_button.clicked.connect(self.show_reports)

        stock_button = QPushButton(
            "Stock Report"
        )

        stock_button.clicked.connect(
            self.show_stock_report
        )

        maintenance_button = QPushButton(
            "Maintenance Report"
        )

        maintenance_button.clicked.connect(
            self.show_maintenance_report
        )

        reports_layout.addWidget(
            departments_button
        )

        reports_layout.addWidget(equipment_reports_button)

        reports_layout.addWidget(
            stock_button
        )

        reports_layout.addWidget(
            maintenance_button
        )

        layout.addWidget(
            reports_group
        )

        refresh_button = QPushButton(
            "Refresh"
        )

        refresh_button.clicked.connect(
            self.show_reports
        )

        layout.addWidget(
            refresh_button
        )

        layout.addStretch()

        self.workspace.addWidget(
            widget
        )

        self.workspace.setCurrentWidget(
            widget
        )

    # =====================================================
    # DEPARTMENTS REPORT
    # =====================================================

    def show_dynamic_departments_report(self):
        session = SessionLocal()
        try:
            rows = get_enabled_department_summary(session)
            enabled_types = get_enabled_report_types(session)
        finally:
            session.close()
        columns = ["Department"] + [item.name for item in enabled_types]
        widget = QWidget()
        layout = QVBoxLayout(widget)
        title = QLabel("Department Equipment Report")
        title.setStyleSheet("font-size: 26px; font-weight: 700;")
        layout.addWidget(title)
        table = QTableWidget()
        table.setColumnCount(len(columns))
        table.setHorizontalHeaderLabels(columns)
        table.setRowCount(len(rows))
        for row_index, row in enumerate(rows):
            for column_index, column in enumerate(columns):
                table.setItem(row_index, column_index, QTableWidgetItem(str(row.get(column, 0) if column != "Department" else row["Department"])))
        table.resizeColumnsToContents()
        table.horizontalHeader().setStretchLastSection(True)
        layout.addWidget(table)
        buttons = QHBoxLayout()
        for label, callback in (
            ("Export Excel", lambda: self.export_report_excel(rows, "department_equipment_report")),
            ("Export CSV", lambda: self.export_report_csv(rows, "department_equipment_report")),
            ("Export PDF", lambda: self.export_report_pdf(rows, "department_equipment_report", "Department Equipment Report")),
            ("Export Word", lambda: self.export_report_word(rows, "department_equipment_report", "Department Equipment Report")),
            ("Refresh", self.show_dynamic_departments_report),
        ):
            button = QPushButton(label)
            button.clicked.connect(callback)
            buttons.addWidget(button)
        buttons.addStretch()
        layout.addLayout(buttons)
        self.workspace.addWidget(widget)
        self.workspace.setCurrentWidget(widget)

    def show_departments_report(self):

        session = SessionLocal()

        try:

            departments = (
                session.query(
                    Department
                )
                .order_by(
                    Department.name
                )
                .all()
            )

            rows = []

            for department in departments:

                computers = department.computers

                assigned = sum(
                    1 for computer in computers
                    if computer.status
                    == ComputerStatus.ASSIGNED.value
                )

                in_stock = sum(
                    1 for computer in computers
                    if computer.status
                    == ComputerStatus.IN_STOCK.value
                )

                reserved = sum(
                    1 for computer in computers
                    if computer.status
                    == ComputerStatus.RESERVED.value
                )

                maintenance = sum(
                    1 for computer in computers
                    if computer.status
                    == ComputerStatus.MAINTENANCE.value
                )

                retired = sum(
                    1 for computer in computers
                    if computer.status
                    == ComputerStatus.RETIRED.value
                )

                rows.append(
                    {
                        "Department": department.name,
                        "Total": len(computers),
                        "Assigned": assigned,
                        "In Stock": in_stock,
                        "Reserved": reserved,
                        "Maintenance": maintenance,
                        "Retired": retired,
                    }
                )

        finally:
            session.close()

        widget = QWidget()

        layout = QVBoxLayout(
            widget
        )

        title = QLabel(
            "Departments Report"
        )

        title.setStyleSheet(
            "font-size: 26px; "
            "font-weight: 700;"
        )

        layout.addWidget(
            title
        )

        table = QTableWidget()

        table.setColumnCount(7)

        table.setHorizontalHeaderLabels(
            [
                "Department",
                "Total",
                "Assigned",
                "In Stock",
                "Reserved",
                "Maintenance",
                "Retired",
            ]
        )

        table.setRowCount(
            len(rows)
        )

        for row_index, row in enumerate(rows):

            table.setItem(
                row_index,
                0,
                QTableWidgetItem(
                    str(row["Department"])
                )
            )

            table.setItem(
                row_index,
                1,
                QTableWidgetItem(
                    str(row["Total"])
                )
            )

            table.setItem(
                row_index,
                2,
                QTableWidgetItem(
                    str(row["Assigned"])
                )
            )

            table.setItem(
                row_index,
                3,
                QTableWidgetItem(
                    str(row["In Stock"])
                )
            )

            table.setItem(
                row_index,
                4,
                QTableWidgetItem(
                    str(row["Reserved"])
                )
            )

            table.setItem(
                row_index,
                5,
                QTableWidgetItem(
                    str(row["Maintenance"])
                )
            )

            table.setItem(
                row_index,
                6,
                QTableWidgetItem(
                    str(row["Retired"])
                )
            )

        table.resizeColumnsToContents()

        table.horizontalHeader().setStretchLastSection(
            True
        )

        layout.addWidget(
            table
        )

        buttons = QHBoxLayout()

        export_excel = QPushButton(
            "Export Excel"
        )

        export_excel.clicked.connect(
            lambda: self.export_report_excel(
                rows,
                "departments_report"
            )
        )

        export_csv = QPushButton(
            "Export CSV"
        )

        export_csv.clicked.connect(
            lambda: self.export_report_csv(
                rows,
                "departments_report"
            )
        )

        export_pdf = QPushButton(
            "Export PDF"
        )

        export_pdf.clicked.connect(
            lambda: self.export_report_pdf(
                rows,
                "departments_report",
                "Departments Report",
            )
        )

        export_word = QPushButton(
            "Export Word"
        )

        export_word.clicked.connect(
            lambda: self.export_report_word(
                rows,
                "departments_report",
                "Departments Report",
            )
        )

        refresh = QPushButton(
            "Refresh"
        )

        refresh.clicked.connect(
            self.show_departments_report
        )

        buttons.addWidget(
            export_excel
        )

        buttons.addWidget(
            export_csv
        )

        buttons.addWidget(
            export_pdf
        )

        buttons.addWidget(
            export_word
        )

        buttons.addWidget(
            refresh
        )

        buttons.addStretch()

        layout.addLayout(
            buttons
        )

        self.workspace.addWidget(
            widget
        )

        self.workspace.setCurrentWidget(
            widget
        )

    # =====================================================
    # COMPUTERS REPORT
    # =====================================================

    def show_computers_report(self):

        session = SessionLocal()

        try:

            computers = (
                session.query(
                    Computer
                )
                .options(
                    joinedload(
                        Computer.department
                    )
                )
                .order_by(
                    Computer.computer_name
                )
                .all()
            )

            departments = (
                session.query(
                    Department
                )
                .order_by(
                    Department.name
                )
                .all()
            )

        finally:
            session.close()

        widget = QWidget()

        layout = QVBoxLayout(
            widget
        )

        title = QLabel(
            "Computers Report"
        )

        title.setStyleSheet(
            "font-size: 26px; "
            "font-weight: 700;"
        )

        layout.addWidget(
            title
        )

        filters_layout = QHBoxLayout()

        department_filter_label = QLabel(
            "Department:"
        )

        department_filter = QComboBox()

        department_filter.addItem(
            "All Departments",
            None,
        )

        for department in departments:

            department_filter.addItem(
                department.name,
                department.id,
            )

        status_filter_label = QLabel(
            "Status:"
        )

        status_filter = QComboBox()

        status_filter.addItem(
            "All Statuses",
            None,
        )

        for status in [
            ComputerStatus.IN_STOCK.value,
            ComputerStatus.ASSIGNED.value,
            ComputerStatus.RESERVED.value,
            ComputerStatus.MAINTENANCE.value,
            ComputerStatus.RETIRED.value,
        ]:

            status_filter.addItem(
                status.replace("_", " ").title(),
                status,
            )

        filters_layout.addWidget(
            department_filter_label
        )

        filters_layout.addWidget(
            department_filter
        )

        filters_layout.addWidget(
            status_filter_label
        )

        filters_layout.addWidget(
            status_filter
        )

        filters_layout.addStretch()

        layout.addLayout(
            filters_layout
        )

        table = QTableWidget()

        table.setColumnCount(7)

        table.setHorizontalHeaderLabels(
            [
                "Computer Name",
                "Department",
                "Assigned User",
                "Manufacturer",
                "Model",
                "Status",
                "Location",
            ]
        )

        visible_computers = []

        table.cellDoubleClicked.connect(
            lambda row, column:
            self.open_computer_from_report(
                row,
                visible_computers
            )
        )

        layout.addWidget(
            table
        )

        buttons = QHBoxLayout()

        export_excel = QPushButton(
            "Export Excel"
        )

        export_csv = QPushButton(
            "Export CSV"
        )

        export_pdf = QPushButton(
            "Export PDF"
        )

        export_word = QPushButton(
            "Export Word"
        )

        refresh = QPushButton(
            "Refresh"
        )

        display_rows = []

        def apply_filters():

            nonlocal display_rows, visible_computers

            selected_department_id = (
                department_filter.currentData()
            )

            selected_status = (
                status_filter.currentData()
            )

            filtered_computers = []

            for computer in computers:

                if (
                    selected_department_id is not None
                    and computer.department_id
                    != selected_department_id
                ):
                    continue

                if (
                    selected_status is not None
                    and computer.status
                    != selected_status
                ):
                    continue

                filtered_computers.append(
                    computer
                )

            display_rows = []
            visible_computers = filtered_computers

            for computer in filtered_computers:

                department_name = ""

                if computer.department:
                    department_name = (
                        computer.department.name
                    )

                display_rows.append(
                    {
                        "Computer Name":
                            computer.computer_name,
                        "Department":
                            department_name,
                        "Assigned User":
                            computer.assigned_user or "",
                        "Manufacturer":
                            computer.manufacturer or "",
                        "Model":
                            computer.model or "",
                        "Status":
                            computer.status,
                        "Location":
                            computer.location or "",
                    }
                )

            table.setRowCount(
                len(display_rows)
            )

            for row_index, row in enumerate(
                display_rows
            ):

                values = [
                    row["Computer Name"],
                    row["Department"],
                    row["Assigned User"],
                    row["Manufacturer"],
                    row["Model"],
                    row["Status"],
                    row["Location"],
                ]

                for column_index, value in enumerate(
                    values
                ):

                    table.setItem(
                        row_index,
                        column_index,
                        QTableWidgetItem(
                            str(value)
                        )
                    )

            table.resizeColumnsToContents()

            table.horizontalHeader().setStretchLastSection(
                True
            )

        department_filter.currentIndexChanged.connect(
            lambda _index: apply_filters()
        )

        status_filter.currentIndexChanged.connect(
            lambda _index: apply_filters()
        )

        apply_filters()

        export_excel.clicked.connect(
            lambda: self.export_report_excel(
                display_rows,
                "computers_report"
            )
        )

        export_csv.clicked.connect(
            lambda: self.export_report_csv(
                display_rows,
                "computers_report"
            )
        )

        export_pdf.clicked.connect(
            lambda: self.export_report_pdf(
                display_rows,
                "computers_report",
                "Computers Report",
            )
        )

        export_word.clicked.connect(
            lambda: self.export_report_word(
                display_rows,
                "computers_report",
                "Computers Report",
            )
        )

        refresh.clicked.connect(
            self.show_computers_report
        )

        buttons.addWidget(
            export_excel
        )

        buttons.addWidget(
            export_csv
        )

        buttons.addWidget(
            export_pdf
        )

        buttons.addWidget(
            export_word
        )

        buttons.addWidget(
            refresh
        )

        buttons.addStretch()

        layout.addLayout(
            buttons
        )

        self.workspace.addWidget(
            widget
        )

        self.workspace.setCurrentWidget(
            widget
        )

    # =====================================================
    # OPEN COMPUTER FROM REPORT
    # =====================================================

    def open_computer_from_report(
        self,
        row,
        computers,
    ):

        if row < 0 or row >= len(computers):
            return

        computer = computers[row]

        self.show_computer(
            computer.id
        )

    # =====================================================
    # STOCK REPORT
    # =====================================================

    def show_dynamic_stock_report(self):
        session = SessionLocal()
        try:
            summaries = get_enabled_equipment_summary(session)
        finally:
            session.close()
        rows = []
        for summary in summaries:
            rows.append({
                "Equipment Type": summary["name"],
                "Total": summary["total"],
                "In Stock": summary["by_status"].get("IN_STOCK", 0),
                "Assigned": summary["by_status"].get("ASSIGNED", 0),
                "Reserved": summary["by_status"].get("RESERVED", 0),
                "Maintenance": summary["by_status"].get("MAINTENANCE", 0),
            })
        widget = QWidget()
        layout = QVBoxLayout(widget)
        title = QLabel("Stock Report")
        title.setStyleSheet("font-size: 26px; font-weight: 700;")
        layout.addWidget(title)
        table = QTableWidget()
        columns = ["Equipment Type", "Total", "In Stock", "Assigned", "Reserved", "Maintenance"]
        table.setColumnCount(len(columns))
        table.setHorizontalHeaderLabels(columns)
        table.setRowCount(len(rows))
        for row_index, row in enumerate(rows):
            for column_index, column in enumerate(columns):
                table.setItem(row_index, column_index, QTableWidgetItem(str(row[column])))
        table.resizeColumnsToContents()
        table.horizontalHeader().setStretchLastSection(True)
        layout.addWidget(table)
        buttons = QHBoxLayout()
        for label, callback in (
            ("Export Excel", lambda: self.export_report_excel(rows, "stock_report")),
            ("Export CSV", lambda: self.export_report_csv(rows, "stock_report")),
            ("Export PDF", lambda: self.export_report_pdf(rows, "stock_report", "Stock Report")),
            ("Export Word", lambda: self.export_report_word(rows, "stock_report", "Stock Report")),
            ("Refresh", self.show_dynamic_stock_report),
        ):
            button = QPushButton(label)
            button.clicked.connect(callback)
            buttons.addWidget(button)
        buttons.addStretch()
        layout.addLayout(buttons)
        self.workspace.addWidget(widget)
        self.workspace.setCurrentWidget(widget)

    def show_stock_report(self):

        session = SessionLocal()

        try:

            computers = (
                session.query(
                    Computer
                )
                .filter(
                    Computer.status
                    == ComputerStatus.IN_STOCK.value
                )
                .order_by(
                    Computer.computer_name
                )
                .all()
            )

            rows = []

            for computer in computers:

                rows.append(
                    {
                        "Computer Name":
                            computer.computer_name,

                        "Manufacturer":
                            computer.manufacturer or "",

                        "Model":
                            computer.model or "",

                        "Serial Number":
                            computer.serial_number or "",

                        "Operating System":
                            computer.operating_system or "",

                        "Location":
                            computer.location or "",

                        "Status":
                            computer.status,
                    }
                )

        finally:
            session.close()

        widget = QWidget()

        layout = QVBoxLayout(
            widget
        )

        title = QLabel(
            "Stock Report"
        )

        title.setStyleSheet(
            "font-size: 26px; "
            "font-weight: 700;"
        )

        layout.addWidget(
            title
        )

        summary = QLabel(
            f"Available Computers: {len(rows)}"
        )

        summary.setStyleSheet(
            "font-size: 16px; "
            "font-weight: 600;"
        )

        layout.addWidget(
            summary
        )

        table = QTableWidget()

        table.setColumnCount(7)

        table.setHorizontalHeaderLabels(
            [
                "Computer Name",
                "Manufacturer",
                "Model",
                "Serial Number",
                "Operating System",
                "Location",
                "Status",
            ]
        )

        table.setRowCount(
            len(rows)
        )

        for row_index, row in enumerate(rows):

            values = [
                row["Computer Name"],
                row["Manufacturer"],
                row["Model"],
                row["Serial Number"],
                row["Operating System"],
                row["Location"],
                row["Status"],
            ]

            for column_index, value in enumerate(values):

                table.setItem(
                    row_index,
                    column_index,
                    QTableWidgetItem(
                        str(value)
                    )
                )

        table.resizeColumnsToContents()

        table.horizontalHeader().setStretchLastSection(
            True
        )

        layout.addWidget(
            table
        )

        buttons = QHBoxLayout()

        export_excel = QPushButton(
            "Export Excel"
        )

        export_excel.clicked.connect(
            lambda: self.export_report_excel(
                rows,
                "stock_report"
            )
        )

        export_csv = QPushButton(
            "Export CSV"
        )

        export_csv.clicked.connect(
            lambda: self.export_report_csv(
                rows,
                "stock_report"
            )
        )

        export_pdf = QPushButton(
            "Export PDF"
        )

        export_pdf.clicked.connect(
            lambda: self.export_report_pdf(
                rows,
                "stock_report",
                "Stock Report",
            )
        )

        export_word = QPushButton(
            "Export Word"
        )

        export_word.clicked.connect(
            lambda: self.export_report_word(
                rows,
                "stock_report",
                "Stock Report",
            )
        )

        refresh = QPushButton(
            "Refresh"
        )

        refresh.clicked.connect(
            self.show_stock_report
        )

        buttons.addWidget(
            export_excel
        )

        buttons.addWidget(
            export_csv
        )

        buttons.addWidget(
            export_pdf
        )

        buttons.addWidget(
            export_word
        )

        buttons.addWidget(
            refresh
        )

        buttons.addStretch()

        layout.addLayout(
            buttons
        )

        self.workspace.addWidget(
            widget
        )

        self.workspace.setCurrentWidget(
            widget
        )

    # =====================================================
    # MAINTENANCE REPORT
    # =====================================================

    def show_maintenance_report(self):

        session = SessionLocal()

        try:

            records = (
                session.query(
                    Maintenance
                )
                .order_by(
                    Maintenance.recorded_at.desc()
                )
                .all()
            )

            rows = []

            for record in records:

                rows.append(
                    {
                        "Computer":
                            record.computer.computer_name
                            if record.computer
                            else "-",
                        "Problem":
                            record.problem or "",
                        "Solution":
                            record.solution or "",
                        "Status":
                            record.status,
                        "Recorded At":
                            record.recorded_at.strftime(
                                "%Y-%m-%d %H:%M:%S"
                            )
                            if record.recorded_at
                            else "",
                        "Cost":
                            record.cost or 0,
                    }
                )

        finally:
            session.close()

        widget = QWidget()

        layout = QVBoxLayout(
            widget
        )

        title = QLabel(
            "Maintenance Report"
        )

        title.setStyleSheet(
            "font-size: 26px; "
            "font-weight: 700;"
        )

        layout.addWidget(
            title
        )

        summary = QLabel(
            f"Maintenance Records: {len(rows)}"
        )

        summary.setStyleSheet(
            "font-size: 16px; "
            "font-weight: 600;"
        )

        layout.addWidget(
            summary
        )

        table = QTableWidget()

        table.setColumnCount(6)

        table.setHorizontalHeaderLabels(
            [
                "Computer",
                "Problem",
                "Solution",
                "Status",
                "Recorded At",
                "Cost",
            ]
        )

        table.setRowCount(
            len(rows)
        )

        for row_index, row in enumerate(rows):

            values = [
                row["Computer"],
                row["Problem"],
                row["Solution"],
                row["Status"],
                row["Recorded At"],
                str(row["Cost"]),
            ]

            for column_index, value in enumerate(values):

                table.setItem(
                    row_index,
                    column_index,
                    QTableWidgetItem(
                        str(value)
                    )
                )

        table.resizeColumnsToContents()

        table.horizontalHeader().setStretchLastSection(
            True
        )

        layout.addWidget(
            table
        )

        buttons = QHBoxLayout()

        export_excel = QPushButton(
            "Export Excel"
        )

        export_excel.clicked.connect(
            lambda: self.export_report_excel(
                rows,
                "maintenance_report"
            )
        )

        export_csv = QPushButton(
            "Export CSV"
        )

        export_csv.clicked.connect(
            lambda: self.export_report_csv(
                rows,
                "maintenance_report"
            )
        )

        refresh = QPushButton(
            "Refresh"
        )

        refresh.clicked.connect(
            self.show_maintenance_report
        )

        buttons.addWidget(
            export_excel
        )

        buttons.addWidget(
            export_csv
        )

        buttons.addWidget(
            refresh
        )

        buttons.addStretch()

        layout.addLayout(
            buttons
        )

        self.workspace.addWidget(
            widget
        )

        self.workspace.setCurrentWidget(
            widget
        )

    # =====================================================
    # EXPORT EXCEL
    # =====================================================

    def export_report_excel(
        self,
        rows,
        report_name,
    ):

        if not rows:

            QMessageBox.information(
                self,
                "No Data",
                "There is no data to export."
            )

            return

        file_path, _ = QFileDialog.getSaveFileName(
            self,
            "Export Excel Report",
            f"{report_name}.xlsx",
            "Excel Files (*.xlsx)"
        )

        if not file_path:
            return

        try:

            dataframe = pd.DataFrame(
                rows
            )

            dataframe.to_excel(
                file_path,
                index=False
            )

            QMessageBox.information(
                self,
                "Export Complete",
                "The Excel report was exported successfully."
            )

        except Exception as exc:

            QMessageBox.critical(
                self,
                "Export Error",
                f"Failed to export Excel report:\n\n{exc}"
            )

    # =====================================================
    # EXPORT CSV
    # =====================================================

    def export_report_csv(
        self,
        rows,
        report_name,
    ):

        if not rows:

            QMessageBox.information(
                self,
                "No Data",
                "There is no data to export."
            )

            return

        file_path, _ = QFileDialog.getSaveFileName(
            self,
            "Export CSV Report",
            f"{report_name}.csv",
            "CSV Files (*.csv)"
        )

        if not file_path:
            return

        try:

            dataframe = pd.DataFrame(
                rows
            )

            dataframe.to_csv(
                file_path,
                index=False,
                encoding="utf-8-sig"
            )

            QMessageBox.information(
                self,
                "Export Complete",
                "The CSV report was exported successfully."
            )

        except Exception as exc:

            QMessageBox.critical(
                self,
                "Export Error",
                f"Failed to export CSV report:\n\n{exc}"
            )

    # =====================================================
    # EXPORT PDF
    # =====================================================

    def export_report_pdf(
        self,
        rows,
        report_name,
        title,
    ):

        if not rows:

            QMessageBox.information(
                self,
                "No Data",
                "There is no data to export."
            )

            return

        file_path, _ = QFileDialog.getSaveFileName(
            self,
            "Export PDF Report",
            f"{report_name}.pdf",
            "PDF Files (*.pdf)"
        )

        if not file_path:
            return

        try:

            headers = list(rows[0].keys())

            pdf = canvas.Canvas(
                file_path,
                pagesize=letter,
            )

            pdf.setTitle(title)
            pdf.setFont(
                "Helvetica-Bold",
                16,
            )
            pdf.drawString(
                50,
                760,
                title,
            )

            y_position = 730

            pdf.setFont(
                "Helvetica-Bold",
                9,
            )
            pdf.drawString(
                50,
                y_position,
                " | ".join(headers),
            )

            y_position -= 18

            pdf.setFont(
                "Helvetica",
                8,
            )

            for row in rows:

                line = " | ".join(
                    str(row.get(header, ""))
                    for header in headers
                )

                if y_position < 50:

                    pdf.showPage()
                    y_position = 760

                    pdf.setFont(
                        "Helvetica-Bold",
                        9,
                    )
                    pdf.drawString(
                        50,
                        y_position,
                        " | ".join(headers),
                    )

                    y_position -= 18

                    pdf.setFont(
                        "Helvetica",
                        8,
                    )

                pdf.drawString(
                    50,
                    y_position,
                    line,
                )

                y_position -= 14

            pdf.save()

            QMessageBox.information(
                self,
                "Export Complete",
                "The PDF report was exported successfully."
            )

        except Exception as exc:

            QMessageBox.critical(
                self,
                "Export Error",
                f"Failed to export PDF report:\n\n{exc}"
            )

    # =====================================================
    # EXPORT WORD
    # =====================================================

    def export_report_word(
        self,
        rows,
        report_name,
        title,
    ):

        if not rows:

            QMessageBox.information(
                self,
                "No Data",
                "There is no data to export."
            )

            return

        file_path, _ = QFileDialog.getSaveFileName(
            self,
            "Export Word Report",
            f"{report_name}.docx",
            "Word Files (*.docx)"
        )

        if not file_path:
            return

        try:

            document = Document()
            document.add_heading(
                title,
                level=1,
            )

            headers = list(rows[0].keys())

            table = document.add_table(
                rows=1,
                cols=len(headers),
            )

            table.style = "Table Grid"

            for index, header in enumerate(headers):
                table.rows[0].cells[index].text = header

            for row in rows:

                cells = table.add_row().cells

                for index, header in enumerate(headers):
                    cells[index].text = str(
                        row.get(header, "")
                    )

            document.save(file_path)

            QMessageBox.information(
                self,
                "Export Complete",
                "The Word report was exported successfully."
            )

        except Exception as exc:

            QMessageBox.critical(
                self,
                "Export Error",
                f"Failed to export Word report:\n\n{exc}"
            )
    
    
    

    # =====================================================
    # ADMIN MANAGEMENT
    # =====================================================

    def show_admins(self):

        session = SessionLocal()

        try:

            admins = (
                session.query(
                    Admin
                )
                .order_by(
                    Admin.created_at.asc()
                )
                .all()
            )

        finally:
            session.close()

        widget = QWidget()

        layout = QVBoxLayout(
            widget
        )

        title = QLabel(
            "Administrators"
        )

        title.setStyleSheet(
            "font-size: 26px; "
            "font-weight: 700;"
        )

        layout.addWidget(
            title
        )

        buttons_layout = QHBoxLayout()

        add_button = QPushButton(
            "Add Administrator"
        )

        edit_button = QPushButton(
            "Edit Administrator"
        )

        password_button = QPushButton(
            "Change Password"
        )

        delete_button = QPushButton(
            "Delete Administrator"
        )

        refresh_button = QPushButton(
            "Refresh"
        )

        buttons_layout.addWidget(
            add_button
        )

        buttons_layout.addWidget(
            edit_button
        )

        buttons_layout.addWidget(
            password_button
        )

        buttons_layout.addWidget(
            delete_button
        )

        buttons_layout.addWidget(
            refresh_button
        )

        buttons_layout.addStretch()

        layout.addLayout(
            buttons_layout
        )

        self.admin_table = QTableWidget()
        self.admin_table.setColumnCount(3)
        self.admin_table.setHorizontalHeaderLabels(
            [
                "Username",
                "Role",
                "Created At",
            ]
        )

        self.admin_table.setRowCount(
            len(admins)
        )

        for row_index, admin in enumerate(admins):

            username_item = QTableWidgetItem(
                admin.username
            )
            username_item.setData(
                Qt.UserRole,
                admin.id,
            )

            self.admin_table.setItem(
                row_index,
                0,
                username_item,
            )

            self.admin_table.setItem(
                row_index,
                1,
                QTableWidgetItem(
                    admin.role
                ),
            )

            created_at = ""

            if admin.created_at:
                created_at = (
                    admin.created_at.strftime(
                        "%Y-%m-%d %H:%M:%S"
                    )
                )

            self.admin_table.setItem(
                row_index,
                2,
                QTableWidgetItem(
                    created_at
                ),
            )

        self.admin_table.resizeColumnsToContents()
        self.admin_table.horizontalHeader().setStretchLastSection(
            True
        )

        layout.addWidget(
            self.admin_table
        )

        add_button.clicked.connect(
            self.show_add_admin_form
        )

        edit_button.clicked.connect(
            self.edit_selected_admin
        )

        password_button.clicked.connect(
            self.change_selected_admin_password
        )

        delete_button.clicked.connect(
            self.delete_selected_admin
        )

        refresh_button.clicked.connect(
            self.show_admins
        )

        self.workspace.addWidget(
            widget
        )

        self.workspace.setCurrentWidget(
            widget
        )

    def show_add_admin_form(self):

        dialog = QDialog(
            self
        )

        dialog.setWindowTitle(
            "Add Administrator"
        )

        dialog.setMinimumWidth(
            420
        )

        username_edit = QLineEdit()
        password_edit = QLineEdit()
        confirm_password_edit = QLineEdit()

        password_edit.setEchoMode(
            QLineEdit.Password
        )

        confirm_password_edit.setEchoMode(
            QLineEdit.Password
        )

        role_combo = QComboBox()
        role_combo.addItem(
            "ADMIN",
            "ADMIN",
        )
        role_combo.setEnabled(
            False
        )

        form = QFormLayout()

        form.addRow(
            "Username:",
            username_edit,
        )

        form.addRow(
            "Password:",
            password_edit,
        )

        form.addRow(
            "Confirm Password:",
            confirm_password_edit,
        )

        form.addRow(
            "Role:",
            role_combo,
        )

        buttons = QDialogButtonBox(
            QDialogButtonBox.Save
            | QDialogButtonBox.Cancel
        )

        buttons.accepted.connect(
            dialog.accept
        )

        buttons.rejected.connect(
            dialog.reject
        )

        layout = QVBoxLayout(
            dialog
        )

        layout.addLayout(
            form
        )

        layout.addWidget(
            buttons
        )

        if dialog.exec() != QDialog.Accepted:
            return

        username = (
            username_edit.text().strip()
        )

        password = password_edit.text()

        confirm_password = (
            confirm_password_edit.text()
        )

        if not username:

            QMessageBox.warning(
                self,
                "Invalid Administrator",
                "Username is required.",
            )

            return

        if password != confirm_password:

            QMessageBox.warning(
                self,
                "Password Mismatch",
                "Password and confirm password must match.",
            )

            return

        if len(password) < 6:

            QMessageBox.warning(
                self,
                "Invalid Password",
                "Password must be at least 6 characters long.",
            )

            return

        session = SessionLocal()

        try:

            existing_admin = (
                session.query(Admin)
                .filter(
                    Admin.username
                    == username
                )
                .first()
            )

            if existing_admin is not None:

                QMessageBox.warning(
                    self,
                    "Duplicate Username",
                    "An administrator with this username already exists.",
                )

                return

            admin = Admin(
                username=username,
                password_hash=generate_password_hash(
                    password
                ),
                role="ADMIN",
            )

            session.add(
                admin
            )

            session.commit()

        except Exception as exc:

            session.rollback()

            QMessageBox.critical(
                self,
                "Error",
                f"Could not create administrator:\n\n{exc}",
            )

            return

        finally:
            session.close()

        QMessageBox.information(
            self,
            "Administrator Created",
            f"Administrator '{username}' was created successfully.",
        )

        self.show_admins()

    def edit_selected_admin(self):

        if not hasattr(
            self,
            "admin_table",
        ):
            return

        current_item = (
            self.admin_table.currentItem()
        )

        if current_item is None:

            QMessageBox.information(
                self,
                "Select Administrator",
                "Please select an administrator first.",
            )

            return

        admin_id = current_item.data(
            Qt.UserRole
        )

        session = SessionLocal()

        try:

            admin = session.get(
                Admin,
                admin_id,
            )

        finally:
            session.close()

        if admin is None:
            return

        dialog = QDialog(
            self
        )

        dialog.setWindowTitle(
            "Edit Administrator"
        )

        dialog.setMinimumWidth(
            420
        )

        username_edit = QLineEdit()
        username_edit.setText(
            admin.username
        )

        role_combo = QComboBox()
        role_combo.addItem(
            "ADMIN",
            "ADMIN",
        )
        role_combo.setEnabled(
            False
        )

        form = QFormLayout()

        form.addRow(
            "Username:",
            username_edit,
        )

        form.addRow(
            "Role:",
            role_combo,
        )

        buttons = QDialogButtonBox(
            QDialogButtonBox.Save
            | QDialogButtonBox.Cancel
        )

        buttons.accepted.connect(
            dialog.accept
        )

        buttons.rejected.connect(
            dialog.reject
        )

        layout = QVBoxLayout(
            dialog
        )

        layout.addLayout(
            form
        )

        layout.addWidget(
            buttons
        )

        if dialog.exec() != QDialog.Accepted:
            return

        new_username = (
            username_edit.text().strip()
        )

        if not new_username:

            QMessageBox.warning(
                self,
                "Invalid Administrator",
                "Username is required.",
            )

            return

        session = SessionLocal()

        try:

            existing_admin = (
                session.query(Admin)
                .filter(
                    Admin.username
                    == new_username
                )
                .filter(
                    Admin.id
                    != admin_id
                )
                .first()
            )

            if existing_admin is not None:

                QMessageBox.warning(
                    self,
                    "Duplicate Username",
                    "An administrator with this username already exists.",
                )

                return

            admin.username = new_username
            admin.role = "ADMIN"

            session.commit()

        except Exception as exc:

            session.rollback()

            QMessageBox.critical(
                self,
                "Error",
                f"Could not update administrator:\n\n{exc}",
            )

            return

        finally:
            session.close()

        self.show_admins()

    def change_selected_admin_password(self):

        if not hasattr(
            self,
            "admin_table",
        ):
            return

        current_item = (
            self.admin_table.currentItem()
        )

        if current_item is None:

            QMessageBox.information(
                self,
                "Select Administrator",
                "Please select an administrator first.",
            )

            return

        admin_id = current_item.data(
            Qt.UserRole
        )

        session = SessionLocal()

        try:

            admin = session.get(
                Admin,
                admin_id,
            )

        finally:
            session.close()

        if admin is None:
            return

        dialog = QDialog(
            self
        )

        dialog.setWindowTitle(
            "Change Password"
        )

        dialog.setMinimumWidth(
            420
        )

        password_edit = QLineEdit()
        password_edit.setEchoMode(
            QLineEdit.Password
        )

        confirm_password_edit = QLineEdit()
        confirm_password_edit.setEchoMode(
            QLineEdit.Password
        )

        form = QFormLayout()

        form.addRow(
            "New Password:",
            password_edit,
        )

        form.addRow(
            "Confirm Password:",
            confirm_password_edit,
        )

        buttons = QDialogButtonBox(
            QDialogButtonBox.Save
            | QDialogButtonBox.Cancel
        )

        buttons.accepted.connect(
            dialog.accept
        )

        buttons.rejected.connect(
            dialog.reject
        )

        layout = QVBoxLayout(
            dialog
        )

        layout.addLayout(
            form
        )

        layout.addWidget(
            buttons
        )

        if dialog.exec() != QDialog.Accepted:
            return

        password = password_edit.text()
        confirm_password = (
            confirm_password_edit.text()
        )

        if password != confirm_password:

            QMessageBox.warning(
                self,
                "Password Mismatch",
                "Password and confirm password must match.",
            )

            return

        if len(password) < 6:

            QMessageBox.warning(
                self,
                "Invalid Password",
                "Password must be at least 6 characters long.",
            )

            return

        session = SessionLocal()

        try:

            admin = session.get(
                Admin,
                admin_id,
            )

            if admin is None:
                return

            admin.password_hash = generate_password_hash(
                password
            )

            session.commit()

        except Exception as exc:

            session.rollback()

            QMessageBox.critical(
                self,
                "Error",
                f"Could not update password:\n\n{exc}",
            )

            return

        finally:
            session.close()

        QMessageBox.information(
            self,
            "Password Updated",
            "The administrator password was updated successfully.",
        )

    def delete_selected_admin(self):

        if not hasattr(
            self,
            "admin_table",
        ):
            return

        current_item = (
            self.admin_table.currentItem()
        )

        if current_item is None:

            QMessageBox.information(
                self,
                "Select Administrator",
                "Please select an administrator first.",
            )

            return

        admin_id = current_item.data(
            Qt.UserRole
        )

        session = SessionLocal()

        try:

            admin = session.get(
                Admin,
                admin_id,
            )

            total_admins = (
                session.query(Admin)
                .count()
            )

        finally:
            session.close()

        if admin is None:
            return

        if total_admins <= 1:

            QMessageBox.warning(
                self,
                "Cannot Delete Last Administrator",
                "At least one administrator must remain in the system.",
            )

            return

        confirm = QMessageBox.question(
            self,
            "Delete Administrator",
            f"Are you sure you want to delete administrator '{admin.username}'?",
            QMessageBox.Yes
            | QMessageBox.No,
        )

        if confirm != QMessageBox.Yes:
            return

        session = SessionLocal()

        try:

            admin = session.get(
                Admin,
                admin_id,
            )

            if admin is None:
                return

            session.delete(
                admin
            )
            session.commit()

        except Exception as exc:

            session.rollback()

            QMessageBox.critical(
                self,
                "Error",
                f"Could not delete administrator:\n\n{exc}",
            )

            return

        finally:
            session.close()

        self.show_admins()

    # =====================================================
    # OTHER SECTIONS
    # =====================================================

    def show_network_discovery(self):

        widget = QWidget()

        layout = QVBoxLayout(
            widget
        )

        title = QLabel(
            "Network Discovery"
        )

        title.setStyleSheet(
            "font-size: 28px; "
            "font-weight: 700;"
        )

        subtitle = QLabel(
            "Read-only local network scan for discovering active devices and verifying them against the existing inventory."
        )

        subtitle.setWordWrap(
            True
        )

        self.network_discovery_range_edit = QLineEdit(
            get_default_network_range()
        )

        self.network_discovery_range_edit.setPlaceholderText(
            "e.g. 192.168.1.0/24"
        )

        network_info_layout = QHBoxLayout()

        self.network_discovery_adapter_label = QLabel(
            "Active Adapter: N/A"
        )
        self.network_discovery_ipv4_label = QLabel(
            "IPv4: N/A"
        )
        self.network_discovery_subnet_label = QLabel(
            "Subnet: N/A"
        )
        self.network_discovery_network_label = QLabel(
            "Network: N/A"
        )
        self.network_discovery_mac_label = QLabel(
            "MAC: N/A"
        )
        self.network_discovery_gateway_label = QLabel(
            "Gateway: N/A"
        )

        network_info_layout.addWidget(
            self.network_discovery_adapter_label
        )
        network_info_layout.addWidget(
            self.network_discovery_ipv4_label
        )
        network_info_layout.addWidget(
            self.network_discovery_subnet_label
        )
        network_info_layout.addWidget(
            self.network_discovery_network_label
        )
        network_info_layout.addWidget(
            self.network_discovery_mac_label
        )
        network_info_layout.addWidget(
            self.network_discovery_gateway_label
        )

        controls = QHBoxLayout()

        controls_label = QLabel(
            "Network range:"
        )

        scan_button = QPushButton(
            "Scan Network"
        )

        self.network_discovery_stop_button = QPushButton(
            "Stop Scan"
        )
        self.network_discovery_stop_button.setEnabled(
            False
        )

        refresh_button = QPushButton(
            "Refresh Network Information"
        )

        scan_button.clicked.connect(
            lambda:
                self.start_network_scan(
                    self.network_discovery_range_edit.text().strip()
                )
        )

        self.network_discovery_stop_button.clicked.connect(
            self.stop_network_scan
        )

        refresh_button.clicked.connect(
            self._refresh_network_discovery_info
        )

        controls.addWidget(
            controls_label
        )

        controls.addWidget(
            self.network_discovery_range_edit
        )

        controls.addWidget(
            scan_button
        )

        controls.addWidget(
            self.network_discovery_stop_button
        )

        controls.addWidget(
            refresh_button
        )

        controls.addStretch()

        self._refresh_network_discovery_info()

        summary_layout = QHBoxLayout()

        self.network_discovery_summary_labels = {}

        summary_items = [
            ("Verified", "verified"),
            ("Mismatches", "mismatch"),
            ("Offline", "offline"),
            ("New Devices", "new_device"),
            ("Unknown", "unknown"),
            ("Total", "total"),
        ]

        for label_text, summary_key in summary_items:

            summary_label = QLabel(
                f"{label_text}: 0"
            )

            summary_label.setStyleSheet(
                "font-weight: 600;"
            )

            self.network_discovery_summary_labels[
                summary_key
            ] = summary_label

            summary_layout.addWidget(
                summary_label
            )

        self.network_discovery_status_label = QLabel(
            "Ready. No scan has been performed yet."
        )

        progress_layout = QHBoxLayout()

        self.network_discovery_progress_label = QLabel(
            "Progress: 0%"
        )
        self.network_discovery_phase_label = QLabel(
            "Phase: Idle"
        )
        self.network_discovery_devices_found_label = QLabel(
            "Devices Found: 0"
        )

        progress_layout.addWidget(
            self.network_discovery_progress_label
        )
        progress_layout.addWidget(
            self.network_discovery_phase_label
        )
        progress_layout.addWidget(
            self.network_discovery_devices_found_label
        )

        self.network_discovery_table = QTableWidget()
        self.network_discovery_table.setColumnCount(
            8
        )
        self.network_discovery_table.setHorizontalHeaderLabels(
            [
                "Computer Name",
                "Hostname",
                "IP Address",
                "MAC Address",
                "Status",
                "Department",
                "Notes",
                "Inventory Match",
            ]
        )
        self.network_discovery_table.setAlternatingRowColors(
            True
        )
        self.network_discovery_table.setSortingEnabled(
            True
        )
        self.network_discovery_table.setEditTriggers(
            QAbstractItemView.NoEditTriggers
        )

        self.network_discovery_details = QTextEdit()
        self.network_discovery_details.setReadOnly(
            True
        )
        self.network_discovery_details.setMinimumHeight(
            140
        )
        self.network_discovery_table.itemSelectionChanged.connect(
            self._update_network_discovery_details
        )

        self.network_discovery_scan_button = scan_button

        layout.addWidget(
            title
        )

        layout.addWidget(
            subtitle
        )

        layout.addLayout(
            network_info_layout
        )

        layout.addLayout(
            controls
        )

        layout.addLayout(
            summary_layout
        )

        layout.addLayout(
            progress_layout
        )

        layout.addWidget(
            self.network_discovery_status_label
        )

        layout.addWidget(
            self.network_discovery_table
        )

        layout.addWidget(
            QLabel(
                "Selected Device Details"
            )
        )

        layout.addWidget(
            self.network_discovery_details
        )

        layout.addStretch()

        self.workspace.addWidget(
            widget
        )

        self.workspace.setCurrentWidget(
            widget
        )

        self._update_network_discovery_summary(
            {
                "verified": 0,
                "mismatch": 0,
                "offline": 0,
                "new_device": 0,
                "unknown": 0,
                "total": 0,
            }
        )

    def _refresh_network_discovery_info(self):

        active_adapter = get_active_network_adapter()

        if active_adapter:
            self.network_discovery_adapter_label.setText(
                f"Active Adapter: {active_adapter.get('name', 'N/A')}"
            )
            self.network_discovery_ipv4_label.setText(
                f"IPv4: {active_adapter.get('ipv4') or 'N/A'}"
            )
            self.network_discovery_subnet_label.setText(
                f"Subnet: {active_adapter.get('subnet_mask') or 'N/A'}"
            )
            self.network_discovery_network_label.setText(
                f"Network: {active_adapter.get('network') or 'N/A'}"
            )
            self.network_discovery_mac_label.setText(
                f"MAC: {active_adapter.get('mac') or 'N/A'}"
            )
            self.network_discovery_gateway_label.setText(
                f"Gateway: {active_adapter.get('gateway') or 'N/A'}"
            )

            self.network_discovery_range_edit.setText(
                active_adapter.get("network") or get_default_network_range()
            )
        else:
            self.network_discovery_adapter_label.setText(
                "Active Adapter: N/A"
            )
            self.network_discovery_ipv4_label.setText(
                "IPv4: N/A"
            )
            self.network_discovery_subnet_label.setText(
                "Subnet: N/A"
            )
            self.network_discovery_network_label.setText(
                "Network: N/A"
            )
            self.network_discovery_mac_label.setText(
                "MAC: N/A"
            )
            self.network_discovery_gateway_label.setText(
                "Gateway: N/A"
            )

            self.network_discovery_range_edit.setText(
                get_default_network_range()
            )

    def start_network_scan(self, network_range):

        network_range = (network_range or get_default_network_range()).strip()

        if not network_range:
            network_range = get_default_network_range()

        active_adapter = get_active_network_adapter()
        active_adapter_name = active_adapter.get("name", "N/A") if active_adapter else "N/A"

        self.network_discovery_status_label.setText(
            f"Scanning {active_adapter_name} network range {network_range}. Please wait..."
        )
        self.network_discovery_phase_label.setText(
            "Phase: Initializing"
        )
        self.network_discovery_progress_label.setText(
            "Progress: 0%"
        )
        self.network_discovery_devices_found_label.setText(
            "Devices Found: 0"
        )

        self.network_discovery_scan_button.setEnabled(
            False
        )
        self.network_discovery_stop_button.setEnabled(
            True
        )

        thread = QThread()
        worker = NetworkScanWorker(
            network_range
        )

        worker.moveToThread(
            thread
        )

        thread.started.connect(
            worker.run
        )

        worker.progress.connect(
            self._handle_network_discovery_progress
        )

        worker.live_results.connect(
            self._handle_network_discovery_live_results
        )

        worker.finished.connect(
            self._handle_network_discovery_results
        )

        worker.finished.connect(
            thread.quit
        )

        worker.finished.connect(
            worker.deleteLater
        )

        worker.error.connect(
            self._handle_network_discovery_error
        )

        worker.error.connect(
            thread.quit
        )

        worker.error.connect(
            worker.deleteLater
        )

        thread.finished.connect(
            thread.deleteLater
        )

        self.network_scan_thread = thread
        self.network_scan_worker = worker

        thread.start()

    def stop_network_scan(self):

        if hasattr(
            self,
            "network_scan_worker",
        ) and self.network_scan_worker is not None:
            self.network_scan_worker.request_stop()

        self.network_discovery_stop_button.setEnabled(
            False
        )

        self.network_discovery_status_label.setText(
            "Stopping scan..."
        )

    def _handle_network_discovery_progress(
        self,
        phase,
        progress,
        devices_found,
    ):

        self.network_discovery_phase_label.setText(
            f"Phase: {phase}"
        )
        self.network_discovery_progress_label.setText(
            f"Progress: {progress}%"
        )
        self.network_discovery_devices_found_label.setText(
            f"Devices Found: {devices_found}"
        )

    def _handle_network_discovery_live_results(
        self,
        results,
    ):

        self.network_discovery_results = results

        self._render_network_discovery_table(
            results
        )

    def _render_network_discovery_table(
        self,
        results,
    ):

        self.network_discovery_table.setRowCount(
            len(results)
        )
        self.network_discovery_table.setColumnCount(
            8
        )

        for row_index, item in enumerate(results):

            row_data = [
                item.get("computer_name") or "",
                item.get("hostname") or "",
                item.get("ip_address") or "",
                item.get("mac_address") or "",
                item.get("status") or "UNKNOWN",
                item.get("department") or "Unassigned",
                item.get("notes") or "",
                "Yes" if item.get("is_known_record") else "No",
            ]

            for column_index, value in enumerate(row_data):

                table_item = QTableWidgetItem(
                    str(value)
                )
                self.network_discovery_table.setItem(
                    row_index,
                    column_index,
                    table_item,
                )

        if results:
            self.network_discovery_table.setCurrentCell(
                0,
                0,
            )

        self.network_discovery_table.resizeColumnsToContents()
        self._update_network_discovery_details()

    def _update_network_discovery_details(
        self,
    ):

        if not hasattr(
            self,
            "network_discovery_results",
        ):
            self.network_discovery_details.setPlainText(
                ""
            )
            return

        row_index = self.network_discovery_table.currentRow()

        if row_index < 0 or row_index >= len(
            self.network_discovery_results
        ):
            self.network_discovery_details.setPlainText(
                ""
            )
            return

        item = self.network_discovery_results[row_index]

        enrichment_details = item.get("enrichment_details") or {}
        details = (
            f"Discovered Hostname: {item.get('hostname') or 'N/A'}\n"
            f"IP Address: {item.get('ip_address') or 'N/A'}\n"
            f"MAC Address: {item.get('mac_address') or 'N/A'}\n"
            f"Matched Computer: {item.get('matched_computer') or 'None'}\n"
            f"Verification Status: {item.get('status') or 'UNKNOWN'}\n"
            f"Reason: {item.get('reason') or 'No reason provided'}\n"
            f"Device Type: {item.get('device_type') or 'Unknown'}\n"
            f"Enrichment Status: {item.get('enrichment_status') or 'Unavailable'}\n"
            f"Access Status: {item.get('access_status') or 'Unavailable'}\n"
            f"Manufacturer: {item.get('manufacturer') or 'Unknown'}\n"
            f"Model: {item.get('model') or 'Unknown'}\n"
            f"Operating System: {item.get('operating_system') or 'Unknown'}\n"
            f"Windows Version: {item.get('windows_version') or 'Unknown'}\n"
            f"CPU: {item.get('cpu') or 'Unknown'}\n"
            f"RAM: {item.get('ram') or 'Unknown'}\n"
            f"Logged In User: {item.get('logged_in_user') or 'Unknown'}\n"
            f"Domain: {item.get('domain') or 'Unknown'}\n"
            f"Workgroup: {item.get('workgroup') or 'Unknown'}\n"
            f"Installed Software: {', '.join(str(entry.get('application') or entry.get('name') or 'Unknown') for entry in (item.get('installed_software') or [])) if item.get('installed_software') else 'Unknown'}\n\n"
            f"Enrichment Details:\n"
            + (
                "\n".join(
                    f"- {section}: {details.get('status', 'Unknown')} - {details.get('reason', 'No details available')}"
                    for section, details in enrichment_details.items()
                )
                if enrichment_details
                else "- No detailed enrichment diagnostics available."
            )
        )

        self.network_discovery_details.setPlainText(
            details
        )

    def _handle_network_discovery_results(
        self,
        results,
        summary,
    ):

        self.network_discovery_results = results

        self._update_network_discovery_summary(
            summary
        )

        if getattr(
            self,
            "network_scan_worker",
            None,
        ) and self.network_scan_worker.stop_event.is_set():
            self.network_discovery_status_label.setText(
                "Scan stopped."
            )
        else:
            self.network_discovery_status_label.setText(
                f"Scan complete. {summary.get('total', 0)} devices checked."
            )

        self.network_discovery_scan_button.setEnabled(
            True
        )
        self.network_discovery_stop_button.setEnabled(
            False
        )

        self._render_network_discovery_table(
            results
        )

    def _handle_network_discovery_error(
        self,
        message,
    ):

        self.network_discovery_scan_button.setEnabled(
            True
        )
        self.network_discovery_stop_button.setEnabled(
            False
        )

        self.network_discovery_status_label.setText(
            "Scan failed."
        )

        QMessageBox.critical(
            self,
            "Network Discovery Error",
            message,
        )

    def _update_network_discovery_summary(
        self,
        summary,
    ):

        summary = summary or {
            "verified": 0,
            "mismatch": 0,
            "offline": 0,
            "new_device": 0,
            "unknown": 0,
            "total": 0,
        }

        if hasattr(
            self,
            "network_discovery_summary_labels",
        ):

            self.network_discovery_summary_labels["verified"].setText(
                f"Verified: {summary.get('verified', 0)}"
            )
            self.network_discovery_summary_labels["mismatch"].setText(
                f"Mismatches: {summary.get('mismatch', 0)}"
            )
            self.network_discovery_summary_labels["offline"].setText(
                f"Offline: {summary.get('offline', 0)}"
            )
            self.network_discovery_summary_labels["new_device"].setText(
                f"New Devices: {summary.get('new_device', 0)}"
            )
            self.network_discovery_summary_labels["unknown"].setText(
                f"Unknown: {summary.get('unknown', 0)}"
            )
            self.network_discovery_summary_labels["total"].setText(
                f"Total: {summary.get('total', 0)}"
            )

    def show_settings(self):
        widget = QWidget()
        layout = QVBoxLayout(widget)
        title = QLabel("Settings")
        title.setStyleSheet("font-size: 26px; font-weight: 700;")
        layout.addWidget(title)

        group = QGroupBox("Managed Equipment")
        group_layout = QVBoxLayout(group)
        session = SessionLocal()
        try:
            equipment_types = get_equipment_types(session)
        finally:
            session.close()

        checkboxes = []
        for equipment_type in equipment_types:
            checkbox = QCheckBox(equipment_type.name)
            checkbox.setChecked(equipment_type.enabled)
            checkbox.setToolTip(equipment_type.description or "")
            group_layout.addWidget(checkbox)
            checkboxes.append((equipment_type.key, checkbox))

        save_button = QPushButton("Save Managed Equipment")
        add_button = QPushButton("Add Equipment Type")

        def add_equipment_type():
            dialog = QDialog(self)
            dialog.setWindowTitle("Add Equipment Type")
            key_edit = QLineEdit()
            name_edit = QLineEdit()
            description_edit = QLineEdit()
            enabled_check = QCheckBox("Enable immediately")
            form = QFormLayout(dialog)
            form.addRow("Key:", key_edit)
            form.addRow("Name:", name_edit)
            form.addRow("Description:", description_edit)
            form.addRow(enabled_check)
            buttons = QDialogButtonBox(QDialogButtonBox.Save | QDialogButtonBox.Cancel)
            buttons.accepted.connect(dialog.accept)
            buttons.rejected.connect(dialog.reject)
            form.addRow(buttons)
            if dialog.exec() != QDialog.Accepted:
                return
            session = SessionLocal()
            try:
                create_equipment_type(
                    session,
                    key_edit.text(),
                    name_edit.text(),
                    description_edit.text(),
                    enabled_check.isChecked(),
                )
                session.commit()
            except Exception as exc:
                session.rollback()
                QMessageBox.critical(self, "Settings", f"Could not add equipment type:\n\n{exc}")
                return
            finally:
                session.close()
            self.refresh_tree()
            self.show_settings()

        add_button.clicked.connect(add_equipment_type)

        def save_settings():
            session = SessionLocal()
            try:
                for key, checkbox in checkboxes:
                    set_equipment_type_enabled(session, key, checkbox.isChecked())
                session.commit()
            except Exception as exc:
                session.rollback()
                QMessageBox.critical(self, "Settings", f"Could not save settings:\n\n{exc}")
                return
            finally:
                session.close()
            self.refresh_tree()
            self.show_settings()

        save_button.clicked.connect(save_settings)
        layout.addWidget(group)
        layout.addWidget(add_button)
        layout.addWidget(save_button)
        layout.addStretch()
        self.workspace.addWidget(widget)
        self.workspace.setCurrentWidget(widget)

    def show_text(
        self,
        kind,
    ):

        widget = QWidget()

        layout = QVBoxLayout(
            widget
        )

        labels = {

            "maintenance": (
                "Maintenance"
            ),

            "reports": (
                "Reports"
            ),

            "admins": (
                "Admin Management"
            ),

            "settings": (
                "Settings"
            ),
        }

        descriptions = {

            "maintenance": (
                "Maintenance management will be "
                "implemented in the next phase."
            ),

            "reports": (
                "Reports and exports will be "
                "implemented in the next phase."
            ),

            "admins": (
                "Administrator management will be "
                "implemented in the next phase."
            ),

            "settings": (
                "Application settings will be "
                "implemented in the next phase."
            ),
        }

        title = QLabel(
            labels.get(
                kind,
                kind.title(),
            )
        )

        title.setStyleSheet(
            "font-size: 26px; "
            "font-weight: 700;"
        )

        description = QLabel(
            descriptions.get(
                kind,
                "",
            )
        )

        layout.addWidget(
            title
        )

        layout.addWidget(
            description
        )

        layout.addStretch()

        self.workspace.addWidget(
            widget
        )

        self.workspace.setCurrentWidget(
            widget
        )
